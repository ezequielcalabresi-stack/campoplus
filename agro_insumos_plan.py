# -*- coding: utf-8 -*-
"""
CAmpo+ — Insumos necesarios para una campaña planificada.

Toma el manejo real de las últimas campañas (líneas de costo con campo y lote), calcula
cuánto se usó de cada insumo por hectárea sembrada de cada cultivo y lo aplica a las
hectáreas planificadas de cada lote. Compara contra el stock del almacén para saber qué
comprar y lo valoriza al último precio en U$S. El ingeniero puede ajustar dosis,
excluir insumos o agregar nuevos por cultivo; los ajustes quedan guardados por campaña.
"""
from __future__ import annotations

import re
import statistics
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime
from typing import Dict, List, Optional, Tuple

from agro_campania import anio_inicio_codigo_campania, normalizar_cultivo

MESES = ("", "Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic")


def _fold(s) -> str:
    s = unicodedata.normalize("NFKD", str(s or ""))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return " ".join(s.lower().split())


def _clave(s) -> str:
    return " ".join(str(s or "").upper().split())


def _num(v) -> float:
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


def _asegurar_tabla(cur) -> None:
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS plan_insumos_ajustes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            empresa_id INTEGER NOT NULL,
            campania_id INTEGER NOT NULL,
            cultivo TEXT NOT NULL,
            producto_clave TEXT NOT NULL,
            producto TEXT NOT NULL,
            dosis_ha REAL,
            excluido INTEGER DEFAULT 0,
            manual INTEGER DEFAULT 0,
            unidad TEXT,
            usuario TEXT,
            fecha TEXT,
            UNIQUE(empresa_id, campania_id, cultivo, producto_clave)
        );
        """
    )


def campanias_base(cur, empresa_id: int, codigo_objetivo: str, n: int) -> List[str]:
    """Las N campañas con datos inmediatamente anteriores a la planificada (códigos AA-BB válidos)."""
    objetivo = anio_inicio_codigo_campania(codigo_objetivo)
    cods = []
    for (c,) in cur.execute(
        """
        SELECT DISTINCT campania_codigo FROM margenes_access
        WHERE empresa_id = ? AND TRIM(COALESCE(campo,'')) != '' AND TRIM(COALESCE(producto,'')) != '';
        """,
        (empresa_id,),
    ).fetchall():
        m = re.match(r"^(\d{2})-(\d{2})$", c or "")
        if not m or int(m.group(2)) != (int(m.group(1)) + 1) % 100:
            continue
        a = int(m.group(1))
        if objetivo >= 0 and a >= objetivo:
            continue
        cods.append((a, c))
    cods.sort(reverse=True)
    return [c for _, c in cods[: max(1, int(n or 1))]]


def _almacen(cur, empresa_id: int) -> Dict[str, dict]:
    try:
        rows = [dict(r) for r in cur.execute(
            """
            SELECT id, nombre, categoria, unidad, presentacion, stock_cantidad, costo_promedio_usd, unificado_en
            FROM almacen_items WHERE empresa_id = ?;
            """,
            (empresa_id,),
        ).fetchall()]
    except Exception:
        return {}
    por_id = {r["id"]: r for r in rows}
    out = {}
    for r in rows:
        maestro = por_id.get(r.get("unificado_en")) or r
        out.setdefault(_fold(r["nombre"]), maestro)
    return out


def _plan(cur, empresa_id: int, campania_id: int) -> List[dict]:
    return [dict(r) for r in cur.execute(
        """
        SELECT p.lote_id, p.superficie,
               COALESCE(NULLIF(TRIM(p.cultivo_planificado),''), p.cultivo_sugerido) AS cultivo,
               l.nombre AS lote, c.nombre AS campo
        FROM planificacion_lote p
        JOIN lotes_agro l ON l.id = p.lote_id
        JOIN campos_agro c ON c.id = l.campo_id
        WHERE p.campania_id = ? AND c.empresa_id = ?
          AND COALESCE(c.baja,0) = 0 AND COALESCE(l.baja,0) = 0
        ORDER BY c.nombre COLLATE NOCASE, l.orden, l.nombre COLLATE NOCASE;
        """,
        (campania_id, empresa_id),
    ).fetchall()]


def _historial(cur, empresa_id: int, base: List[str]) -> dict:
    marcas = ",".join("?" * len(base))
    filas = [dict(r) for r in cur.execute(
        f"""
        SELECT campania_codigo AS camp, cultivo, campo, lote, cantidad_has, producto, dosis_ha,
               cantidad_total, costo_usd, fecha_aplicacion, nro_orden, laboreo, labor_cultural
        FROM margenes_access
        WHERE empresa_id = ? AND campania_codigo IN ({marcas}) AND TRIM(COALESCE(campo,'')) != '';
        """,
        (empresa_id, *base),
    ).fetchall()]

    labor = {}
    cultivos_lote: Dict[Tuple, Counter] = defaultdict(Counter)
    for f in filas:
        k = (f["camp"], _clave(f["campo"]), _clave(f["lote"]))
        cult = normalizar_cultivo(f["cultivo"])
        if cult:
            cultivos_lote[k][cult] += 1
        if (f["laboreo"] or "").strip():
            labor.setdefault((f["camp"], f["nro_orden"], k[1], k[2]),
                             (f["labor_cultural"] or f["laboreo"] or "").strip().capitalize())

    def cultivo_de(f):
        cult = normalizar_cultivo(f["cultivo"])
        if cult:
            return cult
        cand = cultivos_lote.get((f["camp"], _clave(f["campo"]), _clave(f["lote"])))
        if not cand:
            return ""
        primeros = [c for c, _ in cand.most_common() if c != "Soja 2ª"]
        return primeros[0] if primeros else cand.most_common(1)[0][0]

    area_lote: Dict[Tuple, float] = {}           # (camp, cult, CAMPO, LOTE) -> has
    qty_emp: Dict[Tuple, float] = defaultdict(float)       # (cult, prod) -> cantidad
    qty_lote: Dict[Tuple, float] = defaultdict(float)      # (CAMPO, LOTE, cult, prod) -> cantidad
    lotes_aplic: Dict[Tuple, set] = defaultdict(set)       # (cult, prod) -> {(camp, CAMPO, LOTE)}
    camps_prod: Dict[Tuple, set] = defaultdict(set)        # (cult, prod) -> {camp}
    nombres: Dict[str, Counter] = defaultdict(Counter)
    momentos: Dict[str, Counter] = defaultdict(Counter)
    meses: Dict[str, Counter] = defaultdict(Counter)
    precios: Dict[str, List[Tuple[str, float]]] = defaultdict(list)

    for f in filas:
        cult = cultivo_de(f)
        if not cult:
            continue
        k = (f["camp"], cult, _clave(f["campo"]), _clave(f["lote"]))
        area_lote[k] = max(area_lote.get(k, 0.0), _num(f["cantidad_has"]))
    for f in filas:
        prod = (f["producto"] or "").strip()
        cult = cultivo_de(f)
        if not prod or not cult or len(_fold(prod)) < 2:
            continue
        qty = _num(f["cantidad_total"]) or _num(f["dosis_ha"]) * _num(f["cantidad_has"])
        if qty <= 0:
            continue
        p = _fold(prod)
        campo, lote = _clave(f["campo"]), _clave(f["lote"])
        qty_emp[(cult, p)] += qty
        qty_lote[(campo, lote, cult, p)] += qty
        lotes_aplic[(cult, p)].add((f["camp"], campo, lote))
        camps_prod[(cult, p)].add(f["camp"])
        nombres[p][prod] += 1
        mom = labor.get((f["camp"], f["nro_orden"], campo, lote))
        if mom:
            momentos[p][mom] += 1
        fa = str(f["fecha_aplicacion"] or "")
        if re.match(r"^\d{4}-\d{2}", fa):
            meses[p][int(fa[5:7])] += 1
        if _num(f["costo_usd"]) > 0:
            precios[p].append((f["camp"], _num(f["costo_usd"]) / qty))

    area_cult: Dict[str, float] = defaultdict(float)
    camps_cult: Dict[str, set] = defaultdict(set)
    area_lote_cult: Dict[Tuple, float] = defaultdict(float)   # (CAMPO, LOTE, cult) -> has acumuladas
    for (camp, cult, campo, lote), has in area_lote.items():
        area_cult[cult] += has
        camps_cult[cult].add(camp)
        area_lote_cult[(campo, lote, cult)] += has

    return {
        "area_lote": area_lote, "qty_emp": qty_emp, "qty_lote": qty_lote, "lotes_aplic": lotes_aplic,
        "camps_prod": camps_prod, "nombres": nombres, "momentos": momentos, "meses": meses,
        "precios": precios, "area_cult": area_cult, "camps_cult": camps_cult, "area_lote_cult": area_lote_cult,
    }


def _unidades_compra(cur, empresa_id: int) -> Dict[str, str]:
    cuenta: Dict[str, Counter] = defaultdict(Counter)
    for prod, uni in cur.execute(
        """
        SELECT producto, unidad FROM margenes_access
        WHERE empresa_id = ? AND TRIM(COALESCE(unidad,'')) != '' AND TRIM(COALESCE(producto,'')) != '';
        """,
        (empresa_id,),
    ).fetchall():
        cuenta[_fold(prod)][uni.strip()] += 1
    return {p: c.most_common(1)[0][0] for p, c in cuenta.items()}


def ajustes(cur, empresa_id: int, campania_id: int) -> Dict[Tuple[str, str], dict]:
    _asegurar_tabla(cur)
    return {
        (r["cultivo"], r["producto_clave"]): dict(r)
        for r in cur.execute(
            "SELECT * FROM plan_insumos_ajustes WHERE empresa_id=? AND campania_id=?;",
            (empresa_id, campania_id),
        ).fetchall()
    }


def calcular(cur, empresa_id: int, campania_id: int, n_campanias: int = 3, modo: str = "empresa",
             min_area: float = 0) -> dict:
    row = cur.execute(
        "SELECT codigo FROM campanias_agro WHERE id=? AND empresa_id=?;", (campania_id, empresa_id)
    ).fetchone()
    if not row:
        raise ValueError("Campaña no encontrada")
    codigo = row[0]
    base = campanias_base(cur, empresa_id, codigo, n_campanias)
    plan = _plan(cur, empresa_id, campania_id)
    aj = ajustes(cur, empresa_id, campania_id)
    if not base:
        h = _historial(cur, empresa_id, ["--"])
    else:
        h = _historial(cur, empresa_id, base)
    alm = _almacen(cur, empresa_id)
    uni_compra = _unidades_compra(cur, empresa_id)

    # Dosis por ha sembrada de cada cultivo (promedio de la empresa) y ajustes del ingeniero
    def pct_area(cult, p):
        if not h["area_cult"].get(cult):
            return None
        aplic = sum(h["area_lote"].get((camp, cult, campo, lote), 0)
                    for camp, campo, lote in h["lotes_aplic"].get((cult, p), set()))
        return 100 * aplic / h["area_cult"][cult]

    poco_usados = set()
    ratio_emp: Dict[str, Dict[str, float]] = defaultdict(dict)
    for (cult, p), q in h["qty_emp"].items():
        if not h["area_cult"].get(cult):
            continue
        if min_area and (pct_area(cult, p) or 0) < min_area and (cult, p) not in aj:
            poco_usados.add((cult, p))
            continue
        ratio_emp[cult][p] = q / h["area_cult"][cult]

    def nombre(p):
        if p in h["nombres"]:
            return h["nombres"][p].most_common(1)[0][0]
        manual = next((a["producto"] for (c, k), a in aj.items() if k == p), None)
        return manual or p

    def aplicar_ajustes(cult, ratios: Dict[str, float]) -> Dict[str, float]:
        out = dict(ratios)
        for (c, p), a in aj.items():
            if c != cult:
                continue
            if int(a.get("excluido") or 0):
                out.pop(p, None)
            elif a.get("dosis_ha") is not None:
                out[p] = float(a["dosis_ha"])
        return out

    cultivos_plan = {normalizar_cultivo(r["cultivo"]) for r in plan if (r["cultivo"] or "").strip()}
    lotes_out = []
    total_prod: Dict[str, float] = defaultdict(float)
    por_cult_prod: Dict[Tuple[str, str], float] = defaultdict(float)
    has_cult: Dict[str, float] = defaultdict(float)
    lotes_cult: Dict[str, int] = defaultdict(int)
    sin_cultivo = 0
    fuente_lote = Counter()

    for r in plan:
        cult = normalizar_cultivo(r["cultivo"]) if r["cultivo"] else ""
        has = _num(r["superficie"])
        if not cult or has <= 0:
            sin_cultivo += 1
            continue
        has_cult[cult] += has
        lotes_cult[cult] += 1
        campo, lote = _clave(r["campo"]), _clave(r["lote"])
        fuente = "empresa"
        ratios = ratio_emp.get(cult, {})
        area_l = h["area_lote_cult"].get((campo, lote, cult), 0)
        if modo == "lote" and area_l > 0:
            fuente = "lote"
            ratios = {p: q / area_l for (c1, l1, c2, p), q in h["qty_lote"].items()
                      if c1 == campo and l1 == lote and c2 == cult and (cult, p) not in poco_usados}
        if not ratios and not any(c == cult for c, _ in aj):
            fuente = "sin_historial"
        ratios = aplicar_ajustes(cult, ratios)
        fuente_lote[fuente] += 1
        items = []
        for p, dosis in ratios.items():
            q = dosis * has
            if q <= 0:
                continue
            total_prod[p] += q
            por_cult_prod[(cult, p)] += q
            items.append({"producto": nombre(p), "clave": p, "dosis_ha": round(dosis, 4), "cantidad": round(q, 2)})
        items.sort(key=lambda x: -x["cantidad"])
        lotes_out.append({
            "campo": r["campo"], "lote": r["lote"], "lote_id": r["lote_id"], "cultivo": cult,
            "has": has, "fuente": fuente, "insumos": items,
        })

    def info(p):
        item = alm.get(p)
        precios = h["precios"].get(p) or []
        precio = None
        if precios:
            ultima = max(c for c, _ in precios)
            precio = statistics.median(v for c, v in precios if c == ultima)
        elif item and _num(item.get("costo_promedio_usd")) > 0:
            precio = _num(item["costo_promedio_usd"])
        unidad = uni_compra.get(p) or (item or {}).get("presentacion") or (item or {}).get("unidad") or ""
        manual_uni = next((a.get("unidad") for (c, k), a in aj.items() if k == p and a.get("unidad")), None)
        meses = [MESES[m] for m, _ in h["meses"][p].most_common(2)] if p in h["meses"] else []
        return {
            "categoria": (item or {}).get("categoria") or "",
            "unidad": manual_uni or unidad,
            "stock": _num((item or {}).get("stock_cantidad")) if item else None,
            "almacen_item_id": (item or {}).get("id"),
            "precio_usd": round(precio, 4) if precio else None,
            "momento": h["momentos"][p].most_common(1)[0][0] if p in h["momentos"] else "",
            "meses": "–".join(meses),
        }

    insumos = []
    for p, total in total_prod.items():
        i = info(p)
        stock = i["stock"] or 0
        a_comprar = max(0.0, total - stock)
        insumos.append({
            "clave": p, "producto": nombre(p), **i,
            "total": round(total, 2),
            "a_comprar": round(a_comprar, 2),
            "costo_usd": round(total * i["precio_usd"], 2) if i["precio_usd"] else None,
            "costo_compra_usd": round(a_comprar * i["precio_usd"], 2) if i["precio_usd"] else None,
            "por_cultivo": {c: round(q, 2) for (c, k), q in por_cult_prod.items() if k == p},
        })
    insumos.sort(key=lambda x: (-(x["costo_usd"] or 0), -x["total"]))

    cultivos = []
    for cult in sorted(set(has_cult) | set(cultivos_plan)):
        cult_n = normalizar_cultivo(cult)
        base_ratios = ratio_emp.get(cult_n, {})
        finales = aplicar_ajustes(cult_n, base_ratios)
        filas = []
        for p in sorted(set(base_ratios) | set(finales) | {k for c, k in aj if c == cult_n}):
            a = aj.get((cult_n, p)) or {}
            pa = pct_area(cult_n, p)
            i = info(p)
            filas.append({
                "clave": p, "producto": nombre(p), "unidad": i["unidad"], "categoria": i["categoria"],
                "momento": i["momento"], "meses": i["meses"],
                "dosis_hist": round(base_ratios[p], 4) if p in base_ratios else None,
                "dosis_ha": round(finales[p], 4) if p in finales else None,
                "pct_area": round(pa, 1) if pa is not None else None,
                "campanias_uso": len(h["camps_prod"].get((cult_n, p), ())),
                "ajustado": bool(a) and not int(a.get("excluido") or 0),
                "excluido": bool(int(a.get("excluido") or 0)),
                "manual": bool(int(a.get("manual") or 0)),
                "total": round(por_cult_prod.get((cult_n, p), 0), 2),
                "precio_usd": i["precio_usd"],
            })
        filas.sort(key=lambda x: (x["excluido"], -(x["total"] or 0), x["producto"].lower()))
        costo_ha = sum((f["dosis_ha"] or 0) * (f["precio_usd"] or 0) for f in filas if not f["excluido"])
        cultivos.append({
            "cultivo": cult_n,
            "has_plan": round(has_cult.get(cult_n, 0), 2),
            "lotes": lotes_cult.get(cult_n, 0),
            "has_historicas": round(h["area_cult"].get(cult_n, 0), 1),
            "campanias_hist": sorted(h["camps_cult"].get(cult_n, set()), reverse=True),
            "costo_ha_usd": round(costo_ha, 2),
            "insumos": filas,
        })

    return {
        "campania": codigo,
        "base": base,
        "modo": modo,
        "min_area": min_area,
        "poco_usados": len(poco_usados),
        "lotes_plan": len(plan),
        "lotes_sin_cultivo": sin_cultivo,
        "fuentes": dict(fuente_lote),
        "cultivos": cultivos,
        "insumos": insumos,
        "lotes": lotes_out,
        "totales": {
            "has": round(sum(has_cult.values()), 2),
            "costo_usd": round(sum(i["costo_usd"] or 0 for i in insumos), 2),
            "costo_compra_usd": round(sum(i["costo_compra_usd"] or 0 for i in insumos), 2),
            "sin_precio": sum(1 for i in insumos if not i["precio_usd"]),
        },
    }


def guardar_ajuste(cur, empresa_id: int, campania_id: int, data: dict, usuario: str = "") -> None:
    _asegurar_tabla(cur)
    cultivo = normalizar_cultivo(data.get("cultivo") or "")
    producto = (data.get("producto") or "").strip()
    if not cultivo or not producto:
        raise ValueError("Falta el cultivo o el insumo.")
    clave = (data.get("clave") or "").strip() or _fold(producto)
    dosis = data.get("dosis_ha")
    dosis = None if dosis in (None, "") else float(dosis)
    if dosis is not None and dosis < 0:
        raise ValueError("La dosis no puede ser negativa.")
    cur.execute(
        """
        INSERT INTO plan_insumos_ajustes
            (empresa_id, campania_id, cultivo, producto_clave, producto, dosis_ha, excluido, manual, unidad, usuario, fecha)
        VALUES (?,?,?,?,?,?,?,?,?,?,?)
        ON CONFLICT(empresa_id, campania_id, cultivo, producto_clave) DO UPDATE SET
            producto = excluded.producto, dosis_ha = excluded.dosis_ha, excluido = excluded.excluido,
            manual = MAX(plan_insumos_ajustes.manual, excluded.manual),
            unidad = COALESCE(NULLIF(excluded.unidad,''), plan_insumos_ajustes.unidad),
            usuario = excluded.usuario, fecha = excluded.fecha;
        """,
        (
            empresa_id, campania_id, cultivo, clave, producto, dosis,
            1 if data.get("excluido") else 0, 1 if data.get("manual") else 0,
            (data.get("unidad") or "").strip(), usuario, datetime.now().isoformat(timespec="seconds"),
        ),
    )


def borrar_ajuste(cur, empresa_id: int, campania_id: int, cultivo: str, clave: str) -> int:
    _asegurar_tabla(cur)
    cur.execute(
        "DELETE FROM plan_insumos_ajustes WHERE empresa_id=? AND campania_id=? AND cultivo=? AND producto_clave=?;",
        (empresa_id, campania_id, normalizar_cultivo(cultivo), clave),
    )
    return cur.rowcount


def excel(cur, empresa_id: int, campania_id: int, n_campanias: int = 3, modo: str = "empresa",
          min_area: float = 0) -> bytes:
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    d = calcular(cur, empresa_id, campania_id, n_campanias, modo, min_area)
    wb = Workbook()
    negrita = Font(bold=True)
    fondo = PatternFill("solid", fgColor="E2E8F0")

    def hoja(ws, titulo, sub, cols, filas, anchos):
        ws["A1"] = titulo
        ws["A1"].font = Font(bold=True, size=13)
        ws["A2"] = sub
        for j, t in enumerate(cols, 1):
            c = ws.cell(row=4, column=j, value=t)
            c.font, c.fill = negrita, fondo
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for i, fila in enumerate(filas, 5):
            for j, v in enumerate(fila, 1):
                c = ws.cell(row=i, column=j, value=v)
                if isinstance(v, float):
                    c.number_format = "#,##0.00"
        for j, w in enumerate(anchos, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.freeze_panes = "A5"

    sub = (f"Base: campañas {', '.join(d['base']) or '—'} · {'historial de cada lote' if modo == 'lote' else 'promedio de la empresa por cultivo'}"
           f" · {d['totales']['has']:,.0f} ha planificadas" + (f" · sin insumos usados en menos del {min_area:g}% del área" if min_area else ""))
    cultivos = [c["cultivo"] for c in d["cultivos"]]
    ws = wb.active
    ws.title = "Compras"
    hoja(ws, f"Insumos necesarios · Campaña {d['campania']}", sub,
         ["Insumo", "Categoría", "Unidad", "Momento", "Meses", *cultivos, "Total necesario", "Stock almacén", "A comprar", "Precio U$S", "Costo total U$S", "Costo compra U$S"],
         [[i["producto"], i["categoria"], i["unidad"], i["momento"], i["meses"], *[i["por_cultivo"].get(c) for c in cultivos],
           i["total"], i["stock"], i["a_comprar"], i["precio_usd"], i["costo_usd"], i["costo_compra_usd"]] for i in d["insumos"]],
         [38, 14, 9, 16, 10, *[12] * len(cultivos), 14, 12, 12, 10, 14, 14])
    for c in d["cultivos"]:
        ws = wb.create_sheet(re.sub(r"[\\/*?:\[\]]", "", c["cultivo"])[:28] or "Cultivo")
        hoja(ws, f"{c['cultivo']} · {c['has_plan']:,.0f} ha en {c['lotes']} lotes",
             f"Historial: {c['has_historicas']:,.0f} ha en campañas {', '.join(c['campanias_hist']) or '—'} · costo estimado U$S {c['costo_ha_usd']:,.2f}/ha",
             ["Insumo", "Unidad", "Momento", "Dosis histórica /ha", "Dosis plan /ha", "% del área aplicada", "Campañas en que se usó", "Total", "Ajuste"],
             [[f["producto"], f["unidad"], f["momento"], f["dosis_hist"], f["dosis_ha"], f["pct_area"], f["campanias_uso"], f["total"],
               "Excluido" if f["excluido"] else ("Agregado" if f["manual"] else ("Dosis ajustada" if f["ajustado"] else ""))] for f in c["insumos"]],
             [38, 9, 16, 12, 12, 12, 12, 12, 14])
    ws = wb.create_sheet("Por lote")
    hoja(ws, f"Insumos por lote · Campaña {d['campania']}", sub,
         ["Campo", "Lote", "Cultivo", "Has", "Base", "Insumo", "Dosis /ha", "Cantidad"],
         [[l["campo"], l["lote"], l["cultivo"], l["has"], {"lote": "Historial del lote", "empresa": "Promedio empresa"}.get(l["fuente"], "Sin historial"),
           it["producto"], it["dosis_ha"], it["cantidad"]] for l in d["lotes"] for it in (l["insumos"] or [{"producto": "", "dosis_ha": None, "cantidad": None}])],
         [22, 14, 10, 9, 18, 38, 10, 12])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
