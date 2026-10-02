# -*- coding: utf-8 -*-
"""
CAmpo+ — Almacén agropecuario (insumos + laboreos de terceros).
Los productos ingresan a costo NETO (sin impuestos).
Los laboreos de terceros también se 'almacenan' como unidades de costo
para imputarlos a lotes vía Órdenes de Trabajo.
"""
from __future__ import annotations

import bisect
import sqlite3
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple


def init_almacen_schema(cursor) -> None:
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS almacen_items (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            empresa_id INTEGER DEFAULT 1,
            tipo TEXT NOT NULL DEFAULT 'producto',
            -- producto | laboreo
            codigo TEXT,
            nombre TEXT NOT NULL,
            categoria TEXT,
            -- Semillas, Agroquímicos, Fertilizantes, Combustible, Laboreo
            unidad TEXT DEFAULT 'Kg',
            -- Kg, Lt, Un, Has, Hs
            stock_cantidad REAL DEFAULT 0,
            costo_promedio_neto REAL DEFAULT 0,
            -- siempre sin impuestos
            activo INTEGER DEFAULT 1,
            notas TEXT
        );
    """)
    cursor.execute("""
        CREATE UNIQUE INDEX IF NOT EXISTS idx_almacen_item_empresa_codigo
        ON almacen_items(empresa_id, codigo) WHERE codigo IS NOT NULL AND TRIM(codigo) != '';
    """)

    # Catálogo de categorías (extensible) + flag aplica_ot
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS almacen_categorias (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            codigo TEXT NOT NULL UNIQUE,
            nombre TEXT NOT NULL,
            aplica_ot INTEGER NOT NULL DEFAULT 1,
            orden INTEGER DEFAULT 0,
            activo INTEGER DEFAULT 1
        );
    """)
    cats_seed = [
        ("agroquimicos", "Agroquímicos", 1, 10),
        ("fertilizantes", "Fertilizantes", 1, 20),
        ("semillas", "Semillas", 1, 30),
        ("laboreos", "Laboreos / Servicios", 1, 40),
        ("maquinarias", "Maquinarias (lubricantes, combustibles)", 0, 50),
        ("ganaderia", "Ganadería (vacunas, sanidad)", 0, 60),
        ("varios", "Varios", 0, 90),
    ]
    for codigo, nombre, aplica, orden in cats_seed:
        cursor.execute(
            """
            INSERT OR IGNORE INTO almacen_categorias (codigo, nombre, aplica_ot, orden, activo)
            VALUES (?, ?, ?, ?, 1);
            """,
            (codigo, nombre, aplica, orden),
        )

    # asegurar columna categoria_codigo en items
    cursor.execute("PRAGMA table_info(almacen_items);")
    cols = {r[1] for r in cursor.fetchall()}
    if "categoria_codigo" not in cols:
        cursor.execute("ALTER TABLE almacen_items ADD COLUMN categoria_codigo TEXT;")
    if "clasificacion_manual" not in cols:
        cursor.execute("ALTER TABLE almacen_items ADD COLUMN clasificacion_manual INTEGER DEFAULT 0;")
    _reclasificar_categorias_items(cursor)

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS almacen_movimientos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            empresa_id INTEGER DEFAULT 1,
            item_id INTEGER NOT NULL,
            fecha TEXT NOT NULL,
            tipo_mov TEXT NOT NULL,
            -- ingreso | egreso_ot | ajuste
            cantidad REAL NOT NULL,
            precio_unitario_neto REAL DEFAULT 0,
            -- ingreso: costo unitario sin IVA/impuestos
            importe_neto REAL DEFAULT 0,
            stock_resultante REAL DEFAULT 0,
            costo_prom_resultante REAL DEFAULT 0,
            proveedor_cuit TEXT,
            proveedor_nombre TEXT,
            nro_comprobante TEXT,
            campania_id INTEGER,
            lote_id INTEGER,
            ot_id INTEGER,
            observaciones TEXT,
            FOREIGN KEY(item_id) REFERENCES almacen_items(id)
        );
    """)

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS ordenes_trabajo (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            empresa_id INTEGER DEFAULT 1,
            nro_ot TEXT,
            fecha TEXT NOT NULL,
            tipo_labor TEXT,
            contratista_cuit TEXT,
            contratista_nombre TEXT,
            campania_id INTEGER,
            campo_id INTEGER,
            lote_id INTEGER,
            superficie_has REAL DEFAULT 0,
            cultivo TEXT,
            estado TEXT DEFAULT 'Emitida',
            observaciones TEXT,
            costo_insumos_neto REAL DEFAULT 0,
            costo_laboreos_neto REAL DEFAULT 0,
            costo_total_neto REAL DEFAULT 0
        );
    """)
    cursor.execute("""
        CREATE UNIQUE INDEX IF NOT EXISTS idx_ot_empresa_nro
        ON ordenes_trabajo(empresa_id, nro_ot) WHERE nro_ot IS NOT NULL AND TRIM(nro_ot) != '';
    """)

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS ot_consumos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            ot_id INTEGER NOT NULL,
            item_id INTEGER NOT NULL,
            tipo_item TEXT,
            -- producto | laboreo (denormalizado)
            dosis_por_ha REAL DEFAULT 0,
            cantidad REAL NOT NULL,
            unidad TEXT,
            costo_unitario_neto REAL DEFAULT 0,
            costo_total_neto REAL DEFAULT 0,
            movimiento_id INTEGER,
            FOREIGN KEY(ot_id) REFERENCES ordenes_trabajo(id),
            FOREIGN KEY(item_id) REFERENCES almacen_items(id)
        );
    """)
    cursor.execute("PRAGMA table_info(almacen_items);")
    cols_item = {r[1] for r in cursor.fetchall()}
    if "costo_promedio_usd" not in cols_item:
        cursor.execute("ALTER TABLE almacen_items ADD COLUMN costo_promedio_usd REAL DEFAULT 0;")
    if "presentacion" not in cols_item:
        cursor.execute("ALTER TABLE almacen_items ADD COLUMN presentacion TEXT;")
    if "detalle" not in cols_item:
        cursor.execute("ALTER TABLE almacen_items ADD COLUMN detalle TEXT;")
    if "stock_manual" not in cols_item:
        cursor.execute("ALTER TABLE almacen_items ADD COLUMN stock_manual INTEGER DEFAULT 0;")
    cursor.execute("PRAGMA table_info(almacen_movimientos);")
    cols_mov = {r[1] for r in cursor.fetchall()}
    if "precio_unitario_usd" not in cols_mov:
        cursor.execute("ALTER TABLE almacen_movimientos ADD COLUMN precio_unitario_usd REAL DEFAULT 0;")
    cursor.execute("PRAGMA table_info(ordenes_trabajo);")
    cols_ot = {r[1] for r in cursor.fetchall()}
    if "fecha_aplicacion" not in cols_ot:
        cursor.execute("ALTER TABLE ordenes_trabajo ADD COLUMN fecha_aplicacion TEXT;")
    if "labor_cultural" not in cols_ot:
        cursor.execute("ALTER TABLE ordenes_trabajo ADD COLUMN labor_cultural TEXT;")
    if "confirmada" not in cols_ot:
        cursor.execute("ALTER TABLE ordenes_trabajo ADD COLUMN confirmada INTEGER DEFAULT 0;")
    if "fecha_confirmacion" not in cols_ot:
        cursor.execute("ALTER TABLE ordenes_trabajo ADD COLUMN fecha_confirmacion TEXT;")
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS ot_destinos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            ot_id INTEGER NOT NULL,
            orden INTEGER DEFAULT 0,
            campo_id INTEGER,
            lote_id INTEGER,
            superficie_has REAL DEFAULT 0,
            cultivo TEXT
        );
    """)
    cursor.execute("CREATE INDEX IF NOT EXISTS idx_ot_destinos_ot ON ot_destinos(ot_id);")
    cursor.execute("PRAGMA table_info(ot_consumos);")
    if "destino_id" not in {r[1] for r in cursor.fetchall()}:
        cursor.execute("ALTER TABLE ot_consumos ADD COLUMN destino_id INTEGER;")
    init_lineas_ot_schema(cursor)
    asegurar_schema_unificacion(cursor)
    _aplicar_costos_usd_access(cursor)


def init_lineas_ot_schema(cursor) -> None:
    """margenes_access guarda las líneas de OT (históricas de Access y nuevas)."""
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS margenes_access (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            empresa_id INTEGER DEFAULT 1,
            id_access INTEGER,
            campania_codigo TEXT,
            cultivo TEXT,
            nro_orden INTEGER,
            campo TEXT,
            lote TEXT,
            cantidad_has REAL DEFAULT 0,
            fecha_orden TEXT,
            laboreo TEXT,
            labor_cultural TEXT,
            contratista TEXT,
            producto TEXT,
            tipo TEXT,
            dosis_ha REAL DEFAULT 0,
            cantidad_total REAL DEFAULT 0,
            fecha_aplicacion TEXT,
            precio REAL DEFAULT 0,
            costo_ars REAL DEFAULT 0,
            tc REAL DEFAULT 0,
            costo_usd REAL DEFAULT 0,
            varios TEXT,
            fecha_compra TEXT,
            nro_remito TEXT,
            proveedor TEXT,
            cantidad_ingresada REAL DEFAULT 0,
            unidad TEXT,
            empresa_nombre TEXT,
            UNIQUE(empresa_id, id_access)
        );
        """
    )
    cols = {r[1] for r in cursor.execute("PRAGMA table_info(margenes_access)").fetchall()}
    for col, ddl in (
        ("almacen_item_id", "INTEGER"),
        ("ot_id", "INTEGER"),
        ("ot_consumo_id", "INTEGER"),
        ("moneda_costo", "TEXT"),
        ("origen_costo", "TEXT"),
        ("mov_ingreso_id", "INTEGER"),
        ("fecha_costeo", "TEXT"),
    ):
        if col not in cols:
            cursor.execute(f"ALTER TABLE margenes_access ADD COLUMN {col} {ddl};")
    cursor.execute(
        "CREATE INDEX IF NOT EXISTS idx_margenes_nro_orden ON margenes_access(empresa_id, nro_orden);"
    )


CRITERIOS_COSTO = {
    "peps": "Primero entrado, primero salido",
    "ueps": "Último entrado, primero salido",
    "ppp": "Precio promedio ponderado",
    "ultima": "Precio de la última compra",
}


def criterio_costo_empresa(cursor) -> str:
    try:
        cursor.execute("SELECT criterio_costo FROM configuracion_empresa WHERE id=1;")
        row = cursor.fetchone()
    except sqlite3.OperationalError:
        return "peps"
    if not row:
        return "peps"
    valor = row["criterio_costo"] if isinstance(row, sqlite3.Row) else row[0]
    valor = (valor or "peps").strip().lower()
    return valor if valor in CRITERIOS_COSTO else "peps"


def _serie_tc(cursor, empresa_id: int) -> Tuple[List[str], List[float]]:
    """Tipo de cambio por fecha: el de las líneas costeadas y, si no hay, el de compras con precio en $ y U$S."""
    pares: Dict[str, float] = {}
    cursor.execute(
        """
        SELECT substr(fecha_aplicacion, 1, 10), tc FROM margenes_access
        WHERE empresa_id=? AND COALESCE(tc,0) > 0 AND COALESCE(fecha_aplicacion,'') <> ''
        ORDER BY fecha_aplicacion, id;
        """,
        (empresa_id,),
    )
    for f, tc in cursor.fetchall():
        pares[f] = float(tc)
    cursor.execute(
        """
        SELECT substr(fecha, 1, 10), precio_unitario_neto, precio_unitario_usd FROM almacen_movimientos
        WHERE empresa_id=? AND tipo_mov='ingreso'
          AND COALESCE(precio_unitario_usd,0) > 0 AND COALESCE(precio_unitario_neto,0) > 0;
        """,
        (empresa_id,),
    )
    for f, ars, usd in cursor.fetchall():
        pares.setdefault(f, float(ars) / float(usd))
    fechas = sorted(pares)
    return fechas, [pares[f] for f in fechas]


def _tc_en(serie: Tuple[List[str], List[float]], fecha: str) -> float:
    fechas, tcs = serie
    if not fechas:
        return 0.0
    i = bisect.bisect_right(fechas, (fecha or "")[:10]) - 1
    return tcs[max(i, 0)]


def _movimientos_peps(cursor, empresa_id: int, item_ids: Optional[List[int]] = None) -> Dict[int, Dict[str, list]]:
    """Ingresos y salidas de cada ítem en orden de fecha."""
    filtro, args = "", [empresa_id]
    if item_ids is not None:
        if not item_ids:
            return {}
        filtro = f" AND item_id IN ({','.join('?' * len(item_ids))})"
        args += [int(i) for i in item_ids]
    cursor.execute(
        f"""
        SELECT id, item_id, substr(fecha, 1, 10) AS fecha, tipo_mov, ABS(cantidad) AS qty,
               COALESCE(precio_unitario_neto, 0) AS ars, COALESCE(precio_unitario_usd, 0) AS usd
        FROM almacen_movimientos
        WHERE empresa_id=? AND ABS(COALESCE(cantidad, 0)) > 0
          AND (tipo_mov='ingreso' OR tipo_mov LIKE 'egreso%'){filtro}
        ORDER BY fecha, id;
        """,
        args,
    )
    out: Dict[int, Dict[str, list]] = {}
    for r in cursor.fetchall():
        d = out.setdefault(int(r["item_id"]), {"ingresos": [], "salidas": []})
        (d["ingresos"] if r["tipo_mov"] == "ingreso" else d["salidas"]).append(dict(r))
    return out


def _capas_peps(item: Dict[str, Any], movs: Optional[Dict[str, list]], serie) -> Tuple[List[Dict[str, Any]], float]:
    """Capas de compra en orden de entrada (U$S por unidad) y cantidad a descontar de las más viejas.

    Las compras solo en pesos se pasan a dólares con el tipo de cambio de su fecha. Si el stock real
    supera lo que explican los movimientos (stock inicial traído de Access), esa diferencia es la capa
    más vieja, al costo promedio del ítem; si es menor, se descuenta de las capas más viejas.
    """
    movs = movs or {}
    capas: List[Dict[str, Any]] = []
    for m in movs.get("ingresos", []):
        usd, ars = float(m["usd"] or 0), float(m["ars"] or 0)
        if usd <= 0 < ars:
            tc = _tc_en(serie, m["fecha"])
            usd = ars / tc if tc > 0 else 0.0
        capas.append({"id": m["id"], "fecha": m["fecha"], "qty": float(m["qty"]), "usd": usd, "ars": ars})
    salido = sum(float(s["qty"]) for s in movs.get("salidas", []))
    dif = float(item.get("stock_cantidad") or 0) - (sum(c["qty"] for c in capas) - salido)
    if dif > 1e-6:
        capas.insert(0, {
            "id": 0, "fecha": "", "qty": dif, "inicial": True,
            "usd": float(item.get("costo_promedio_usd") or 0), "ars": float(item.get("costo_promedio_neto") or 0),
        })
    return capas, (-dif if dif < -1e-6 else 0.0)


def _capas_restantes(capas: List[Dict[str, Any]], consumido: float) -> List[Dict[str, Any]]:
    """Capas que quedan después de sacar `consumido` de las más viejas."""
    out = []
    for c in capas:
        qty = c["qty"]
        if consumido >= qty - 1e-9:
            consumido -= qty
            continue
        out.append(dict(c, qty=qty - consumido))
        consumido = 0.0
    return out


def _capas_en_stock(item: Dict[str, Any], movs: Optional[Dict[str, list]], serie):
    """(capas con precio que siguen en stock, compras con precio), en orden de entrada."""
    capas, previo = _capas_peps(item, movs, serie)
    salido = previo + sum(float(s["qty"]) for s in (movs or {}).get("salidas", []))
    quedan = [c for c in _capas_restantes(capas, salido) if c["usd"] > 0]
    compras = [c for c in capas if c["usd"] > 0 and not c.get("inicial")]
    return quedan, compras


def precios_peps(item: Dict[str, Any], movs: Optional[Dict[str, list]], serie) -> Dict[str, Any]:
    """Primero entrado que queda en stock (el próximo en salir por PEPS) y último entrado, en U$S."""
    return _precios_de_capas(*_capas_en_stock(item, movs, serie))


def precios_peps_unificados(cursor, empresa_id: int, items: List[Dict[str, Any]]) -> Dict[str, Any]:
    """Primero y último entrado del producto que resulta de juntar varios ítems."""
    serie = _serie_tc(cursor, empresa_id)
    movs = _movimientos_peps(cursor, empresa_id, [int(i["id"]) for i in items])
    quedan, compras = [], []
    for it in items:
        q, c = _capas_en_stock(it, movs.get(int(it["id"])), serie)
        quedan += q
        compras += c
    orden = lambda c: (c["fecha"], c["id"])
    return _precios_de_capas(sorted(quedan, key=orden), sorted(compras, key=orden))


def _precios_de_capas(quedan: List[Dict[str, Any]], compras: List[Dict[str, Any]]) -> Dict[str, Any]:
    primero = quedan[0] if quedan else None
    ultimo = compras[-1] if compras else None
    return {
        "primero_usd": round(primero["usd"], 4) if primero else None,
        "primero_ars": round(primero["ars"], 4) if primero and primero["ars"] else None,
        "primero_fecha": (primero["fecha"] or "stock inicial") if primero else None,
        "ultimo_usd": round(ultimo["usd"], 4) if ultimo else None,
        "ultimo_ars": round(ultimo["ars"], 4) if ultimo and ultimo["ars"] else None,
        "ultimo_fecha": ultimo["fecha"] if ultimo else None,
    }


def precios_peps_items(cursor, empresa_id: int, items: List[Dict[str, Any]]) -> Dict[int, Dict[str, Any]]:
    serie = _serie_tc(cursor, empresa_id)
    movs = _movimientos_peps(cursor, empresa_id)
    return {int(i["id"]): precios_peps(i, movs.get(int(i["id"])), serie) for i in items}


def valuar_salida_insumo(
    cursor,
    empresa_id: int,
    item_id: int,
    cantidad: float,
    metodo: Optional[str] = None,
    movimiento_id: Optional[int] = None,
    fecha: str = "",
) -> Dict[str, Any]:
    """Costo de una salida de insumo según el criterio de la empresa. Insumos en U$S.

    Por PEPS, la salida `movimiento_id` toma las compras más viejas que quedaban en stock después de
    las salidas anteriores a ella. Sin movimiento se usan las salidas anteriores a `fecha` y, sin fecha,
    todas (la próxima salida).
    """
    cant = float(cantidad or 0)
    metodo = (metodo or criterio_costo_empresa(cursor) or "peps").strip().lower()
    if metodo not in CRITERIOS_COSTO:
        metodo = "peps"
    cursor.execute(
        "SELECT nombre, tipo, costo_promedio_usd, costo_promedio_neto, stock_cantidad FROM almacen_items WHERE id=? AND empresa_id=?;",
        (item_id, empresa_id),
    )
    item = cursor.fetchone()
    if not item:
        raise ValueError("Ítem no encontrado en el almacén.")
    item = dict(item)
    promedio = float(item.get("costo_promedio_usd") or 0)
    es_lab = (item.get("tipo") or "") == "laboreo"
    base = {"criterio": metodo, "criterio_nombre": CRITERIOS_COSTO[metodo], "cantidad": cant}
    if es_lab or metodo == "ppp":
        en_usd = not es_lab and promedio > 0
        unit = promedio if en_usd else float(item.get("costo_promedio_neto") or 0)
        return dict(base, moneda="USD" if en_usd else "ARS", costo_unitario=round(unit, 4),
                    costo_total=round(cant * unit, 2), capas=[])

    serie = _serie_tc(cursor, empresa_id)
    movs = _movimientos_peps(cursor, empresa_id, [item_id]).get(int(item_id), {"ingresos": [], "salidas": []})
    capas, previo = _capas_peps(item, movs, serie)
    compras = [c for c in capas if c["usd"] > 0 and not c.get("inicial")]
    respaldo = promedio or (compras[-1]["usd"] if compras else 0.0)
    if metodo == "ultima":
        unit = compras[-1]["usd"] if compras else promedio
        capas_u = [{"fecha": compras[-1]["fecha"], "cantidad": cant, "costo_unitario": round(unit, 4)}] if compras else []
        return dict(base, moneda="USD", costo_unitario=round(unit, 4), costo_total=round(cant * unit, 2), capas=capas_u)

    salidas = movs["salidas"]
    ref = next((s for s in salidas if movimiento_id and int(s["id"]) == int(movimiento_id)), None)
    if ref:
        antes = [s for s in salidas if (s["fecha"], s["id"]) < (ref["fecha"], ref["id"])]
    elif fecha:
        antes = [s for s in salidas if s["fecha"] < fecha[:10]]
    else:
        antes = salidas
    consumido = previo + sum(float(s["qty"]) for s in antes)
    disponibles = _capas_restantes(capas, consumido)
    if metodo == "ueps":
        disponibles.reverse()
    tomar, total, usadas = cant, 0.0, []
    for c in disponibles:
        if tomar <= 1e-9:
            break
        uso = min(tomar, c["qty"])
        usd = c["usd"] or respaldo
        total += uso * usd
        usadas.append({
            "fecha": c["fecha"] or "stock inicial", "cantidad": round(uso, 4), "costo_unitario": round(usd, 4),
            "pesos": round(c["ars"], 4) if c["ars"] and not c.get("inicial") else None,
        })
        tomar -= uso
    if tomar > 1e-6:
        total += tomar * respaldo
        usadas.append({"fecha": "sin compra registrada", "cantidad": round(tomar, 4), "costo_unitario": round(respaldo, 4), "pesos": None})
    unit = (total / cant) if cant else 0
    return dict(base, moneda="USD", costo_unitario=round(unit, 4), costo_total=round(total, 2), capas=usadas)


def _aplicar_costos_usd_access(cursor) -> None:
    """Completa el costo en dólares de los productos que vinieron de Access.

    El promedio en pesos sigue en costo_promedio_neto. La OT de insumos usa
    esta moneda constante y no el valor histórico en pesos.
    """
    import json
    from pathlib import Path

    path = Path(__file__).with_name("datos") / "costos_usd_almacen.json"
    if not path.exists():
        return
    try:
        mapa = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return
    if not isinstance(mapa, dict) or not mapa:
        return
    for nombre, usd in mapa.items():
        try:
            valor = float(usd or 0)
        except (TypeError, ValueError):
            continue
        if valor <= 0 or not str(nombre).strip():
            continue
        cursor.execute(
            """
            UPDATE almacen_items
            SET costo_promedio_usd = ?
            WHERE UPPER(TRIM(nombre)) = UPPER(TRIM(?))
              AND COALESCE(costo_promedio_usd, 0) = 0;
            """,
            (round(valor, 4), str(nombre).strip()),
        )
    capas = Path(__file__).with_name("datos") / "capas_usd_almacen.json"
    if not capas.exists():
        return
    try:
        filas = json.loads(capas.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return
    if not isinstance(filas, list):
        return
    for fila in filas:
        if not isinstance(fila, (list, tuple)) or len(fila) < 2:
            continue
        try:
            ida = int(fila[0])
            usd = float(fila[1] or 0)
        except (TypeError, ValueError):
            continue
        if ida <= 0 or usd <= 0:
            continue
        cursor.execute(
            """
            UPDATE almacen_movimientos
            SET precio_unitario_usd = ?
            WHERE id_access = ? AND COALESCE(precio_unitario_usd, 0) = 0;
            """,
            (round(usd, 4), ida),
        )


def _inferir_categoria_codigo(nombre: str = "", tipo: str = "", categoria: str = "") -> str:
    """Heurística para mapear ítems legacy a códigos de catálogo."""
    n = f"{nombre or ''} {categoria or ''}".lower()
    t = (tipo or "").strip().lower()
    cat = (categoria or "").strip().lower()

    known = {
        "agroquimicos", "fertilizantes", "semillas", "laboreos",
        "maquinarias", "ganaderia",
    }
    if cat in known:
        return cat
    # "varios" / "insumo" no son definitivos: seguir por nombre
    if cat == "varios":
        cat = ""
        n = f"{nombre or ''}".lower()
    if "fertiliz" in cat:
        return "fertilizantes"
    if "semill" in cat:
        return "semillas"
    if "laboreo" in cat or "servicio" in cat:
        return "laboreos"
    if "maquin" in cat or "lubric" in cat or "combust" in cat:
        return "maquinarias"
    if "ganader" in cat or "vacun" in cat or "sanidad" in cat:
        return "ganaderia"
    if "agroqu" in cat:
        return "agroquimicos"

    if t == "laboreo":
        return "laboreos"

    # combustibles / lubricantes (maquinaria)
    if any(k in n for k in (
        "gas oil", "gasoil", "diesel", "combustible", "lubricante",
        "engras", "nafta", "fuel oil", "aceite hidraul", "aceite hidrául",
        "aceite motor", "grasa ",
    )):
        return "maquinarias"
    # ganadería
    if any(k in n for k in (
        "vacuna", "ivermectin", "caravana", "sanidad animal", "antiparas",
        "afthosa", "aftosa", "brucel", "ganader", "tambo", "reproductor",
    )):
        return "ganaderia"
    # fertilizantes
    if any(k in n for k in (
        "urea", "fertiliz", "fosfato", " map", "map ", " dap", "dap ",
        "sulfato", "nitrato", "amon", "npk", "superfosfato", "fosforo",
        "fósforo", "potasio", "mezcla fisica", "mezcla física",
    )):
        return "fertilizantes"
    # semillas
    if any(k in n for k in ("semilla", "seed", "hibrido", "híbrido", "bolsa de")):
        return "semillas"
    # agroquímicos / coadyuvantes de pulverización
    if any(k in n for k in (
        "glifo", "herbicid", "insecticid", "fungicid", "agroquim", "power plus",
        "2,4", "2.4", "atrazina", "dicamba", "paraquat", "insect", "fungi",
        "coadyuvan", "surfactan", "desecante", "spray", "pulveriz",
        "abamectin", "lambda", "cipermetr", "imidaclop", "clethodim",
        "haloxifop", "quizalofop", "metsulfuron", "diclosulam", "fomesafen",
        "bentazon", "picloram", "fluroxipir", "clorpirif", "dimetoato",
        "aceite rizo", "aceite mineral", "aceite agric", "aceite agrí",
    )):
        return "agroquimicos"
    if any(k in n for k in ("siembra", "cosecha", "laboreo", "rastreo", "arada")):
        return "laboreos"
    # default productos de almacén agro → agroquímicos (OT sí); el usuario puede corregir
    if t == "producto" or not t:
        return "agroquimicos"
    return "varios"


def _nombre_categoria(cursor, codigo: str) -> str:
    cursor.execute(
        "SELECT nombre FROM almacen_categorias WHERE codigo=? LIMIT 1;",
        (codigo,),
    )
    row = cursor.fetchone()
    if row:
        return row["nombre"] if isinstance(row, dict) or hasattr(row, "keys") else row[0]
    return codigo or ""


def _reclasificar_categorias_items(cursor, forzar: bool = False) -> None:
    """Asigna categoria_codigo a ítems sin código válido (o a todos si forzar)."""
    cursor.execute("SELECT codigo FROM almacen_categorias WHERE COALESCE(activo,1)=1;")
    valid = {r["codigo"] for r in cursor.fetchall()}
    if not valid:
        return
    if forzar:
        cursor.execute(
            "SELECT id, nombre, tipo, categoria, categoria_codigo, clasificacion_manual FROM almacen_items;"
        )
    else:
        placeholders = ",".join("?" * len(valid))
        cursor.execute(
            f"""
            SELECT id, nombre, tipo, categoria, categoria_codigo, clasificacion_manual
            FROM almacen_items
            WHERE COALESCE(categoria_codigo,'') = ''
               OR categoria_codigo NOT IN ({placeholders})
               OR categoria_codigo = 'varios'
            """,
            tuple(valid),
        )
    rows = cursor.fetchall()
    for r in rows:
        d = dict(r)
        # no pisar categorías explícitas distintas de vacías/varios NI edición manual
        prev = (d.get("categoria_codigo") or "").strip()
        if int(d.get("clasificacion_manual") or 0) == 1 and not forzar:
            continue
        if prev and prev != "varios" and prev in valid and not forzar:
            continue
        codigo = _inferir_categoria_codigo(
            d.get("nombre") or "", d.get("tipo") or "", d.get("categoria") or ""
        )
        if codigo not in valid:
            codigo = "agroquimicos" if (d.get("tipo") or "producto") == "producto" else "varios"
        # si el ítem ya estaba en varios y la inferencia sigue en varios, ok
        nombre_cat = _nombre_categoria(cursor, codigo)
        cursor.execute(
            """
            UPDATE almacen_items
            SET categoria_codigo=?, categoria=?
            WHERE id=?;
            """,
            (codigo, nombre_cat, d["id"]),
        )


def listar_categorias_almacen(cursor, solo_aplica_ot: Optional[bool] = None) -> List[Dict[str, Any]]:
    q = "SELECT * FROM almacen_categorias WHERE COALESCE(activo,1)=1"
    params: List[Any] = []
    if solo_aplica_ot is True:
        q += " AND COALESCE(aplica_ot,0)=1"
    elif solo_aplica_ot is False:
        q += " AND COALESCE(aplica_ot,0)=0"
    q += " ORDER BY orden ASC, nombre COLLATE NOCASE ASC;"
    cursor.execute(q, params)
    return [dict(r) for r in cursor.fetchall()]


def crear_categoria_almacen(
    cursor, *, codigo: str, nombre: str, aplica_ot: bool = False, orden: int = 100
) -> Dict[str, Any]:
    codigo = (codigo or "").strip().lower().replace(" ", "_")
    nombre = (nombre or "").strip()
    if not codigo or not nombre:
        raise ValueError("Código y nombre de categoría son obligatorios.")
    cursor.execute(
        """
        INSERT INTO almacen_categorias (codigo, nombre, aplica_ot, orden, activo)
        VALUES (?, ?, ?, ?, 1);
        """,
        (codigo, nombre, 1 if aplica_ot else 0, orden or 100),
    )
    cursor.execute("SELECT * FROM almacen_categorias WHERE id=?;", (cursor.lastrowid,))
    return dict(cursor.fetchone())


def actualizar_item_categoria(
    cursor, empresa_id: int, item_id: int, categoria_codigo: str
) -> Dict[str, Any]:
    return actualizar_item_clasificacion(
        cursor, empresa_id, item_id, categoria_codigo=categoria_codigo
    )


def actualizar_item_clasificacion(
    cursor,
    empresa_id: int,
    item_id: int,
    *,
    categoria_codigo: Optional[str] = None,
    tipo: Optional[str] = None,
    unidad: Optional[str] = None,
) -> Dict[str, Any]:
    """Edición manual de tipo/categoría/unidad. Marca clasificacion_manual=1."""
    cursor.execute(
        "SELECT * FROM almacen_items WHERE id=? AND empresa_id=?;",
        (item_id, empresa_id),
    )
    item = cursor.fetchone()
    if not item:
        raise ValueError("Ítem no encontrado.")
    item = dict(item)

    nuevo_tipo = (tipo or item.get("tipo") or "producto").strip().lower()
    if nuevo_tipo not in ("producto", "laboreo"):
        nuevo_tipo = "producto"

    cat_cod = (categoria_codigo if categoria_codigo is not None else item.get("categoria_codigo") or "").strip().lower()
    if not cat_cod:
        cat_cod = _inferir_categoria_codigo(item.get("nombre") or "", nuevo_tipo, item.get("categoria") or "")
    cursor.execute(
        "SELECT * FROM almacen_categorias WHERE codigo=? AND COALESCE(activo,1)=1;",
        (cat_cod,),
    )
    cat = cursor.fetchone()
    if not cat:
        raise ValueError(f"Categoría '{cat_cod}' no encontrada.")
    cat = dict(cat)

    nueva_unidad = (unidad if unidad is not None else item.get("unidad") or "").strip()
    if not nueva_unidad:
        nueva_unidad = "Has" if nuevo_tipo == "laboreo" else (item.get("unidad") or "Kg")

    # asegurar columna clasificacion_manual
    cursor.execute("PRAGMA table_info(almacen_items);")
    cols = {r[1] for r in cursor.fetchall()}
    if "clasificacion_manual" not in cols:
        cursor.execute("ALTER TABLE almacen_items ADD COLUMN clasificacion_manual INTEGER DEFAULT 0;")

    cursor.execute(
        """
        UPDATE almacen_items
        SET tipo=?, categoria_codigo=?, categoria=?, unidad=?, clasificacion_manual=1
        WHERE id=? AND empresa_id=?;
        """,
        (nuevo_tipo, cat_cod, cat["nombre"], nueva_unidad, item_id, empresa_id),
    )
    cursor.execute(
        "SELECT * FROM almacen_items WHERE id=? AND empresa_id=?;",
        (item_id, empresa_id),
    )
    return dict(cursor.fetchone())


def guardar_ficha_item(
    cursor,
    empresa_id: int,
    item_id: int,
    *,
    nombre: str,
    presentacion: str = "",
    detalle: str = "",
    tipo: Optional[str] = None,
    categoria_codigo: Optional[str] = None,
    unidad: Optional[str] = None,
    stock_cantidad: float = 0,
    costo_usd: float = 0,
    costo_ars: float = 0,
    tipo_cambio: float = 0,
) -> Dict[str, Any]:
    """Ficha del producto: nombre, presentación, stock y costos en U$S y pesos."""
    cursor.execute(
        "SELECT * FROM almacen_items WHERE id=? AND empresa_id=?;",
        (item_id, empresa_id),
    )
    item = cursor.fetchone()
    if not item:
        raise ValueError("Ítem no encontrado.")
    item = dict(item)
    nombre = (nombre or "").strip()
    if not nombre:
        raise ValueError("El nombre es obligatorio.")
    usd = round(float(costo_usd or 0), 4)
    ars = round(float(costo_ars or 0), 4)
    tc = round(float(tipo_cambio or 0), 4)
    if tc > 0 and usd > 0 and ars <= 0:
        ars = round(usd * tc, 4)
    elif tc > 0 and ars > 0 and usd <= 0:
        usd = round(ars / tc, 4)
    elif usd > 0 and ars > 0 and tc <= 0:
        tc = round(ars / usd, 4)

    actualizar_item_clasificacion(
        cursor, empresa_id, item_id,
        categoria_codigo=categoria_codigo, tipo=tipo, unidad=unidad,
    )
    stock_nuevo = round(float(stock_cantidad or 0), 4)
    stock_viejo = round(float(item.get("stock_cantidad") or 0), 4)
    manual = int(item.get("stock_manual") or 0)
    if abs(stock_nuevo - stock_viejo) > 0.0001:
        manual = 1
        delta = round(stock_nuevo - stock_viejo, 4)
        cursor.execute("PRAGMA table_info(almacen_movimientos);")
        if "precio_unitario_usd" not in {r[1] for r in cursor.fetchall()}:
            cursor.execute("ALTER TABLE almacen_movimientos ADD COLUMN precio_unitario_usd REAL DEFAULT 0;")
        cursor.execute(
            """
            INSERT INTO almacen_movimientos (
                empresa_id, item_id, fecha, tipo_mov, cantidad,
                precio_unitario_neto, precio_unitario_usd, importe_neto,
                stock_resultante, costo_prom_resultante, observaciones
            ) VALUES (?, ?, ?, 'ajuste', ?, ?, ?, ?, ?, ?, ?);
            """,
            (
                empresa_id, item_id, datetime.now().strftime("%Y-%m-%d"),
                delta, ars, usd, round(abs(delta) * ars, 2),
                stock_nuevo, usd or ars,
                "Corrección manual de stock",
            ),
        )
    cursor.execute("PRAGMA table_info(almacen_items);")
    cols = {r[1] for r in cursor.fetchall()}
    if "presentacion" not in cols:
        cursor.execute("ALTER TABLE almacen_items ADD COLUMN presentacion TEXT;")
    if "detalle" not in cols:
        cursor.execute("ALTER TABLE almacen_items ADD COLUMN detalle TEXT;")
    if "stock_manual" not in cols:
        cursor.execute("ALTER TABLE almacen_items ADD COLUMN stock_manual INTEGER DEFAULT 0;")
    cursor.execute(
        """
        UPDATE almacen_items
        SET nombre=?, presentacion=?, detalle=?,
            stock_cantidad=?, stock_manual=?,
            costo_promedio_usd=?, costo_promedio_neto=?
        WHERE id=? AND empresa_id=?;
        """,
        (nombre, (presentacion or "").strip(), (detalle or "").strip(),
         stock_nuevo, manual, usd, ars, item_id, empresa_id),
    )
    cursor.execute("SELECT * FROM almacen_items WHERE id=? AND empresa_id=?;", (item_id, empresa_id))
    row = dict(cursor.fetchone())
    row["tipo_cambio"] = tc
    return row


def listar_items_almacen(
    cursor,
    empresa_id: int,
    tipo: Optional[str] = None,
    solo_con_stock: bool = False,
    orden: str = "nombre",
    categoria_codigo: Optional[str] = None,
    solo_aplica_ot: bool = False,
) -> List[Dict[str, Any]]:
    q = """
        SELECT i.*,
               c.nombre AS categoria_nombre,
               COALESCE(c.aplica_ot, 0) AS aplica_ot
        FROM almacen_items i
        LEFT JOIN almacen_categorias c ON c.codigo = i.categoria_codigo
        WHERE i.empresa_id = ? AND COALESCE(i.activo,1)=1
    """
    params: List[Any] = [empresa_id]
    if tipo in ("producto", "laboreo"):
        q += " AND i.tipo = ?"
        params.append(tipo)
    if solo_con_stock:
        q += " AND ABS(COALESCE(i.stock_cantidad,0)) > 0.0001"
    if categoria_codigo:
        q += " AND i.categoria_codigo = ?"
        params.append(categoria_codigo.strip().lower())
    if solo_aplica_ot:
        q += " AND COALESCE(c.aplica_ot,0)=1"
    orden_n = (orden or "nombre").strip().lower()
    if orden_n in ("stock", "cantidad"):
        q += " ORDER BY ABS(COALESCE(i.stock_cantidad,0)) DESC, i.nombre COLLATE NOCASE ASC;"
    else:
        # default: alfabético
        q += " ORDER BY i.nombre COLLATE NOCASE ASC, i.tipo ASC;"
    cursor.execute(q, params)
    out = []
    for r in cursor.fetchall():
        d = dict(r)
        # preferir nombre del catálogo si existe
        if d.get("categoria_nombre"):
            d["categoria"] = d["categoria_nombre"]
        out.append(d)
    return out


def _recalcular_promedio(stock_prev: float, costo_prev: float, cant_in: float, precio_in: float) -> float:
    """Promedio ponderado móvil (solo neto)."""
    stock_prev = float(stock_prev or 0)
    costo_prev = float(costo_prev or 0)
    cant_in = float(cant_in or 0)
    precio_in = float(precio_in or 0)
    if stock_prev <= 0 or costo_prev <= 0:
        return precio_in
    nuevo_stock = stock_prev + cant_in
    if nuevo_stock <= 0:
        return precio_in
    return round(((stock_prev * costo_prev) + (cant_in * precio_in)) / nuevo_stock, 6)


def ingresar_almacen(
    cursor,
    *,
    empresa_id: int,
    item_id: Optional[int] = None,
    tipo: str = "producto",
    codigo: str = "",
    nombre: str = "",
    categoria: str = "",
    categoria_codigo: str = "",
    unidad: str = "Kg",
    fecha: str = "",
    cantidad: float = 0.0,
    precio_unitario_neto: float = 0.0,
    proveedor_cuit: str = "",
    proveedor_nombre: str = "",
    nro_comprobante: str = "",
    observaciones: str = "",
) -> Dict[str, Any]:
    """
    Ingreso a almacén a costo NETO (sin impuestos).
    tipo=laboreo: p.ej. pulverización, siembra (unidad Has/Hs).
    """
    if cantidad <= 0:
        raise ValueError("La cantidad de ingreso debe ser mayor a cero.")
    if precio_unitario_neto < 0:
        raise ValueError("El precio unitario neto no puede ser negativo.")

    fecha = (fecha or datetime.now().strftime("%Y-%m-%d"))[:10]
    tipo = (tipo or "producto").strip().lower()
    if tipo not in ("producto", "laboreo"):
        tipo = "producto"

    cat_cod = (categoria_codigo or "").strip().lower()
    if not cat_cod:
        cat_cod = _inferir_categoria_codigo(nombre, tipo, categoria)
    cat_nombre = _nombre_categoria(cursor, cat_cod) or categoria or cat_cod

    redirigido = False
    if item_id:
        cursor.execute(
            "SELECT * FROM almacen_items WHERE id=? AND empresa_id=?;",
            (item_id, empresa_id),
        )
        item = cursor.fetchone()
        if not item:
            raise ValueError("Ítem de almacén no encontrado.")
        item = dict(item)
        if item.get("unificado_en"):
            item = _seguir_unificacion(cursor, item)
            item_id = item["id"]
            redirigido = True
    else:
        if not (nombre or "").strip():
            raise ValueError("Indicá el nombre del producto o laboreo.")
        # Buscar por código, por nombre (o alias de un producto unificado) o crear
        if codigo:
            cursor.execute(
                "SELECT * FROM almacen_items WHERE empresa_id=? AND codigo=?;",
                (empresa_id, codigo.strip()),
            )
            item = cursor.fetchone()
        else:
            item = None
        if not item:
            item = resolver_item_por_nombre(cursor, empresa_id, nombre)
            redirigido = bool(item)
        if item:
            item = dict(item)
            item_id = item["id"]
        else:
            cursor.execute(
                """
                INSERT INTO almacen_items
                (empresa_id, tipo, codigo, nombre, categoria, categoria_codigo, unidad,
                 stock_cantidad, costo_promedio_neto, activo)
                VALUES (?, ?, ?, ?, ?, ?, ?, 0, 0, 1);
                """,
                (
                    empresa_id, tipo, (codigo or "").strip() or None,
                    nombre.strip(), cat_nombre, cat_cod,
                    unidad or ("Has" if tipo == "laboreo" else "Kg"),
                ),
            )
            item_id = cursor.lastrowid
            cursor.execute("SELECT * FROM almacen_items WHERE id=?;", (item_id,))
            item = dict(cursor.fetchone())

    stock_prev = float(item.get("stock_cantidad") or 0)
    costo_prev = float(item.get("costo_promedio_neto") or 0)
    nuevo_prom = _recalcular_promedio(stock_prev, costo_prev, cantidad, precio_unitario_neto)
    nuevo_stock = round(stock_prev + cantidad, 6)
    importe = round(cantidad * precio_unitario_neto, 2)

    cursor.execute(
        """
        UPDATE almacen_items
        SET stock_cantidad=?, costo_promedio_neto=?,
            nombre=COALESCE(NULLIF(?,''), nombre),
            categoria=COALESCE(NULLIF(?,''), categoria),
            categoria_codigo=COALESCE(NULLIF(?,''), categoria_codigo),
            unidad=COALESCE(NULLIF(?,''), unidad),
            tipo=?
        WHERE id=?;
        """,
        (
            nuevo_stock, nuevo_prom, "" if redirigido else (nombre or ""), cat_nombre or "",
            cat_cod or "", "" if redirigido else (unidad or ""),
            (item.get("tipo") or tipo) if redirigido else tipo, item_id,
        ),
    )
    cursor.execute(
        """
        INSERT INTO almacen_movimientos
        (empresa_id, item_id, fecha, tipo_mov, cantidad, precio_unitario_neto, importe_neto,
         stock_resultante, costo_prom_resultante, proveedor_cuit, proveedor_nombre,
         nro_comprobante, observaciones)
        VALUES (?, ?, ?, 'ingreso', ?, ?, ?, ?, ?, ?, ?, ?, ?);
        """,
        (
            empresa_id, item_id, fecha, cantidad, precio_unitario_neto, importe,
            nuevo_stock, nuevo_prom, proveedor_cuit or "", proveedor_nombre or "",
            nro_comprobante or "", observaciones or "",
        ),
    )
    mov_id = cursor.lastrowid
    return {
        "status": "success",
        "item_id": item_id,
        "movimiento_id": mov_id,
        "stock_cantidad": nuevo_stock,
        "costo_promedio_neto": nuevo_prom,
        "importe_neto": importe,
        "message": "Ingreso a almacén registrado a costo neto (sin impuestos).",
    }


def egresar_almacen_nc(
    cursor,
    *,
    empresa_id: int,
    item_id: Optional[int] = None,
    nombre: str = "",
    fecha: str = "",
    cantidad: float = 0.0,
    precio_unitario_neto: float = 0.0,
    proveedor_cuit: str = "",
    nro_comprobante: str = "",
    observaciones: str = "",
) -> Dict[str, Any]:
    """
    Devolución / egreso por Nota de Crédito de compra (baja stock).
    Si no hay stock suficiente, deja stock en 0 y registra el egreso parcial.
    """
    if cantidad <= 0:
        raise ValueError("La cantidad de egreso NC debe ser mayor a cero.")
    fecha = (fecha or datetime.now().strftime("%Y-%m-%d"))[:10]

    item = None
    if item_id:
        cursor.execute(
            "SELECT * FROM almacen_items WHERE id=? AND empresa_id=?;",
            (item_id, empresa_id),
        )
        item = cursor.fetchone()
        if item and item["unificado_en"]:
            item = _seguir_unificacion(cursor, dict(item))
    if not item and (nombre or "").strip():
        item = resolver_item_por_nombre(cursor, empresa_id, nombre)
    if not item:
        raise ValueError("No se encontró el ítem de almacén para egresar por NC.")

    item = dict(item)
    item_id = int(item["id"])
    stock = float(item.get("stock_cantidad") or 0)
    costo_u = float(item.get("costo_promedio_neto") or 0) or float(precio_unitario_neto or 0)
    cant = min(cantidad, stock) if stock > 0 else cantidad
    nuevo_stock = round(max(0.0, stock - cant), 6)
    importe = round(cant * costo_u, 2)

    cursor.execute(
        "UPDATE almacen_items SET stock_cantidad=? WHERE id=?;",
        (nuevo_stock, item_id),
    )
    cursor.execute(
        """
        INSERT INTO almacen_movimientos
        (empresa_id, item_id, fecha, tipo_mov, cantidad, precio_unitario_neto, importe_neto,
         stock_resultante, costo_prom_resultante, proveedor_cuit, proveedor_nombre,
         nro_comprobante, observaciones)
        VALUES (?, ?, ?, 'egreso_nc', ?, ?, ?, ?, ?, ?, '', ?, ?);
        """,
        (
            empresa_id, item_id, fecha, cant, costo_u, importe,
            nuevo_stock, costo_u, proveedor_cuit or "",
            nro_comprobante or "", observaciones or "",
        ),
    )
    return {
        "status": "success",
        "item_id": item_id,
        "movimiento_id": cursor.lastrowid,
        "stock_cantidad": nuevo_stock,
        "cantidad_egresada": cant,
        "message": "Egreso por Nota de Crédito registrado.",
    }


NRO_OT_INICIAL = 3274


def _ultimo_nro_ot(cursor, empresa_id: int) -> int:
    ultimo = NRO_OT_INICIAL - 1
    cursor.execute(
        """
        SELECT MAX(CAST(nro_ot AS INTEGER)) FROM ordenes_trabajo
        WHERE empresa_id=? AND TRIM(COALESCE(nro_ot,'')) != '' AND nro_ot NOT GLOB '*[^0-9]*';
        """,
        (empresa_id,),
    )
    row = cursor.fetchone()
    if row and row[0]:
        ultimo = max(ultimo, int(row[0]))
    cursor.execute("SELECT MAX(nro_orden) FROM margenes_access WHERE empresa_id=?;", (empresa_id,))
    row = cursor.fetchone()
    if row and row[0]:
        ultimo = max(ultimo, int(row[0]))
    return ultimo


def siguiente_nro_ot(cursor, empresa_id: int, anio: Optional[int] = None) -> str:
    return str(_ultimo_nro_ot(cursor, empresa_id) + 1)


def _nro_ot_en_uso(cursor, empresa_id: int, nro: str) -> bool:
    cursor.execute(
        "SELECT 1 FROM ordenes_trabajo WHERE empresa_id=? AND TRIM(nro_ot)=? LIMIT 1;",
        (empresa_id, nro),
    )
    if cursor.fetchone():
        return True
    if nro.isdigit():
        cursor.execute(
            "SELECT 1 FROM margenes_access WHERE empresa_id=? AND nro_orden=? LIMIT 1;",
            (empresa_id, int(nro)),
        )
        return cursor.fetchone() is not None
    return False


# ---- Participación con otras sociedades (por campo y campaña) ----

def asegurar_schema_participaciones(cursor) -> None:
    """Las bases por empresa se clonan una sola vez: se completa acá lo que falte."""
    cambios = False
    cursor.execute(
        "SELECT 1 FROM sqlite_master WHERE type='table' AND name='campo_participaciones';"
    )
    if not cursor.fetchone():
        cursor.execute(
            """
            CREATE TABLE campo_participaciones (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                empresa_id INTEGER DEFAULT 1,
                campo_id INTEGER NOT NULL,
                campania_id INTEGER NOT NULL,
                socio_nombre TEXT DEFAULT '',
                socio_cuit TEXT DEFAULT '',
                porcentaje REAL DEFAULT 0
            );
            """
        )
        cursor.execute(
            "CREATE INDEX IF NOT EXISTS idx_campo_part ON campo_participaciones(empresa_id, campo_id, campania_id);"
        )
        cambios = True
    cursor.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='socios_agro';")
    if not cursor.fetchone():
        cursor.execute(
            """
            CREATE TABLE socios_agro (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                empresa_id INTEGER DEFAULT 1,
                nombre TEXT NOT NULL,
                cuit TEXT DEFAULT '',
                contacto TEXT DEFAULT '',
                telefono TEXT DEFAULT '',
                email TEXT DEFAULT '',
                observaciones TEXT DEFAULT '',
                activo INTEGER DEFAULT 1
            );
            """
        )
        cambios = True
    for tabla, columnas in (
        ("campo_participaciones", (("socio_id", "INTEGER"), ("aporta_labores", "INTEGER DEFAULT 1"),
                                   ("aporta_insumos", "INTEGER DEFAULT 1"))),
        ("ot_destinos", (("participacion_pct", "REAL DEFAULT 100"), ("socios_txt", "TEXT DEFAULT ''"))),
        ("ot_consumos", (("aporte_pct", "REAL DEFAULT 100"), ("participacion_pct", "REAL DEFAULT 100"),
                         ("aporta", "TEXT DEFAULT ''"))),
    ):
        existentes = {r[1] for r in cursor.execute(f"PRAGMA table_info({tabla});").fetchall()}
        if not existentes:
            continue
        for col, ddl in columnas:
            if col not in existentes:
                cursor.execute(f"ALTER TABLE {tabla} ADD COLUMN {col} {ddl};")
                cambios = True
    if cambios:
        cursor.connection.commit()


def participacion_campo(cursor, empresa_id: int, campo_id: Optional[int], campania_id: Optional[int]) -> Dict[str, Any]:
    """% de la empresa y socios del campo en la campaña. Sin carga propia, toma la
    última campaña anterior que tenga carga. Una fila sin socio marca 'sin participación'."""
    vacio = {"pct_empresa": 100.0, "socios": [], "socios_txt": "", "campania_id_origen": None, "campania_codigo_origen": ""}
    if not campo_id or not campania_id:
        return vacio
    cursor.execute("SELECT codigo FROM campanias_agro WHERE id=?;", (campania_id,))
    r = cursor.fetchone()
    codigo = (r["codigo"] if r else "") or ""
    cursor.execute(
        """
        SELECT p.campania_id, ca.codigo
        FROM campo_participaciones p
        JOIN campanias_agro ca ON ca.id = p.campania_id
        WHERE p.empresa_id=? AND p.campo_id=? AND (p.campania_id=? OR ca.codigo <= ?)
        ORDER BY CASE WHEN p.campania_id=? THEN 0 ELSE 1 END, ca.codigo DESC
        LIMIT 1;
        """,
        (empresa_id, campo_id, campania_id, codigo, campania_id),
    )
    origen = cursor.fetchone()
    if not origen:
        return vacio
    cursor.execute(
        """
        SELECT socio_id, socio_nombre, socio_cuit, porcentaje,
               COALESCE(aporta_labores, 1) AS aporta_labores, COALESCE(aporta_insumos, 1) AS aporta_insumos
        FROM campo_participaciones
        WHERE empresa_id=? AND campo_id=? AND campania_id=?
          AND TRIM(COALESCE(socio_nombre,'')) != '' AND COALESCE(porcentaje,0) > 0
        ORDER BY id;
        """,
        (empresa_id, campo_id, origen["campania_id"]),
    )
    socios = [dict(x) for x in cursor.fetchall()]
    pct_socios = sum(float(s["porcentaje"] or 0) for s in socios)
    return {
        "pct_empresa": round(100.0 - pct_socios, 4),
        "socios": socios,
        "socios_txt": "; ".join(f"{s['socio_nombre']} {float(s['porcentaje']):g}%" for s in socios),
        "aporta_labores": any(int(s["aporta_labores"]) for s in socios),
        "aporta_insumos": any(int(s["aporta_insumos"]) for s in socios),
        "campania_id_origen": int(origen["campania_id"]),
        "campania_codigo_origen": origen["codigo"] or "",
    }


def guardar_participacion_campo(
    cursor, empresa_id: int, campo_id: int, campania_id: int, socios: List[Dict[str, Any]]
) -> Dict[str, Any]:
    # Quien guarda sin indicar la modalidad (ej. el modal de Margenes) conserva la anterior del socio.
    previos = {
        (s["socio_nombre"] or "").strip().lower(): s
        for s in participacion_campo(cursor, empresa_id, campo_id, campania_id)["socios"]
    }
    limpios = []
    for s in socios or []:
        socio_id = int(s.get("socio_id") or 0) or None
        nombre = (s.get("socio_nombre") or "").strip()
        cuit = (s.get("socio_cuit") or "").strip()
        if not socio_id and nombre:
            socio_id = previos.get(nombre.lower(), {}).get("socio_id")
        if socio_id:
            cursor.execute("SELECT nombre, cuit FROM socios_agro WHERE id=? AND empresa_id=?;", (socio_id, empresa_id))
            reg = cursor.fetchone()
            if not reg:
                raise ValueError("Socio inexistente.")
            nombre, cuit = reg["nombre"].strip(), (reg["cuit"] or "").strip()
        pct = float(s.get("porcentaje") or 0)
        if not nombre and pct <= 0:
            continue
        if not nombre:
            raise ValueError("Indicá el nombre de cada sociedad participante.")
        if pct <= 0 or pct >= 100:
            raise ValueError(f"El % de {nombre} tiene que estar entre 0 y 100.")
        prev = previos.get(nombre.lower(), {})
        flags = {}
        for k in ("aporta_labores", "aporta_insumos"):
            v = s.get(k)
            flags[k] = int(prev.get(k, 1)) if v is None else (1 if v else 0)
        limpios.append({"socio_id": socio_id or prev.get("socio_id"), "socio_nombre": nombre,
                        "socio_cuit": cuit, "porcentaje": pct, **flags})
    total = sum(s["porcentaje"] for s in limpios)
    if total >= 100:
        raise ValueError("La suma de los socios tiene que dejar un % para la empresa.")
    cursor.execute(
        "DELETE FROM campo_participaciones WHERE empresa_id=? AND campo_id=? AND campania_id=?;",
        (empresa_id, campo_id, campania_id),
    )
    vacio = {"socio_id": None, "socio_nombre": "", "socio_cuit": "", "porcentaje": 0,
             "aporta_labores": 1, "aporta_insumos": 1}
    for s in limpios or [vacio]:
        cursor.execute(
            """
            INSERT INTO campo_participaciones (empresa_id, campo_id, campania_id, socio_id, socio_nombre,
                                               socio_cuit, porcentaje, aporta_labores, aporta_insumos)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?);
            """,
            (empresa_id, campo_id, campania_id, s["socio_id"], s["socio_nombre"], s["socio_cuit"],
             s["porcentaje"], s["aporta_labores"], s["aporta_insumos"]),
        )
    return participacion_campo(cursor, empresa_id, campo_id, campania_id)


def listar_socios_agro(cursor, empresa_id: int, incluir_inactivos: bool = False) -> List[Dict[str, Any]]:
    cursor.execute(
        f"""
        SELECT s.*,
               (SELECT COUNT(DISTINCT p.campo_id) FROM campo_participaciones p
                WHERE p.empresa_id = s.empresa_id AND (p.socio_id = s.id
                      OR (p.socio_id IS NULL AND LOWER(TRIM(p.socio_nombre)) = LOWER(TRIM(s.nombre))))) AS campos
        FROM socios_agro s
        WHERE s.empresa_id = ? {'' if incluir_inactivos else 'AND COALESCE(s.activo, 1) = 1'}
        ORDER BY s.nombre COLLATE NOCASE;
        """,
        (empresa_id,),
    )
    return [dict(r) for r in cursor.fetchall()]


def guardar_socio_agro(cursor, empresa_id: int, data: Dict[str, Any], socio_id: Optional[int] = None) -> Dict[str, Any]:
    nombre = (data.get("nombre") or "").strip()
    if not nombre:
        raise ValueError("El nombre del socio es obligatorio.")
    cursor.execute(
        "SELECT id FROM socios_agro WHERE empresa_id=? AND LOWER(TRIM(nombre))=LOWER(?) AND id != COALESCE(?, -1);",
        (empresa_id, nombre, socio_id),
    )
    if cursor.fetchone():
        raise ValueError("Ya existe un socio con ese nombre.")
    campos = {k: (data.get(k) or "").strip() for k in ("cuit", "contacto", "telefono", "email", "observaciones")}
    campos["nombre"] = nombre
    campos["activo"] = 1 if data.get("activo", True) else 0
    if socio_id:
        cursor.execute("SELECT nombre FROM socios_agro WHERE id=? AND empresa_id=?;", (socio_id, empresa_id))
        prev = cursor.fetchone()
        if not prev:
            raise ValueError("Socio inexistente.")
        cursor.execute(
            f"UPDATE socios_agro SET {', '.join(f'{k}=?' for k in campos)} WHERE id=?;",
            (*campos.values(), socio_id),
        )
        cursor.execute(
            """
            UPDATE campo_participaciones SET socio_id=?, socio_nombre=?, socio_cuit=?
            WHERE empresa_id=? AND (socio_id=? OR (socio_id IS NULL AND LOWER(TRIM(socio_nombre))=LOWER(TRIM(?))));
            """,
            (socio_id, nombre, campos["cuit"], empresa_id, socio_id, prev["nombre"]),
        )
    else:
        cursor.execute(
            f"INSERT INTO socios_agro (empresa_id, {', '.join(campos)}) VALUES (?, {', '.join('?' for _ in campos)});",
            (empresa_id, *campos.values()),
        )
        socio_id = cursor.lastrowid
        cursor.execute(
            """
            UPDATE campo_participaciones SET socio_id=?
            WHERE empresa_id=? AND socio_id IS NULL AND LOWER(TRIM(socio_nombre))=LOWER(?);
            """,
            (socio_id, empresa_id, nombre),
        )
    cursor.execute("SELECT * FROM socios_agro WHERE id=?;", (socio_id,))
    return dict(cursor.fetchone())


def borrar_socio_agro(cursor, empresa_id: int, socio_id: int) -> str:
    """Si ya participa en algún campo se da de baja (inactivo) para no perder la historia."""
    cursor.execute("SELECT nombre FROM socios_agro WHERE id=? AND empresa_id=?;", (socio_id, empresa_id))
    reg = cursor.fetchone()
    if not reg:
        raise ValueError("Socio inexistente.")
    cursor.execute(
        """
        SELECT 1 FROM campo_participaciones
        WHERE empresa_id=? AND (socio_id=? OR LOWER(TRIM(socio_nombre))=LOWER(TRIM(?))) LIMIT 1;
        """,
        (empresa_id, socio_id, reg["nombre"]),
    )
    if cursor.fetchone():
        cursor.execute("UPDATE socios_agro SET activo=0 WHERE id=?;", (socio_id,))
        return "inactivo"
    cursor.execute("DELETE FROM socios_agro WHERE id=?;", (socio_id,))
    return "borrado"


def participaciones_campania(cursor, empresa_id: int, campania_id: int) -> Dict[int, Dict[str, Any]]:
    """Participación vigente en la campaña de cada campo que la tenga."""
    cursor.execute(
        "SELECT DISTINCT campo_id FROM campo_participaciones WHERE empresa_id=?;",
        (empresa_id,),
    )
    res = {}
    for r in cursor.fetchall():
        p = participacion_campo(cursor, empresa_id, int(r["campo_id"]), campania_id)
        if p["pct_empresa"] < 100:
            res[int(r["campo_id"])] = p
    return res


def emitir_ot_consumiendo_almacen(
    cursor,
    *,
    empresa_id: int,
    fecha: str,
    tipo_labor: str,
    contratista_nombre: str = "",
    contratista_cuit: str = "",
    campania_id: Optional[int] = None,
    campo_id: Optional[int] = None,
    lote_id: Optional[int] = None,
    superficie_has: float = 0.0,
    cultivo: str = "",
    observaciones: str = "",
    consumos: Optional[List[Dict[str, Any]]] = None,
    nro_ot: Optional[str] = None,
    fecha_aplicacion: str = "",
    labor_cultural: str = "",
    destinos: Optional[List[Dict[str, Any]]] = None,
) -> Dict[str, Any]:
    """
    Emite la OT: registra las líneas y la salida física del almacén, sin costo.
    El costo se asigna después con costear_linea_ot. No se bloquea por falta de stock.
    cada consumo: {item_id, cantidad?, dosis_por_ha?, aporte_pct?}
    destinos: [{campo_id, lote_id, superficie_has, cultivo}]; la cantidad de cada
    consumo se reparte entre los destinos en proporción a sus hectáreas.
    En campos con participación, superficie_has son las hectáreas físicas del lote.
    """
    asegurar_schema_participaciones(cursor)
    consumos = consumos or []
    fecha = (fecha or datetime.now().strftime("%Y-%m-%d"))[:10]
    fecha_aplic = (fecha_aplicacion or fecha)[:10]
    nro = str(nro_ot or "").strip() or siguiente_nro_ot(cursor, empresa_id)
    if _nro_ot_en_uso(cursor, empresa_id, nro):
        raise ValueError(f"La OT Nº {nro} ya existe.")
    if not any(int(c.get("item_id") or 0) for c in consumos):
        raise ValueError("Agregá al menos un producto o labor.")
    dests = _normalizar_destinos_ot(cursor, destinos, campo_id, lote_id, superficie_has, cultivo)
    total_has = round(sum(d["superficie_has"] for d in dests), 4)

    cursor.execute(
        """
        INSERT INTO ordenes_trabajo
        (empresa_id, nro_ot, fecha, tipo_labor, contratista_cuit, contratista_nombre,
         campania_id, campo_id, lote_id, superficie_has, cultivo, estado, observaciones,
         fecha_aplicacion, labor_cultural, confirmada)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'Pendiente de costeo', ?, ?, ?, 0);
        """,
        (
            empresa_id, nro, fecha, tipo_labor or "",
            contratista_cuit or "", contratista_nombre or "",
            campania_id, dests[0]["campo_id"], dests[0]["lote_id"], total_has,
            dests[0]["cultivo"] or cultivo or "",
            observaciones or "", fecha_aplic, labor_cultural or "",
        ),
    )
    ot_id = cursor.lastrowid
    lineas, sin_stock = _crear_lineas_ot(
        cursor, empresa_id, ot_id, nro,
        fecha=fecha, fecha_aplic=fecha_aplic, tipo_labor=tipo_labor,
        labor_cultural=labor_cultural, contratista_nombre=contratista_nombre,
        campania_id=campania_id, destinos=dests, consumos=consumos,
    )
    msg = f"OT {nro} emitida con {lineas} línea(s)"
    if len(dests) > 1:
        msg += f" en {len(dests)} lotes"
    msg += ". Queda pendiente de costeo."
    if sin_stock:
        msg += " Atención: algunos productos quedaron con stock negativo en el almacén."
    return {
        "status": "success",
        "ot_id": ot_id,
        "nro_ot": nro,
        "lineas": lineas,
        "sin_stock": sin_stock,
        "message": msg,
    }


def _normalizar_destinos_ot(
    cursor,
    destinos: Optional[List[Dict[str, Any]]],
    campo_id: Optional[int],
    lote_id: Optional[int],
    superficie_has: float,
    cultivo: str,
) -> List[Dict[str, Any]]:
    lista: List[Dict[str, Any]] = []
    for d in destinos or []:
        cid = int(d.get("campo_id") or 0) or None
        lid = int(d.get("lote_id") or 0) or None
        has = float(d.get("superficie_has") or 0)
        if not cid and not lid and has <= 0:
            continue
        if has < 0:
            raise ValueError("Las hectáreas de un destino no pueden ser negativas.")
        lista.append({
            "campo_id": cid, "lote_id": lid, "superficie_has": has,
            "cultivo": (d.get("cultivo") or cultivo or "").strip(),
        })
    if not lista:
        lista.append({
            "campo_id": int(campo_id or 0) or None, "lote_id": int(lote_id or 0) or None,
            "superficie_has": float(superficie_has or 0), "cultivo": (cultivo or "").strip(),
        })
    vistos = set()
    for d in lista:
        if d["lote_id"]:
            if d["lote_id"] in vistos:
                raise ValueError("Hay un lote repetido en los destinos de la OT.")
            vistos.add(d["lote_id"])
        d["campo_nombre"] = ""
        d["lote_nombre"] = ""
        if d["campo_id"]:
            cursor.execute("SELECT nombre FROM campos_agro WHERE id=?;", (d["campo_id"],))
            r = cursor.fetchone()
            d["campo_nombre"] = (r["nombre"] if r else "") or ""
        if d["lote_id"]:
            cursor.execute("SELECT nombre, campo_id FROM lotes_agro WHERE id=?;", (d["lote_id"],))
            r = cursor.fetchone()
            if r:
                d["lote_nombre"] = r["nombre"] or ""
                if not d["campo_id"] and r["campo_id"]:
                    d["campo_id"] = int(r["campo_id"])
                    cursor.execute("SELECT nombre FROM campos_agro WHERE id=?;", (d["campo_id"],))
                    rc = cursor.fetchone()
                    d["campo_nombre"] = (rc["nombre"] if rc else "") or ""
    return lista


def _crear_lineas_ot(
    cursor,
    empresa_id: int,
    ot_id: int,
    nro: str,
    *,
    fecha: str,
    fecha_aplic: str,
    tipo_labor: str,
    labor_cultural: str,
    contratista_nombre: str,
    campania_id: Optional[int],
    destinos: List[Dict[str, Any]],
    consumos: List[Dict[str, Any]],
):
    """Graba destinos, salidas de almacén, ot_consumos y líneas de costos de una OT.
    En campos con participación: el almacén descuenta lo que aporta la empresa
    (aporte_pct del consumo) y la línea de costos lleva lo que le corresponde
    según su % en el campo; la cantidad física queda en ot_consumos."""
    campania_codigo = ""
    if campania_id:
        cursor.execute("SELECT codigo FROM campanias_agro WHERE id=?;", (campania_id,))
        r = cursor.fetchone()
        campania_codigo = (r["codigo"] if r else "") or ""
    for i, d in enumerate(destinos):
        part = participacion_campo(cursor, empresa_id, d["campo_id"], campania_id)
        d["participacion_pct"] = part["pct_empresa"]
        cursor.execute(
            """
            INSERT INTO ot_destinos (ot_id, orden, campo_id, lote_id, superficie_has, cultivo,
                                     participacion_pct, socios_txt)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?);
            """,
            (ot_id, i, d["campo_id"], d["lote_id"], d["superficie_has"], d["cultivo"],
             part["pct_empresa"], part["socios_txt"]),
        )
        d["id"] = cursor.lastrowid
    total_has = sum(d["superficie_has"] for d in destinos)
    sin_stock = []
    lineas = 0

    for c in consumos:
        item_id = int(c.get("item_id") or 0)
        if not item_id:
            continue
        cursor.execute(
            "SELECT * FROM almacen_items WHERE id=? AND empresa_id=?;",
            (item_id, empresa_id),
        )
        item = cursor.fetchone()
        if not item:
            raise ValueError(f"Ítem {item_id} no encontrado en almacén.")
        item = dict(item)
        # solo categorías habilitadas para OT
        cat_cod = (item.get("categoria_codigo") or "").strip()
        if cat_cod:
            cursor.execute(
                "SELECT aplica_ot, nombre FROM almacen_categorias WHERE codigo=? LIMIT 1;",
                (cat_cod,),
            )
            cat = cursor.fetchone()
            if cat and not int(cat["aplica_ot"] or 0):
                raise ValueError(
                    f"«{item.get('nombre')}» es categoría «{cat['nombre']}» y no aplica a OT."
                )
        dosis = float(c.get("dosis_por_ha") or 0)
        cant = float(c.get("cantidad") or 0)
        if cant <= 0 and dosis > 0 and total_has > 0:
            cant = round(dosis * total_has, 6)
        if cant <= 0:
            continue
        tipo_item = item.get("tipo") or "producto"
        es_lab = tipo_item == "laboreo"
        aporte_pedido = c.get("aporte_pct")
        aporte_pedido = 100.0 if aporte_pedido is None or aporte_pedido == "" else float(aporte_pedido)
        if aporte_pedido < 0 or aporte_pedido > 100:
            raise ValueError(f"El % de aporte de «{item.get('nombre')}» tiene que estar entre 0 y 100.")
        aporta_txt = (c.get("aporta") or "").strip()[:120]

        if total_has > 0:
            partes = [round(cant * d["superficie_has"] / total_has, 6) for d in destinos]
            partes[-1] = round(cant - sum(partes[:-1]), 6)
        else:
            partes = [cant] + [0.0] * (len(destinos) - 1)
        stock_previo = float(item.get("stock_cantidad") or 0)
        stock = stock_previo

        for d, cant_d in zip(destinos, partes):
            if cant_d <= 0:
                continue
            part_pct = float(d.get("participacion_pct") or 100)
            aporte = aporte_pedido if part_pct < 100 else 100.0
            cant_stock = round(cant_d * aporte / 100, 6)
            cant_emp = round(cant_d * part_pct / 100, 6)
            mov_id = None
            if cant_stock > 0:
                stock = round(stock - cant_stock, 6)
                cursor.execute(
                    "UPDATE almacen_items SET stock_cantidad=? WHERE id=?;",
                    (stock, item_id),
                )
                obs = f"Consumo OT {nro}"
                if part_pct < 100:
                    obs += f" (aporte {aporte:g}% de {cant_d:g})"
                cursor.execute(
                    """
                    INSERT INTO almacen_movimientos
                    (empresa_id, item_id, fecha, tipo_mov, cantidad, precio_unitario_neto, importe_neto,
                     stock_resultante, costo_prom_resultante, campania_id, lote_id, ot_id, observaciones)
                    VALUES (?, ?, ?, 'egreso_ot', ?, 0, 0, ?, ?, ?, ?, ?, ?);
                    """,
                    (
                        empresa_id, item_id, fecha_aplic, cant_stock,
                        stock, float(item.get("costo_promedio_neto") or 0),
                        campania_id, d["lote_id"], ot_id, obs,
                    ),
                )
                mov_id = cursor.lastrowid
            cursor.execute(
                """
                INSERT INTO ot_consumos
                (ot_id, item_id, tipo_item, dosis_por_ha, cantidad, unidad,
                 costo_unitario_neto, costo_total_neto, movimiento_id, destino_id,
                 aporte_pct, participacion_pct, aporta)
                VALUES (?, ?, ?, ?, ?, ?, 0, 0, ?, ?, ?, ?, ?);
                """,
                (ot_id, item_id, tipo_item, dosis, cant_d, item.get("unidad") or "", mov_id, d["id"],
                 aporte, part_pct, aporta_txt if part_pct < 100 else ""),
            )
            consumo_id = cursor.lastrowid
            cursor.execute(
                """
                INSERT INTO margenes_access (
                    empresa_id, campania_codigo, cultivo, nro_orden, campo, lote, cantidad_has,
                    fecha_orden, laboreo, labor_cultural, contratista, producto, tipo, dosis_ha,
                    cantidad_total, fecha_aplicacion, unidad, almacen_item_id, ot_id, ot_consumo_id
                ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?);
                """,
                (
                    empresa_id, campania_codigo, d["cultivo"] or "",
                    int(nro) if nro.isdigit() else None,
                    d["campo_nombre"], d["lote_nombre"],
                    round(float(d["superficie_has"] or 0) * part_pct / 100, 4), fecha,
                    (item.get("nombre") or "") if es_lab else "",
                    (labor_cultural or tipo_labor or "") if es_lab else "",
                    (contratista_nombre or "") if es_lab else "",
                    "" if es_lab else (item.get("nombre") or ""),
                    item.get("categoria") or "", dosis, cant_emp, fecha_aplic,
                    item.get("unidad") or "", item_id, ot_id, consumo_id,
                ),
            )
            lineas += 1

        if not es_lab and stock < -1e-6:
            sin_stock.append({
                "nombre": item.get("nombre"),
                "stock_previo": round(stock_previo, 4),
                "pedido": round(stock_previo - stock, 4),
                "unidad": item.get("unidad") or "",
            })

    if not lineas:
        raise ValueError("Ninguna línea tiene cantidad. Indicá dosis y hectáreas o la cantidad.")
    return lineas, sin_stock


def _revertir_lineas_ot(cursor, ot_id: int) -> Dict[Any, Dict[str, Any]]:
    """Devuelve el stock y borra líneas, consumos y destinos de la OT.
    Retorna el costeo por (ítem, campo, lote) (None si esa línea no estaba costeada)
    y por ítem, para lotes que no existían antes de la edición."""
    costeos: Dict[Any, Optional[Dict[str, Any]]] = {}
    cursor.execute("SELECT * FROM margenes_access WHERE ot_id=? ORDER BY id;", (ot_id,))
    for r in cursor.fetchall():
        m = dict(r)
        if not (float(m.get("costo_ars") or 0) or float(m.get("costo_usd") or 0) or m.get("fecha_costeo")):
            costeos.setdefault((m.get("almacen_item_id"), m.get("campo") or "", m.get("lote") or ""), None)
            continue
        c = {
            "precio": float(m.get("precio") or 0),
            "moneda": m.get("moneda_costo") or "USD",
            "tc": float(m.get("tc") or 0),
            "origen": m.get("origen_costo") or "manual",
            "mov_ingreso_id": m.get("mov_ingreso_id"),
            "proveedor": m.get("proveedor") or "",
            "nro_comprobante": m.get("nro_remito") or "",
            "fecha_compra": m.get("fecha_compra") or "",
        }
        item = m.get("almacen_item_id")
        costeos[(item, m.get("campo") or "", m.get("lote") or "")] = c
        costeos.setdefault(item, c)
    cursor.execute(
        "SELECT id, item_id, cantidad FROM almacen_movimientos WHERE ot_id=? AND tipo_mov='egreso_ot';",
        (ot_id,),
    )
    for m in cursor.fetchall():
        cursor.execute(
            "UPDATE almacen_items SET stock_cantidad=ROUND(COALESCE(stock_cantidad,0)+?, 6) WHERE id=?;",
            (abs(float(m["cantidad"] or 0)), m["item_id"]),
        )
        cursor.execute("DELETE FROM almacen_movimientos WHERE id=?;", (m["id"],))
    cursor.execute("DELETE FROM margenes_access WHERE ot_id=?;", (ot_id,))
    cursor.execute("DELETE FROM ot_consumos WHERE ot_id=?;", (ot_id,))
    cursor.execute("DELETE FROM ot_destinos WHERE ot_id=?;", (ot_id,))
    return costeos


def editar_ot(
    cursor,
    empresa_id: int,
    ot_id: int,
    *,
    fecha: str,
    tipo_labor: str,
    contratista_nombre: str = "",
    contratista_cuit: str = "",
    campania_id: Optional[int] = None,
    campo_id: Optional[int] = None,
    lote_id: Optional[int] = None,
    superficie_has: float = 0.0,
    cultivo: str = "",
    observaciones: str = "",
    consumos: Optional[List[Dict[str, Any]]] = None,
    nro_ot: Optional[str] = None,
    fecha_aplicacion: str = "",
    labor_cultural: str = "",
    destinos: Optional[List[Dict[str, Any]]] = None,
) -> Dict[str, Any]:
    """Corrige una OT emitida y no confirmada: rehace stock y líneas, conservando el costeo."""
    asegurar_schema_participaciones(cursor)
    cursor.execute("SELECT * FROM ordenes_trabajo WHERE id=? AND empresa_id=?;", (ot_id, empresa_id))
    ot = cursor.fetchone()
    if not ot:
        raise ValueError("OT no encontrada.")
    ot = dict(ot)
    if (ot.get("estado") or "") == "Anulada":
        raise ValueError("La OT está anulada; no se puede editar.")
    if int(ot.get("confirmada") or 0):
        raise ValueError("La OT está confirmada. Reabrila para poder editarla.")
    consumos = consumos or []
    if not any(int(c.get("item_id") or 0) for c in consumos):
        raise ValueError("Agregá al menos un producto o labor.")
    nro_actual = str(ot.get("nro_ot") or "").strip()
    nro = str(nro_ot or "").strip() or nro_actual
    if nro != nro_actual and _nro_ot_en_uso(cursor, empresa_id, nro):
        raise ValueError(f"La OT Nº {nro} ya existe.")
    fecha = (fecha or ot.get("fecha") or datetime.now().strftime("%Y-%m-%d"))[:10]
    fecha_aplic = (fecha_aplicacion or fecha)[:10]
    dests = _normalizar_destinos_ot(cursor, destinos, campo_id, lote_id, superficie_has, cultivo)
    total_has = round(sum(d["superficie_has"] for d in dests), 4)

    costeos = _revertir_lineas_ot(cursor, ot_id)
    cursor.execute(
        """
        UPDATE ordenes_trabajo
        SET nro_ot=?, fecha=?, tipo_labor=?, contratista_cuit=?, contratista_nombre=?,
            campania_id=?, campo_id=?, lote_id=?, superficie_has=?, cultivo=?,
            observaciones=?, fecha_aplicacion=?, labor_cultural=?,
            estado='Pendiente de costeo', costo_insumos_neto=0, costo_laboreos_neto=0, costo_total_neto=0
        WHERE id=?;
        """,
        (
            nro, fecha, tipo_labor or "", contratista_cuit or "", contratista_nombre or "",
            campania_id, dests[0]["campo_id"], dests[0]["lote_id"], total_has,
            dests[0]["cultivo"] or cultivo or "",
            observaciones or "", fecha_aplic, labor_cultural or "", ot_id,
        ),
    )
    lineas, sin_stock = _crear_lineas_ot(
        cursor, empresa_id, ot_id, nro,
        fecha=fecha, fecha_aplic=fecha_aplic, tipo_labor=tipo_labor,
        labor_cultural=labor_cultural, contratista_nombre=contratista_nombre,
        campania_id=campania_id, destinos=dests, consumos=consumos,
    )
    recosteadas = 0
    if costeos:
        cursor.execute("SELECT * FROM margenes_access WHERE ot_id=? ORDER BY id;", (ot_id,))
        for r in cursor.fetchall():
            m = dict(r)
            item = m.get("almacen_item_id")
            clave = (item, m.get("campo") or "", m.get("lote") or "")
            c = costeos[clave] if clave in costeos else costeos.get(item)
            if not c:
                continue
            try:
                costear_linea_ot(cursor, empresa_id, int(m["id"]), **c)
                recosteadas += 1
            except ValueError:
                pass
    _recalcular_totales_ot(cursor, ot_id)
    msg = f"OT {nro} actualizada con {lineas} línea(s)"
    if len(dests) > 1:
        msg += f" en {len(dests)} lotes"
    msg += "."
    if recosteadas:
        msg += f" Se conservó el costeo de {recosteadas} línea(s)."
    if sin_stock:
        msg += " Atención: algunos productos quedaron con stock negativo en el almacén."
    return {
        "status": "success",
        "ot_id": ot_id,
        "nro_ot": nro,
        "lineas": lineas,
        "sin_stock": sin_stock,
        "message": msg,
    }


def confirmar_ot(cursor, empresa_id: int, ot_id: int, confirmar: bool = True) -> Dict[str, Any]:
    cursor.execute("SELECT * FROM ordenes_trabajo WHERE id=? AND empresa_id=?;", (ot_id, empresa_id))
    ot = cursor.fetchone()
    if not ot:
        raise ValueError("OT no encontrada.")
    if (ot["estado"] or "") == "Anulada":
        raise ValueError("La OT está anulada.")
    cursor.execute(
        "UPDATE ordenes_trabajo SET confirmada=?, fecha_confirmacion=? WHERE id=?;",
        (
            1 if confirmar else 0,
            datetime.now().strftime("%Y-%m-%d %H:%M") if confirmar else None,
            ot_id,
        ),
    )
    accion = "confirmada" if confirmar else "reabierta para edición"
    return {"status": "success", "message": f"OT {ot['nro_ot']} {accion}."}


def detalle_ot(cursor, empresa_id: int, ot_id: int) -> Dict[str, Any]:
    """Cabecera, destinos y consumos (agrupados por ítem) para cargar la OT en el formulario."""
    asegurar_schema_participaciones(cursor)
    cursor.execute("SELECT * FROM ordenes_trabajo WHERE id=? AND empresa_id=?;", (ot_id, empresa_id))
    ot = cursor.fetchone()
    if not ot:
        raise ValueError("OT no encontrada.")
    data = dict(ot)
    cursor.execute(
        """
        SELECT d.campo_id, d.lote_id, d.superficie_has, d.cultivo,
               COALESCE(d.participacion_pct, 100) AS participacion_pct, COALESCE(d.socios_txt, '') AS socios_txt,
               c.nombre AS campo_nombre, l.nombre AS lote_nombre
        FROM ot_destinos d
        LEFT JOIN campos_agro c ON c.id = d.campo_id
        LEFT JOIN lotes_agro l ON l.id = d.lote_id
        WHERE d.ot_id=? ORDER BY d.orden, d.id;
        """,
        (ot_id,),
    )
    destinos = [dict(r) for r in cursor.fetchall()]
    if not destinos:
        destinos = [{
            "campo_id": data.get("campo_id"), "lote_id": data.get("lote_id"),
            "superficie_has": data.get("superficie_has") or 0, "cultivo": data.get("cultivo") or "",
        }]
    cursor.execute(
        """
        SELECT oc.item_id, MAX(oc.tipo_item) AS tipo_item, MAX(oc.unidad) AS unidad,
               MAX(oc.dosis_por_ha) AS dosis_por_ha, SUM(oc.cantidad) AS cantidad,
               COALESCE(MAX(CASE WHEN COALESCE(oc.participacion_pct, 100) < 100 THEN oc.aporte_pct END), 100) AS aporte_pct,
               COALESCE(MAX(CASE WHEN COALESCE(oc.participacion_pct, 100) < 100 THEN oc.aporta END), '') AS aporta,
               MAX(i.nombre) AS nombre, MIN(oc.id) AS orden
        FROM ot_consumos oc
        LEFT JOIN almacen_items i ON i.id = oc.item_id
        WHERE oc.ot_id=?
        GROUP BY oc.item_id
        ORDER BY orden;
        """,
        (ot_id,),
    )
    data["destinos"] = destinos
    data["consumos"] = [dict(r) for r in cursor.fetchall()]
    return data


def _recalcular_totales_ot(cursor, ot_id: int) -> None:
    cursor.execute(
        """
        SELECT COUNT(*) AS n,
               SUM(CASE WHEN COALESCE(costo_ars,0) != 0 OR COALESCE(costo_usd,0) != 0 THEN 1 ELSE 0 END) AS costeadas,
               COALESCE(SUM(CASE WHEN TRIM(COALESCE(producto,'')) != '' THEN costo_usd END), 0) AS insumos_usd,
               COALESCE(SUM(CASE WHEN TRIM(COALESCE(producto,'')) = '' THEN costo_ars END), 0) AS labores_ars,
               COALESCE(SUM(costo_usd), 0) AS total_usd
        FROM margenes_access WHERE ot_id=?;
        """,
        (ot_id,),
    )
    r = dict(cursor.fetchone())
    n = int(r["n"] or 0)
    costeadas = int(r["costeadas"] or 0)
    if n and costeadas >= n:
        estado = "Costeada"
    elif costeadas:
        estado = "Costeo parcial"
    else:
        estado = "Pendiente de costeo"
    cursor.execute(
        """
        UPDATE ordenes_trabajo
        SET costo_insumos_neto=?, costo_laboreos_neto=?, costo_total_neto=?, estado=?
        WHERE id=? AND COALESCE(estado,'') != 'Anulada';
        """,
        (
            round(float(r["insumos_usd"] or 0), 2), round(float(r["labores_ars"] or 0), 2),
            round(float(r["total_usd"] or 0), 2), estado, ot_id,
        ),
    )


def _es_linea_laboreo(linea: Dict[str, Any]) -> bool:
    return not (linea.get("producto") or "").strip()


def _item_de_linea(cursor, empresa_id: int, linea: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    if linea.get("almacen_item_id"):
        cursor.execute(
            "SELECT * FROM almacen_items WHERE id=? AND empresa_id=?;",
            (linea["almacen_item_id"], empresa_id),
        )
        r = cursor.fetchone()
        if r:
            return _seguir_unificacion(cursor, dict(r))
    nombre = (linea.get("producto") or linea.get("laboreo") or "").strip()
    if not nombre:
        return None
    return resolver_item_por_nombre(cursor, empresa_id, nombre)


def _cuit_contratista(cursor, empresa_id: int, linea: Dict[str, Any]) -> str:
    if linea.get("ot_id"):
        cursor.execute("SELECT contratista_cuit FROM ordenes_trabajo WHERE id=?;", (linea["ot_id"],))
        r = cursor.fetchone()
        if r and (r[0] or "").strip():
            return r[0].strip()
    nombre = (linea.get("contratista") or "").strip()
    if not nombre:
        return ""
    cursor.execute(
        """
        SELECT cuit FROM entidades
        WHERE COALESCE(empresa_id,1)=? AND COALESCE(es_proveedor,0)=1
          AND (UPPER(TRIM(razon_social))=UPPER(TRIM(?)) OR UPPER(TRIM(nombre_fantasia))=UPPER(TRIM(?)))
        LIMIT 1;
        """,
        (empresa_id, nombre, nombre),
    )
    r = cursor.fetchone()
    return (r[0] or "").strip() if r else ""


def _tc_sugerido(cursor, empresa_id: int, fecha: str) -> float:
    cursor.execute(
        """
        SELECT tc FROM margenes_access
        WHERE empresa_id=? AND COALESCE(tc,0) > 0 AND COALESCE(fecha_aplicacion,'') <= ?
        ORDER BY fecha_aplicacion DESC, id DESC LIMIT 1;
        """,
        (empresa_id, fecha or "9999-12-31"),
    )
    r = cursor.fetchone()
    return float(r[0]) if r else 0.0


def listar_lineas_ot(
    cursor,
    empresa_id: int,
    *,
    solo_pendientes: bool = True,
    nro_orden: Optional[int] = None,
    campania_codigo: str = "",
    ultimas_ot: int = 40,
) -> List[Dict[str, Any]]:
    """Líneas de OT (Access + nuevas) agrupables por Nº de orden, para costear."""
    asegurar_schema_participaciones(cursor)
    where = " WHERE m.empresa_id=? AND m.nro_orden IS NOT NULL "
    params: List[Any] = [empresa_id]
    if nro_orden:
        where += " AND m.nro_orden=? "
        params.append(int(nro_orden))
    if campania_codigo:
        where += " AND m.campania_codigo=? "
        params.append(campania_codigo)
    pend = " AND COALESCE(m.costo_ars,0)=0 AND COALESCE(m.costo_usd,0)=0 "
    if solo_pendientes:
        where += pend
    anuladas = " AND (m.ot_id IS NULL OR COALESCE(o.estado,'') != 'Anulada') "
    cursor.execute(
        f"""
        SELECT DISTINCT m.nro_orden FROM margenes_access m
        LEFT JOIN ordenes_trabajo o ON o.id = m.ot_id
        {where} {anuladas}
        ORDER BY m.nro_orden DESC LIMIT ?;
        """,
        params + [max(1, min(int(ultimas_ot or 40), 500))],
    )
    nros = [r[0] for r in cursor.fetchall()]
    if not nros:
        return []
    marcas = ",".join("?" for _ in nros)
    cursor.execute(
        f"""
        SELECT m.id, m.nro_orden, m.campania_codigo, m.cultivo, m.campo, m.lote, m.cantidad_has,
               m.fecha_orden, m.fecha_aplicacion, m.laboreo, m.labor_cultural, m.contratista,
               m.producto, m.dosis_ha, m.cantidad_total, m.unidad, m.precio, m.tc,
               m.costo_ars, m.costo_usd, m.moneda_costo, m.origen_costo, m.proveedor,
               m.nro_remito, m.fecha_compra, m.almacen_item_id, m.ot_id, m.id_access,
               o.estado AS estado_ot,
               oc.cantidad AS cantidad_fisica,
               COALESCE(oc.participacion_pct, 100) AS participacion_pct,
               COALESCE(oc.aporte_pct, 100) AS aporte_pct
        FROM margenes_access m
        LEFT JOIN ordenes_trabajo o ON o.id = m.ot_id
        LEFT JOIN ot_consumos oc ON oc.id = m.ot_consumo_id
        {where} {anuladas} AND m.nro_orden IN ({marcas})
        ORDER BY m.nro_orden DESC, CASE WHEN TRIM(COALESCE(m.producto,''))='' THEN 0 ELSE 1 END, m.id;
        """,
        params + nros,
    )
    return [dict(r) for r in cursor.fetchall()]


def _sugerido_linea(cursor, empresa_id: int, linea: Dict[str, Any], item: Dict[str, Any],
                    movimiento_id: Optional[int]) -> Optional[Dict[str, Any]]:
    """Costo por el criterio de la empresa para la salida de almacén de esa línea."""
    cant = float(linea.get("cantidad_total") or 0)
    if movimiento_id:
        cursor.execute("SELECT ABS(cantidad) FROM almacen_movimientos WHERE id=?;", (movimiento_id,))
        r = cursor.fetchone()
        if r and float(r[0] or 0) > 0:
            cant = float(r[0])
    if cant <= 0:
        return None
    fecha = (linea.get("fecha_aplicacion") or linea.get("fecha_orden") or "")[:10]
    try:
        return valuar_salida_insumo(cursor, empresa_id, int(item["id"]), cant,
                                    movimiento_id=movimiento_id, fecha=fecha)
    except ValueError:
        return None


def opciones_costeo_linea(cursor, empresa_id: int, linea_id: int) -> Dict[str, Any]:
    """Compras del almacén y facturas del contratista para elegir el costo de una línea."""
    asegurar_schema_participaciones(cursor)
    cursor.execute("SELECT * FROM margenes_access WHERE id=? AND empresa_id=?;", (linea_id, empresa_id))
    r = cursor.fetchone()
    if not r:
        raise ValueError("Línea de OT no encontrada.")
    linea = dict(r)
    linea["cantidad_fisica"] = None
    linea["participacion_pct"] = 100.0
    linea["aporte_pct"] = 100.0
    movimiento_id = None
    if linea.get("ot_consumo_id"):
        cursor.execute(
            "SELECT cantidad, participacion_pct, aporte_pct, movimiento_id FROM ot_consumos WHERE id=?;",
            (linea["ot_consumo_id"],),
        )
        oc = cursor.fetchone()
        if oc:
            linea["cantidad_fisica"] = oc["cantidad"]
            linea["participacion_pct"] = float(oc["participacion_pct"] if oc["participacion_pct"] is not None else 100)
            linea["aporte_pct"] = float(oc["aporte_pct"] if oc["aporte_pct"] is not None else 100)
            movimiento_id = oc["movimiento_id"]
    es_lab = _es_linea_laboreo(linea)
    item = _item_de_linea(cursor, empresa_id, linea)
    fecha = (linea.get("fecha_aplicacion") or linea.get("fecha_orden") or "")[:10]
    cuit = _cuit_contratista(cursor, empresa_id, linea) if es_lab else ""

    ingresos: List[Dict[str, Any]] = []
    vistos = set()

    def _agregar_ingresos(sql: str, args: tuple) -> None:
        cursor.execute(sql, args)
        for m in cursor.fetchall():
            d = dict(m)
            if d["id"] in vistos:
                continue
            vistos.add(d["id"])
            ars = float(d.get("precio_unitario_neto") or 0)
            usd = float(d.get("precio_unitario_usd") or 0)
            d["tc"] = round(ars / usd, 4) if usd > 0 and ars > 0 else 0
            d["posterior"] = bool(fecha and (d.get("fecha") or "") > fecha)
            ingresos.append(d)

    base_sql = """
        SELECT m.id, m.fecha, m.cantidad, m.precio_unitario_neto,
               COALESCE(m.precio_unitario_usd,0) AS precio_unitario_usd,
               m.proveedor_cuit, m.proveedor_nombre, m.nro_comprobante,
               i.id AS item_id, i.nombre AS item_nombre, i.unidad
        FROM almacen_movimientos m JOIN almacen_items i ON i.id = m.item_id
        WHERE m.empresa_id=? AND m.tipo_mov='ingreso' AND COALESCE(m.precio_unitario_neto,0) > 0
    """
    if item:
        _agregar_ingresos(base_sql + " AND m.item_id=? ORDER BY m.fecha DESC, m.id DESC LIMIT 40;",
                          (empresa_id, item["id"]))
    if es_lab and cuit:
        _agregar_ingresos(
            base_sql + """ AND i.tipo='laboreo'
              AND REPLACE(COALESCE(m.proveedor_cuit,''),'-','')=REPLACE(?,'-','')
            ORDER BY m.fecha DESC, m.id DESC LIMIT 40;""",
            (empresa_id, cuit),
        )
    ingresos.sort(key=lambda d: (d.get("fecha") or "", d["id"]), reverse=True)

    facturas: List[Dict[str, Any]] = []
    if es_lab and cuit:
        cursor.execute(
            """
            SELECT id, fecha, tipo_comprobante, numero_comprobante, neto, iva, debe, total, observaciones
            FROM cuentas_corrientes
            WHERE COALESCE(empresa_id,1)=? AND REPLACE(entidad_id,'-','')=REPLACE(?,'-','')
              AND COALESCE(debe,0) > 0
              AND COALESCE(tipo_comprobante,'') NOT LIKE 'Ajuste%'
              AND COALESCE(tipo_comprobante,'') NOT LIKE 'Pago%'
            ORDER BY fecha DESC, id DESC LIMIT 20;
            """,
            (empresa_id, cuit),
        )
        for f in cursor.fetchall():
            d = dict(f)
            neto = float(d.get("neto") or 0)
            d["neto_estimado"] = neto <= 0
            d["neto_usar"] = round(neto if neto > 0 else float(d.get("debe") or 0) / 1.21, 2)
            d["comprobante"] = (d.get("numero_comprobante") or d.get("tipo_comprobante") or "").strip()
            facturas.append(d)

    sugerido = _sugerido_linea(cursor, empresa_id, linea, item, movimiento_id) if item and not es_lab else None

    return {
        "linea": linea,
        "es_laboreo": es_lab,
        "item": {k: item.get(k) for k in ("id", "nombre", "unidad", "costo_promedio_neto", "costo_promedio_usd", "stock_cantidad")} if item else None,
        "contratista_cuit": cuit,
        "ingresos": ingresos,
        "facturas": facturas,
        "sugerido_criterio": sugerido,
        "tc_sugerido": _tc_sugerido(cursor, empresa_id, fecha),
    }


def costear_linea_ot(
    cursor,
    empresa_id: int,
    linea_id: int,
    *,
    precio: float,
    moneda: str = "USD",
    tc: float = 0.0,
    origen: str = "manual",
    mov_ingreso_id: Optional[int] = None,
    proveedor: str = "",
    nro_comprobante: str = "",
    fecha_compra: str = "",
) -> Dict[str, Any]:
    """Asigna el costo a una línea de OT. Costo = cantidad × precio unitario neto."""
    cursor.execute("SELECT * FROM margenes_access WHERE id=? AND empresa_id=?;", (linea_id, empresa_id))
    r = cursor.fetchone()
    if not r:
        raise ValueError("Línea de OT no encontrada.")
    linea = dict(r)
    precio = float(precio or 0)
    tc = float(tc or 0)
    moneda = "ARS" if (moneda or "").upper() in ("ARS", "$", "PESOS") else "USD"
    if precio < 0:
        raise ValueError("El precio no puede ser negativo.")
    cant = float(linea.get("cantidad_total") or 0)
    if cant <= 0:
        has = float(linea.get("cantidad_has") or 0)
        dosis = float(linea.get("dosis_ha") or 0) or (1.0 if _es_linea_laboreo(linea) else 0.0)
        cant = round(has * dosis, 4)
    if cant <= 0:
        raise ValueError("La línea no tiene cantidad (hectáreas × dosis).")
    if precio > 0 and tc <= 0:
        raise ValueError("Indicá el tipo de cambio para calcular pesos y dólares.")
    total = round(cant * precio, 2)
    if moneda == "USD":
        costo_usd, costo_ars = total, round(total * tc, 2)
        precio_ars, precio_usd = round(precio * tc, 4), precio
    else:
        costo_ars, costo_usd = total, round(total / tc, 2) if tc else 0.0
        precio_ars, precio_usd = precio, round(precio / tc, 4) if tc else 0.0

    if mov_ingreso_id:
        cursor.execute(
            "SELECT fecha, proveedor_nombre, nro_comprobante FROM almacen_movimientos WHERE id=? AND empresa_id=?;",
            (int(mov_ingreso_id), empresa_id),
        )
        m = cursor.fetchone()
        if m:
            fecha_compra = fecha_compra or (m["fecha"] or "")
            proveedor = proveedor or (m["proveedor_nombre"] or "")
            nro_comprobante = nro_comprobante or (m["nro_comprobante"] or "")

    cursor.execute(
        """
        UPDATE margenes_access
        SET cantidad_total=?, precio=?, tc=?, costo_ars=?, costo_usd=?, moneda_costo=?,
            origen_costo=?, mov_ingreso_id=?, proveedor=?, nro_remito=?, fecha_compra=?,
            fecha_costeo=?
        WHERE id=?;
        """,
        (
            cant, round(precio, 4), round(tc, 4), costo_ars, costo_usd, moneda,
            (origen or "manual")[:20], int(mov_ingreso_id) if mov_ingreso_id else None,
            proveedor or "", nro_comprobante or "", (fecha_compra or "")[:10],
            datetime.now().strftime("%Y-%m-%d %H:%M"), linea_id,
        ),
    )

    if linea.get("ot_consumo_id"):
        es_lab = _es_linea_laboreo(linea)
        cursor.execute(
            "UPDATE ot_consumos SET costo_unitario_neto=?, costo_total_neto=? WHERE id=?;",
            (
                precio_ars if es_lab else precio_usd,
                costo_ars if es_lab else costo_usd,
                linea["ot_consumo_id"],
            ),
        )
        cursor.execute(
            """
            SELECT m.id, m.cantidad FROM almacen_movimientos m
            JOIN ot_consumos oc ON oc.movimiento_id = m.id
            WHERE oc.id=?;
            """,
            (linea["ot_consumo_id"],),
        )
        mov = cursor.fetchone()
        if mov:
            cant_mov = abs(float(mov["cantidad"] or 0))
            importe_mov = costo_ars if abs(cant_mov - cant) < 1e-6 else round(cant_mov * precio_ars, 2)
            cursor.execute(
                "UPDATE almacen_movimientos SET precio_unitario_neto=?, importe_neto=?, precio_unitario_usd=? WHERE id=?;",
                (precio_ars, importe_mov, precio_usd, mov["id"]),
            )
    if linea.get("ot_id"):
        _recalcular_totales_ot(cursor, int(linea["ot_id"]))
    return {
        "status": "success",
        "id": linea_id,
        "cantidad": cant,
        "precio": round(precio, 4),
        "moneda": moneda,
        "tc": round(tc, 4),
        "costo_ars": costo_ars,
        "costo_usd": costo_usd,
    }


def costear_ot_por_criterio(cursor, empresa_id: int, nro_orden: int) -> Dict[str, Any]:
    """Costea los productos sin costo de una OT con el criterio de la empresa (PEPS por defecto).
    Los laboreos se siguen costeando a mano con la factura del contratista."""
    asegurar_schema_participaciones(cursor)
    cursor.execute(
        """
        SELECT m.* FROM margenes_access m
        LEFT JOIN ordenes_trabajo o ON o.id = m.ot_id
        WHERE m.empresa_id=? AND m.nro_orden=? AND TRIM(COALESCE(m.producto,'')) <> ''
          AND COALESCE(m.costo_ars,0)=0 AND COALESCE(m.costo_usd,0)=0
          AND (m.ot_id IS NULL OR COALESCE(o.estado,'') != 'Anulada')
        ORDER BY m.id;
        """,
        (empresa_id, int(nro_orden)),
    )
    lineas = [dict(r) for r in cursor.fetchall()]
    costeadas, omitidas, criterio = 0, [], None
    for linea in lineas:
        item = _item_de_linea(cursor, empresa_id, linea)
        if not item:
            omitidas.append(f"{linea['producto']}: no está en el almacén")
            continue
        mov_id = None
        if linea.get("ot_consumo_id"):
            cursor.execute("SELECT movimiento_id FROM ot_consumos WHERE id=?;", (linea["ot_consumo_id"],))
            r = cursor.fetchone()
            mov_id = r["movimiento_id"] if r else None
        sug = _sugerido_linea(cursor, empresa_id, linea, item, mov_id)
        if not sug or float(sug.get("costo_unitario") or 0) <= 0:
            omitidas.append(f"{linea['producto']}: sin compras con precio ni costo en el almacén")
            continue
        criterio = sug["criterio_nombre"]
        fecha = (linea.get("fecha_aplicacion") or linea.get("fecha_orden") or "")[:10]
        tc = float(linea.get("tc") or 0) or _tc_sugerido(cursor, empresa_id, fecha)
        if tc <= 0:
            omitidas.append(f"{linea['producto']}: falta el tipo de cambio")
            continue
        costear_linea_ot(cursor, empresa_id, int(linea["id"]), precio=float(sug["costo_unitario"]),
                         moneda=sug["moneda"], tc=tc, origen="criterio")
        costeadas += 1
    return {"status": "ok", "costeadas": costeadas, "omitidas": omitidas, "criterio": criterio}


def anular_ot(cursor, empresa_id: int, ot_id: int) -> Dict[str, Any]:
    """Anula una OT emitida en el sistema: devuelve el stock y borra sus líneas de costos."""
    cursor.execute("SELECT * FROM ordenes_trabajo WHERE id=? AND empresa_id=?;", (ot_id, empresa_id))
    ot = cursor.fetchone()
    if not ot:
        raise ValueError("OT no encontrada.")
    if (ot["estado"] or "") == "Anulada":
        raise ValueError("La OT ya está anulada.")
    if int(ot["confirmada"] or 0):
        raise ValueError("La OT está confirmada. Reabrila antes de anularla.")
    cursor.execute(
        "SELECT id, item_id, cantidad FROM almacen_movimientos WHERE ot_id=? AND tipo_mov='egreso_ot';",
        (ot_id,),
    )
    for m in cursor.fetchall():
        cursor.execute(
            "UPDATE almacen_items SET stock_cantidad=ROUND(COALESCE(stock_cantidad,0)+?, 6) WHERE id=?;",
            (abs(float(m["cantidad"] or 0)), m["item_id"]),
        )
        cursor.execute("DELETE FROM almacen_movimientos WHERE id=?;", (m["id"],))
    cursor.execute("UPDATE ot_consumos SET movimiento_id=NULL WHERE ot_id=?;", (ot_id,))
    cursor.execute("DELETE FROM margenes_access WHERE ot_id=?;", (ot_id,))
    cursor.execute(
        """
        UPDATE ordenes_trabajo
        SET estado='Anulada', costo_insumos_neto=0, costo_laboreos_neto=0, costo_total_neto=0
        WHERE id=?;
        """,
        (ot_id,),
    )
    return {"status": "success", "message": f"OT {ot['nro_ot']} anulada. El stock volvió al almacén."}


def costos_por_lote_campania(cursor, campania_id: int, empresa_id: int) -> List[Dict[str, Any]]:
    """Resumen de costos netos imputados por OT a cada lote de la campaña.
    En campos con participación cuenta solo las hectáreas que le corresponden a la empresa."""
    asegurar_schema_participaciones(cursor)
    filtro = """
        FROM ordenes_trabajo ot
        LEFT JOIN ot_destinos d ON d.ot_id = ot.id
        WHERE ot.empresa_id = ? AND ot.campania_id = ?
          AND COALESCE(ot.estado,'') != 'Anulada'
    """
    cursor.execute(
        f"""
        SELECT COALESCE(d.lote_id, ot.lote_id) AS lote_id,
               COALESCE(d.campo_id, ot.campo_id) AS campo_id,
               COALESCE(SUM(COALESCE(d.superficie_has, ot.superficie_has)
                            * COALESCE(d.participacion_pct, 100) / 100.0), 0) AS has_trabajadas
        {filtro}
        GROUP BY 1, 2;
        """,
        (empresa_id, campania_id),
    )
    grupos: Dict[Any, Dict[str, Any]] = {}
    for r in cursor.fetchall():
        grupos[(r["lote_id"], r["campo_id"])] = {
            "lote_id": r["lote_id"], "campo_id": r["campo_id"],
            "has_trabajadas": float(r["has_trabajadas"] or 0),
            "insumos": 0.0, "laboreos": 0.0, "total": 0.0,
        }
    cursor.execute(
        """
        SELECT COALESCE(d.lote_id, ot.lote_id) AS lote_id,
               COALESCE(d.campo_id, ot.campo_id) AS campo_id,
               COALESCE(SUM(CASE WHEN TRIM(COALESCE(m.producto,'')) != '' THEN m.costo_usd END),0) AS insumos,
               COALESCE(SUM(CASE WHEN TRIM(COALESCE(m.producto,'')) = '' THEN m.costo_ars END),0) AS laboreos,
               COALESCE(SUM(m.costo_usd),0) AS total
        FROM margenes_access m
        JOIN ordenes_trabajo ot ON ot.id = m.ot_id
        LEFT JOIN ot_consumos oc ON oc.id = m.ot_consumo_id
        LEFT JOIN ot_destinos d ON d.id = oc.destino_id
        WHERE ot.empresa_id = ? AND ot.campania_id = ?
          AND COALESCE(ot.estado,'') != 'Anulada'
        GROUP BY 1, 2;
        """,
        (empresa_id, campania_id),
    )
    for r in cursor.fetchall():
        g = grupos.setdefault((r["lote_id"], r["campo_id"]), {
            "lote_id": r["lote_id"], "campo_id": r["campo_id"], "has_trabajadas": 0.0,
            "insumos": 0.0, "laboreos": 0.0, "total": 0.0,
        })
        g["insumos"] += float(r["insumos"] or 0)
        g["laboreos"] += float(r["laboreos"] or 0)
        g["total"] += float(r["total"] or 0)
    rows = []
    for g in grupos.values():
        g["lote_nombre"] = ""
        g["campo_nombre"] = ""
        if g["lote_id"]:
            cursor.execute("SELECT nombre FROM lotes_agro WHERE id=?;", (g["lote_id"],))
            r = cursor.fetchone()
            g["lote_nombre"] = (r["nombre"] if r else "") or ""
        if g["campo_id"]:
            cursor.execute("SELECT nombre FROM campos_agro WHERE id=?;", (g["campo_id"],))
            r = cursor.fetchone()
            g["campo_nombre"] = (r["nombre"] if r else "") or ""
        for k in ("insumos", "laboreos", "total"):
            g[k] = round(g[k], 2)
        has = g["has_trabajadas"]
        g["costo_por_ha"] = round(g["total"] / has, 2) if has > 0 else 0
        rows.append(g)
    rows.sort(key=lambda x: ((x["campo_nombre"] or "").lower(), (x["lote_nombre"] or "").lower()))
    return rows


def resumen_participacion_campania(cursor, empresa_id: int, campania_id: int) -> Dict[str, Any]:
    """Qué aportó y qué le corresponde a la empresa en los campos con participación.
    Solo toma líneas costeadas; diferencia > 0 = el socio le debe a la empresa."""
    asegurar_schema_participaciones(cursor)
    cursor.execute(
        """
        SELECT o.id AS ot_id, o.nro_ot, o.fecha, o.fecha_aplicacion,
               d.campo_id, d.lote_id, c.nombre AS campo, l.nombre AS lote,
               d.superficie_has AS has_fisicas, d.cultivo,
               COALESCE(d.participacion_pct, 100) AS participacion_pct,
               COALESCE(d.socios_txt, '') AS socios_txt,
               oc.id AS consumo_id, oc.tipo_item, oc.cantidad, oc.unidad,
               COALESCE(oc.aporte_pct, 100) AS aporte_pct,
               COALESCE(oc.aporta, '') AS aporta,
               i.nombre AS item,
               m.cantidad_total AS cantidad_empresa, m.costo_usd, m.costo_ars,
               m.proveedor, m.nro_remito
        FROM ot_consumos oc
        JOIN ordenes_trabajo o ON o.id = oc.ot_id
        JOIN ot_destinos d ON d.id = oc.destino_id
        LEFT JOIN campos_agro c ON c.id = d.campo_id
        LEFT JOIN lotes_agro l ON l.id = d.lote_id
        LEFT JOIN almacen_items i ON i.id = oc.item_id
        LEFT JOIN margenes_access m ON m.ot_consumo_id = oc.id
        WHERE o.empresa_id=? AND o.campania_id=? AND COALESCE(o.estado,'') != 'Anulada'
          AND COALESCE(d.participacion_pct, 100) < 100
        ORDER BY c.nombre COLLATE NOCASE, o.fecha, o.nro_ot, l.nombre COLLATE NOCASE, oc.id;
        """,
        (empresa_id, campania_id),
    )
    lineas: List[Dict[str, Any]] = []
    pendientes: List[Dict[str, Any]] = []
    grupos: Dict[Any, Dict[str, Any]] = {}
    for r in cursor.fetchall():
        x = dict(r)
        costo_usd = float(x.get("costo_usd") or 0)
        costo_ars = float(x.get("costo_ars") or 0)
        cant_emp = float(x.get("cantidad_empresa") or 0)
        cant = float(x.get("cantidad") or 0)
        if not (costo_usd or costo_ars) or cant_emp <= 0:
            pendientes.append({k: x.get(k) for k in ("nro_ot", "campo", "lote", "item", "tipo_item")})
            continue
        part = float(x["participacion_pct"])
        aporte = float(x["aporte_pct"])
        for mon, costo in (("usd", costo_usd), ("ars", costo_ars)):
            total = cant * costo / cant_emp
            x[f"total_{mon}"] = round(total, 2)
            x[f"aporte_empresa_{mon}"] = round(total * aporte / 100, 2)
            x[f"corresponde_empresa_{mon}"] = round(total * part / 100, 2)
            x[f"aporte_socio_{mon}"] = round(x[f"total_{mon}"] - x[f"aporte_empresa_{mon}"], 2)
            x[f"corresponde_socio_{mon}"] = round(x[f"total_{mon}"] - x[f"corresponde_empresa_{mon}"], 2)
            x[f"diferencia_{mon}"] = round(x[f"aporte_empresa_{mon}"] - x[f"corresponde_empresa_{mon}"], 2)
        lineas.append(x)
        g = grupos.setdefault((x.get("campo") or "", x["socios_txt"]), {
            "campo": x.get("campo") or "", "socios_txt": x["socios_txt"],
            "participacion_pct": part, "lineas": 0,
            **{f"{k}_{mon}": 0.0 for k in ("total", "aporte_empresa", "corresponde_empresa",
                                           "aporte_socio", "corresponde_socio", "diferencia")
               for mon in ("usd", "ars")},
        })
        g["lineas"] += 1
        for k in list(g.keys()):
            if k.endswith("_usd") or k.endswith("_ars"):
                g[k] = round(g[k] + x[k], 2)
    return {
        "campania_id": campania_id,
        "lineas": lineas,
        "resumen": sorted(grupos.values(), key=lambda g: (g["campo"].lower(), g["socios_txt"].lower())),
        "pendientes": pendientes,
    }


def aporta_legible(aporta: Optional[str], aporte_pct: Any, empresa: str = "Empresa") -> str:
    """'empresa' | 'compartido' | 'socio:Nombre' | 'manual' -> texto para mostrar."""
    a = (aporta or "").strip()
    pct = float(aporte_pct if aporte_pct is not None else 100)
    if a == "empresa" or (not a and pct >= 100):
        return empresa
    if a == "compartido":
        return f"Compartido ({empresa} {pct:g}%)"
    if a.startswith("socio:"):
        return a[6:].strip() or "Socio"
    if not a and pct <= 0:
        return "Socio"
    return f"{empresa} {pct:g}%"


def excel_participacion_campania(cursor, empresa_id: int, campania_id: int) -> bytes:
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    data = resumen_participacion_campania(cursor, empresa_id, campania_id)
    cursor.execute("SELECT codigo FROM campanias_agro WHERE id=?;", (campania_id,))
    r = cursor.fetchone()
    codigo = (r["codigo"] if r else "") or ""
    empresa = "Empresa"
    try:
        cursor.execute("SELECT razon_social FROM empresas WHERE id=?;", (empresa_id,))
        r = cursor.fetchone()
        if r and (r["razon_social"] or "").strip():
            empresa = r["razon_social"].strip()
    except sqlite3.Error:
        pass

    wb = Workbook()
    negrita = Font(bold=True)
    fondo = PatternFill("solid", fgColor="E2E8F0")
    money = '#,##0.00'

    def encabezado(ws, fila, titulos):
        for col, t in enumerate(titulos, 1):
            celda = ws.cell(row=fila, column=col, value=t)
            celda.font = negrita
            celda.fill = fondo
            celda.alignment = Alignment(wrap_text=True, vertical="center")

    def anchos(ws, valores):
        for i, w in enumerate(valores, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    ws = wb.active
    ws.title = "Resumen"
    ws["A1"] = f"Participación con otras sociedades | Campaña {codigo} | {empresa}"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "Solo líneas costeadas. Diferencia positiva: el socio le debe a la empresa; negativa: la empresa le debe al socio."
    cols = ["Campo", "Socios", f"% {empresa}", "Líneas",
            "Total U$S", f"Aportó {empresa} U$S", f"Corresponde {empresa} U$S", "Diferencia U$S",
            "Total $", f"Aportó {empresa} $", f"Corresponde {empresa} $", "Diferencia $"]
    encabezado(ws, 4, cols)
    fila = 5
    for g in data["resumen"]:
        valores = [g["campo"], g["socios_txt"], g["participacion_pct"], g["lineas"],
                   g["total_usd"], g["aporte_empresa_usd"], g["corresponde_empresa_usd"], g["diferencia_usd"],
                   g["total_ars"], g["aporte_empresa_ars"], g["corresponde_empresa_ars"], g["diferencia_ars"]]
        for col, v in enumerate(valores, 1):
            c = ws.cell(row=fila, column=col, value=v)
            if col >= 5:
                c.number_format = money
        fila += 1
    if data["pendientes"]:
        fila += 1
        ws.cell(row=fila, column=1, value=f"Líneas sin costear (no incluidas): {len(data['pendientes'])}").font = negrita
    anchos(ws, [22, 30, 11, 8, 14, 16, 18, 15, 16, 18, 20, 16])

    ws2 = wb.create_sheet("Detalle")
    cols2 = ["OT", "Fecha", "Campo", "Lote", "Cultivo", "Has físicas", "Tipo", "Producto / labor",
             "Cantidad física", "Unidad", "Aporta", f"% aporte {empresa}", f"% {empresa} en el campo", "Socios",
             "Total U$S", f"Aportó {empresa} U$S", f"Corresponde {empresa} U$S", "Diferencia U$S",
             "Total $", f"Aportó {empresa} $", f"Corresponde {empresa} $", "Diferencia $"]
    encabezado(ws2, 1, cols2)
    for fila, x in enumerate(data["lineas"], 2):
        f = (x.get("fecha_aplicacion") or x.get("fecha") or "")[:10]
        valores = [x.get("nro_ot"), "/".join(reversed(f.split("-"))) if f else "", x.get("campo"), x.get("lote"),
                   x.get("cultivo"), x.get("has_fisicas"),
                   "Labor" if x.get("tipo_item") == "laboreo" else "Producto", x.get("item"),
                   x.get("cantidad"), x.get("unidad"), aporta_legible(x.get("aporta"), x.get("aporte_pct"), empresa),
                   x.get("aporte_pct"), x.get("participacion_pct"), x.get("socios_txt"),
                   x["total_usd"], x["aporte_empresa_usd"], x["corresponde_empresa_usd"], x["diferencia_usd"],
                   x["total_ars"], x["aporte_empresa_ars"], x["corresponde_empresa_ars"], x["diferencia_ars"]]
        for col, v in enumerate(valores, 1):
            c = ws2.cell(row=fila, column=col, value=v)
            if col >= 15:
                c.number_format = money
    anchos(ws2, [8, 11, 18, 14, 12, 11, 10, 30, 13, 8, 22, 12, 12, 26, 13, 15, 17, 14, 15, 17, 19, 15])
    ws2.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Unificación de productos: varios nombres comerciales → un solo producto/stock.
# ---------------------------------------------------------------------------

_UNIDADES_EQUIV = {
    "l": "lt", "lt": "lt", "lts": "lt", "litro": "lt", "litros": "lt",
    "kg": "kg", "kgs": "kg", "kilo": "kg", "kilos": "kg",
    "u": "un", "un": "un", "und": "un", "unidad": "un", "unidades": "un",
    "ha": "ha", "has": "ha", "hs": "hs", "hora": "hs", "horas": "hs",
    "g": "gr", "gr": "gr", "grs": "gr",
}


def _unidad_normalizada(u: Any) -> str:
    s = str(u or "").strip().lower().rstrip(".")
    return _UNIDADES_EQUIV.get(s, s)


def asegurar_schema_unificacion(cursor) -> None:
    cols = {r[1] for r in cursor.execute("PRAGMA table_info(almacen_items);").fetchall()}
    if cols and "unificado_en" not in cols:
        cursor.execute("ALTER TABLE almacen_items ADD COLUMN unificado_en INTEGER;")
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS almacen_alias (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            empresa_id INTEGER DEFAULT 1,
            item_id INTEGER NOT NULL,
            nombre TEXT NOT NULL,
            origen_item_id INTEGER,
            unificacion_id INTEGER
        );
        """
    )
    cursor.execute("CREATE INDEX IF NOT EXISTS idx_almacen_alias_nombre ON almacen_alias(empresa_id, nombre COLLATE NOCASE);")
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS almacen_unificaciones (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            empresa_id INTEGER DEFAULT 1,
            fecha TEXT,
            usuario TEXT,
            principal_id INTEGER NOT NULL,
            principal_nombre TEXT,
            unificados TEXT,
            datos TEXT,
            deshecha INTEGER DEFAULT 0
        );
        """
    )
    if cursor.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='margenes_access';").fetchone():
        cols_m = {r[1] for r in cursor.execute("PRAGMA table_info(margenes_access);").fetchall()}
        if "producto_original" not in cols_m:
            cursor.execute("ALTER TABLE margenes_access ADD COLUMN producto_original TEXT;")
        if "campania_original" not in cols_m:
            cursor.execute("ALTER TABLE margenes_access ADD COLUMN campania_original TEXT;")
        if "laboreo_original" not in cols_m:
            cursor.execute("ALTER TABLE margenes_access ADD COLUMN laboreo_original TEXT;")
    for tabla, col in (
        ("almacen_movimientos", "item_id"),
        ("ot_consumos", "item_id"),
        ("factura_imputaciones", "item_almacen_id"),
        ("margenes_access", "almacen_item_id"),
    ):
        if col in {r[1] for r in cursor.execute(f"PRAGMA table_info({tabla});").fetchall()}:
            cursor.execute(f"CREATE INDEX IF NOT EXISTS idx_{tabla}_{col} ON {tabla}({col});")


def _activo(item: Dict[str, Any]) -> bool:
    return item.get("activo") is None or int(item.get("activo") or 0) == 1


def _seguir_unificacion(cursor, item: Optional[Dict[str, Any]]) -> Optional[Dict[str, Any]]:
    for _ in range(10):
        if not item or not item.get("unificado_en"):
            return item
        r = cursor.execute("SELECT * FROM almacen_items WHERE id=?;", (item["unificado_en"],)).fetchone()
        if not r:
            return item
        item = dict(r)
    return item


def resolver_item_por_nombre(cursor, empresa_id: int, nombre: str) -> Optional[Dict[str, Any]]:
    """Ítem por nombre exacto; si ese nombre quedó unificado (o es un alias), devuelve el producto principal."""
    nombre = (nombre or "").strip()
    if not nombre:
        return None
    r = cursor.execute(
        """
        SELECT * FROM almacen_items
        WHERE empresa_id=? AND UPPER(TRIM(nombre))=UPPER(TRIM(?))
        ORDER BY COALESCE(activo,1) DESC, id LIMIT 1;
        """,
        (empresa_id, nombre),
    ).fetchone()
    item = dict(r) if r else None
    if item and _activo(item) and not item.get("unificado_en"):
        return item
    try:
        a = cursor.execute(
            """
            SELECT i.* FROM almacen_alias a JOIN almacen_items i ON i.id = a.item_id
            WHERE a.empresa_id=? AND UPPER(TRIM(a.nombre))=UPPER(TRIM(?))
            ORDER BY a.id DESC LIMIT 1;
            """,
            (empresa_id, nombre),
        ).fetchone()
    except sqlite3.OperationalError:
        a = None
    if a:
        return _seguir_unificacion(cursor, dict(a))
    return _seguir_unificacion(cursor, item)


def _tabla_existe(cursor, nombre: str) -> bool:
    return bool(cursor.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name=?;", (nombre,)).fetchone())


def _items_para_unificar(cursor, empresa_id: int, principal_id: int, otros_ids: List[int]):
    ids = [int(principal_id)] + [int(x) for x in otros_ids if int(x) != int(principal_id)]
    ids = list(dict.fromkeys(ids))
    if len(ids) < 2:
        raise ValueError("Elegí el producto principal y al menos otro para unificar.")
    marcas = ",".join("?" * len(ids))
    filas = cursor.execute(
        f"SELECT * FROM almacen_items WHERE empresa_id=? AND id IN ({marcas});", [empresa_id] + ids
    ).fetchall()
    por_id = {int(r["id"]): dict(r) for r in filas}
    faltan = [i for i in ids if i not in por_id]
    if faltan:
        raise ValueError(f"No se encontraron los productos {faltan}.")
    principal = por_id[ids[0]]
    otros = [por_id[i] for i in ids[1:]]
    return principal, otros


def _columna_nombre_linea(item: Dict[str, Any]) -> str:
    """En las líneas de costo el producto va en 'producto'; un laboreo sin producto, en 'laboreo'."""
    return "laboreo" if (item.get("tipo") or "producto") == "laboreo" else "producto"


def _lineas_campania_de(cursor, empresa_id: int, otros: List[Dict[str, Any]], col: str = "producto") -> List[Dict[str, Any]]:
    if not _tabla_existe(cursor, "margenes_access"):
        return []
    ids = [int(o["id"]) for o in otros]
    nombres = [str(o["nombre"] or "").strip().upper() for o in otros if str(o["nombre"] or "").strip()]
    m_ids = ",".join("?" * len(ids))
    m_nom = ",".join("?" * len(nombres)) or "''"
    solo_labor = " AND TRIM(COALESCE(producto,''))=''" if col == "laboreo" else ""
    filas = cursor.execute(
        f"""
        SELECT id, campania_codigo, {col} AS nombre, almacen_item_id FROM margenes_access
        WHERE empresa_id=? AND (
            almacen_item_id IN ({m_ids})
            OR (almacen_item_id IS NULL AND UPPER(TRIM({col})) IN ({m_nom}){solo_labor})
        );
        """,
        [empresa_id] + ids + nombres,
    ).fetchall()
    return [dict(r) for r in filas]


def vista_previa_unificacion(cursor, empresa_id: int, principal_id: int, otros_ids: List[int]) -> Dict[str, Any]:
    asegurar_schema_unificacion(cursor)
    principal, otros = _items_para_unificar(cursor, empresa_id, principal_id, otros_ids)
    errores = []
    for o in otros:
        if (o.get("tipo") or "producto") != (principal.get("tipo") or "producto"):
            errores.append(f'"{o["nombre"]}" es {o.get("tipo")} y el principal es {principal.get("tipo")}.')
        if _unidad_normalizada(o.get("unidad")) != _unidad_normalizada(principal.get("unidad")):
            errores.append(
                f'"{o["nombre"]}" está en {o.get("unidad") or "?"} y el principal en {principal.get("unidad") or "?"}: '
                "corregí la unidad en la ficha antes de unificar, si no el stock se mezcla mal."
            )
        if o.get("unificado_en"):
            errores.append(f'"{o["nombre"]}" ya está unificado en otro producto.')
    if principal.get("unificado_en") or not _activo(principal):
        errores.append(f'El principal "{principal["nombre"]}" no está activo.')
    ids = [int(o["id"]) for o in otros]
    marcas = ",".join("?" * len(ids))

    def contar(tabla, col):
        if not _tabla_existe(cursor, tabla):
            return 0
        return int(cursor.execute(f"SELECT COUNT(*) FROM {tabla} WHERE {col} IN ({marcas});", ids).fetchone()[0])

    lineas = _lineas_campania_de(cursor, empresa_id, otros, _columna_nombre_linea(principal))
    por_campania: Dict[str, int] = {}
    for l in lineas:
        k = l["campania_codigo"] or "(sin campaña)"
        por_campania[k] = por_campania.get(k, 0) + 1
    stock_total = float(principal.get("stock_cantidad") or 0) + sum(float(o.get("stock_cantidad") or 0) for o in otros)
    precios = precios_peps_unificados(cursor, empresa_id, [principal] + otros)
    return {
        "principal": {k: principal.get(k) for k in ("id", "nombre", "unidad", "stock_cantidad", "costo_promedio_usd", "costo_promedio_neto")},
        "otros": [{k: o.get(k) for k in ("id", "nombre", "unidad", "stock_cantidad", "costo_promedio_usd", "costo_promedio_neto")} for o in otros],
        "movimientos": contar("almacen_movimientos", "item_id"),
        "consumos_ot": contar("ot_consumos", "item_id"),
        "renglones_factura": contar("factura_imputaciones", "item_almacen_id"),
        "lineas_campania": len(lineas),
        "lineas_por_campania": dict(sorted(por_campania.items(), reverse=True)),
        "stock_resultante": round(stock_total, 4),
        "precios_resultantes": precios,
        "costos_resultantes": _costos_unificados(principal, otros, precios),
        "errores": errores,
    }


def _costos_unificados(principal: Dict[str, Any], otros: List[Dict[str, Any]], precios: Dict[str, Any]) -> Dict[str, float]:
    """Sin promediar: el costo queda en el primero entrado (el próximo en salir por PEPS).
    Sin compras con precio, se queda el costo del principal (o el primero con costo)."""
    todos = [principal] + otros
    out = {}
    for campo, clave in (("costo_promedio_neto", "primero_ars"), ("costo_promedio_usd", "primero_usd")):
        valor = float(precios.get(clave) or 0)
        if valor <= 0:
            valor = next((float(i.get(campo) or 0) for i in todos if float(i.get(campo) or 0) > 0), 0.0)
        out[campo] = round(valor, 6)
    return out


def unificar_items(cursor, empresa_id: int, principal_id: int, otros_ids: List[int], usuario: str = "") -> Dict[str, Any]:
    """Pasa movimientos, consumos de OT, renglones de factura y líneas de todas las campañas al principal.
    Los otros quedan inactivos y sus nombres como alias (las próximas facturas con ese nombre van al principal)."""
    import json

    previa = vista_previa_unificacion(cursor, empresa_id, principal_id, otros_ids)
    if previa["errores"]:
        raise ValueError(" ".join(previa["errores"]))
    principal, otros = _items_para_unificar(cursor, empresa_id, principal_id, otros_ids)
    pid = int(principal["id"])
    ids = [int(o["id"]) for o in otros]
    marcas = ",".join("?" * len(ids))
    datos: Dict[str, Any] = {
        "principal": {k: principal.get(k) for k in ("stock_cantidad", "costo_promedio_neto", "costo_promedio_usd")},
        "items": {str(o["id"]): {k: o.get(k) for k in ("nombre", "activo", "stock_cantidad", "costo_promedio_neto", "costo_promedio_usd")} for o in otros},
    }
    for clave, tabla, col in (
        ("movimientos", "almacen_movimientos", "item_id"),
        ("consumos_ot", "ot_consumos", "item_id"),
        ("renglones_factura", "factura_imputaciones", "item_almacen_id"),
        ("alias", "almacen_alias", "item_id"),
    ):
        if not _tabla_existe(cursor, tabla):
            datos[clave] = []
            continue
        filas = cursor.execute(f"SELECT id, {col} FROM {tabla} WHERE {col} IN ({marcas});", ids).fetchall()
        datos[clave] = [[int(r[0]), int(r[1])] for r in filas]
        cursor.execute(f"UPDATE {tabla} SET {col}=? WHERE {col} IN ({marcas});", [pid] + ids)

    col = _columna_nombre_linea(principal)
    lineas = _lineas_campania_de(cursor, empresa_id, otros, col)
    datos["lineas"] = [[l["id"], l["nombre"], l["almacen_item_id"], col] for l in lineas]
    for l in lineas:
        cursor.execute(
            f"""
            UPDATE margenes_access
            SET {col}_original = COALESCE({col}_original, {col}), {col} = ?, almacen_item_id = ?
            WHERE id = ?;
            """,
            (principal["nombre"], pid, l["id"]),
        )

    costos = previa["costos_resultantes"]
    cursor.execute(
        "UPDATE almacen_items SET stock_cantidad=?, costo_promedio_neto=?, costo_promedio_usd=? WHERE id=?;",
        (previa["stock_resultante"], costos["costo_promedio_neto"], costos["costo_promedio_usd"], pid),
    )
    cursor.execute(
        f"UPDATE almacen_items SET activo=0, unificado_en=?, stock_cantidad=0 WHERE id IN ({marcas});",
        [pid] + ids,
    )
    cursor.execute(
        """
        INSERT INTO almacen_unificaciones (empresa_id, fecha, usuario, principal_id, principal_nombre, unificados, datos)
        VALUES (?, ?, ?, ?, ?, ?, ?);
        """,
        (
            empresa_id, datetime.now().strftime("%Y-%m-%d %H:%M:%S"), usuario or "", pid, principal["nombre"],
            json.dumps([o["nombre"] for o in otros], ensure_ascii=False), json.dumps(datos, ensure_ascii=False),
        ),
    )
    uid = cursor.lastrowid
    for o in otros:
        if str(o["nombre"] or "").strip().upper() != str(principal["nombre"] or "").strip().upper():
            cursor.execute(
                "INSERT INTO almacen_alias (empresa_id, item_id, nombre, origen_item_id, unificacion_id) VALUES (?, ?, ?, ?, ?);",
                (empresa_id, pid, str(o["nombre"]).strip(), int(o["id"]), uid),
            )
    return {
        "status": "ok",
        "unificacion_id": uid,
        "principal": principal["nombre"],
        "unificados": [o["nombre"] for o in otros],
        "movimientos": len(datos["movimientos"]),
        "consumos_ot": len(datos["consumos_ot"]),
        "renglones_factura": len(datos["renglones_factura"]),
        "lineas_campania": len(lineas),
        "stock_resultante": previa["stock_resultante"],
    }


def listar_unificaciones(cursor, empresa_id: int) -> List[Dict[str, Any]]:
    import json

    asegurar_schema_unificacion(cursor)
    filas = cursor.execute(
        """
        SELECT id, fecha, usuario, principal_id, principal_nombre, unificados, deshecha
        FROM almacen_unificaciones WHERE empresa_id=? ORDER BY id DESC;
        """,
        (empresa_id,),
    ).fetchall()
    out = []
    for r in filas:
        d = dict(r)
        try:
            d["unificados"] = json.loads(d["unificados"] or "[]")
        except ValueError:
            d["unificados"] = []
        out.append(d)
    return out


def deshacer_unificacion(cursor, empresa_id: int, unificacion_id: int) -> Dict[str, Any]:
    """Vuelve cada movimiento/consumo/renglón/línea a su producto original y reactiva los productos."""
    import json

    asegurar_schema_unificacion(cursor)
    r = cursor.execute(
        "SELECT * FROM almacen_unificaciones WHERE id=? AND empresa_id=?;", (unificacion_id, empresa_id)
    ).fetchone()
    if not r:
        raise ValueError("No se encontró esa unificación.")
    if int(r["deshecha"] or 0):
        raise ValueError("Esa unificación ya se deshizo.")
    datos = json.loads(r["datos"] or "{}")
    pid = int(r["principal_id"])
    principal = cursor.execute("SELECT * FROM almacen_items WHERE id=?;", (pid,)).fetchone()
    if not principal or principal["unificado_en"]:
        raise ValueError("El producto principal se unificó después en otro: deshacé primero esa unificación.")
    for oid in datos.get("items", {}):
        it = cursor.execute("SELECT unificado_en FROM almacen_items WHERE id=?;", (int(oid),)).fetchone()
        if not it or int(it["unificado_en"] or 0) != pid:
            raise ValueError("Algún producto unificado cambió después; no se puede deshacer automáticamente.")
    for clave, tabla, col in (
        ("movimientos", "almacen_movimientos", "item_id"),
        ("consumos_ot", "ot_consumos", "item_id"),
        ("renglones_factura", "factura_imputaciones", "item_almacen_id"),
        ("alias", "almacen_alias", "item_id"),
    ):
        if _tabla_existe(cursor, tabla):
            for fila_id, original in datos.get(clave, []):
                cursor.execute(f"UPDATE {tabla} SET {col}=? WHERE id=?;", (original, fila_id))
    for linea in datos.get("lineas", []):
        lid, valor, item_original = linea[:3]
        col = "laboreo" if len(linea) > 3 and linea[3] == "laboreo" else "producto"
        cursor.execute(
            f"UPDATE margenes_access SET {col}=?, almacen_item_id=?, {col}_original=NULL WHERE id=?;",
            (valor, item_original, lid),
        )
    devuelto = 0.0
    for oid, it in datos.get("items", {}).items():
        devuelto += float(it.get("stock_cantidad") or 0)
        cursor.execute(
            """
            UPDATE almacen_items SET activo=?, unificado_en=NULL, stock_cantidad=?, costo_promedio_neto=?, costo_promedio_usd=?
            WHERE id=?;
            """,
            (it.get("activo") if it.get("activo") is not None else 1, it.get("stock_cantidad") or 0,
             it.get("costo_promedio_neto") or 0, it.get("costo_promedio_usd") or 0, int(oid)),
        )
    p = datos.get("principal", {})
    cursor.execute(
        """
        UPDATE almacen_items SET stock_cantidad=ROUND(COALESCE(stock_cantidad,0)-?, 6),
               costo_promedio_neto=?, costo_promedio_usd=?
        WHERE id=?;
        """,
        (devuelto, p.get("costo_promedio_neto") or 0, p.get("costo_promedio_usd") or 0, pid),
    )
    cursor.execute("DELETE FROM almacen_alias WHERE unificacion_id=?;", (unificacion_id,))
    cursor.execute("UPDATE almacen_unificaciones SET deshecha=1 WHERE id=?;", (unificacion_id,))
    return {"status": "ok", "message": f"Se deshizo la unificación en \"{r['principal_nombre']}\"."}
