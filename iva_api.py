# -*- coding: utf-8 -*-
"""Posición mensual de IVA (esquema F.2051 "IVA Simple").

Débito fiscal − crédito fiscal − saldo técnico a favor anterior = impuesto determinado.
Impuesto determinado − retenciones/percepciones/pagos a cuenta − saldo de libre
disponibilidad anterior (neto de devoluciones) = saldo a pagar o a favor.
"""
from __future__ import annotations

import base64
import json
import re
import unicodedata
from datetime import date, datetime, timedelta
from io import BytesIO
from typing import Optional

from fastapi import HTTPException, Query
from fastapi.responses import Response
from pydantic import BaseModel

import iva_arca

SECCIONES = ("debito", "credito", "retencion", "percepcion", "devolucion")
TIPOS_AJUSTE = SECCIONES + ("pago_cuenta",)
ALICUOTAS = (0.21, 0.105, 0.27, 0.05, 0.025)
ORIGENES_ITEMS = ("cc", "cc_perc", "banco", "granos", "hacienda", "leche", "arca")


# ---------------------------------------------------------------- esquema

def init_iva_schema(cur) -> None:
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS iva_posiciones (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            empresa_id INTEGER NOT NULL,
            periodo TEXT NOT NULL,
            st_anterior REAL,
            sld_anterior REAL,
            pagos_cuenta REAL DEFAULT 0,
            devoluciones REAL DEFAULT 0,
            estado TEXT DEFAULT 'borrador',
            fecha_presentacion TEXT,
            nro_transaccion TEXT,
            importe_pagado REAL DEFAULT 0,
            fecha_pago TEXT,
            observaciones TEXT,
            snapshot_json TEXT,
            actualizado_en TEXT,
            UNIQUE(empresa_id, periodo)
        );
        """
    )
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS iva_ajustes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            empresa_id INTEGER NOT NULL,
            periodo TEXT NOT NULL,
            tipo TEXT NOT NULL,
            concepto TEXT,
            alicuota REAL DEFAULT 0,
            neto REAL DEFAULT 0,
            importe REAL DEFAULT 0,
            creado_en TEXT
        );
        """
    )
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS iva_exclusiones (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            empresa_id INTEGER NOT NULL,
            origen TEXT NOT NULL,
            origen_id INTEGER NOT NULL,
            periodo_destino TEXT DEFAULT '',
            motivo TEXT,
            creado_en TEXT,
            UNIQUE(empresa_id, origen, origen_id)
        );
        """
    )
    iva_arca.init_arca_schema(cur)


# ---------------------------------------------------------------- utilidades

def _norm(s) -> str:
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s.lower()).strip()


def _f(v) -> float:
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


def _r2(v) -> float:
    return round(float(v or 0) + 0.0, 2)


def _fmt(v) -> str:
    return f"{_r2(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _iso(fecha) -> str:
    s = str(fecha or "").strip()
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return m.group(0)
    m = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{4})", s)
    if m:
        return f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"
    return s[:10]


def _periodo_de(fecha) -> str:
    s = str(fecha or "").strip()
    m = re.match(r"^(\d{1,2})/(\d{4})$", s)
    if m:
        return f"{m.group(2)}-{int(m.group(1)):02d}"
    iso = _iso(s)
    return iso[:7] if re.match(r"^\d{4}-\d{2}", iso) else ""


def _validar_periodo(p: str) -> str:
    p = (p or "").strip()
    if not re.match(r"^\d{4}-(0[1-9]|1[0-2])$", p):
        raise HTTPException(400, "Período inválido (formato AAAA-MM).")
    return p


def _periodo_mas(p: str, n: int) -> str:
    y, m = int(p[:4]), int(p[5:7])
    idx = y * 12 + (m - 1) + n
    return f"{idx // 12:04d}-{idx % 12 + 1:02d}"


def _patrones_fecha(p: str):
    return (f"{p}%", f"%/{p[5:7]}/{p[:4]}", f"%/{int(p[5:7])}/{p[:4]}")


def _es_nc(tipo) -> bool:
    t = _norm(tipo)
    return "nota de credito" in t or t.startswith("nc ") or t == "nc" or bool(re.match(r"^nc\d", t))


def _es_comprobante_iva(tipo) -> bool:
    t = _norm(tipo)
    return bool(re.search(r"factura|nota de credito|nota de debito|ticket|liquidacion|^n[cd]\b|^fc\b", t))


def _letra(tipo) -> str:
    m = re.search(r"\b([abcem])\s*$", _norm(tipo))
    return m.group(1).upper() if m else ""


def _alicuota_inferida(neto: float, iva: float) -> Optional[float]:
    if not neto or not iva:
        return 0.0 if neto and not iva else None
    ratio = iva / neto
    for a in ALICUOTAS:
        if abs(ratio - a) <= 0.0015:
            return a
    return None


def _tabla_existe(cur, nombre: str) -> bool:
    cur.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name=?;", (nombre,))
    return cur.fetchone() is not None


def _cols(cur, tabla: str) -> set:
    cur.execute(f"PRAGMA table_info({tabla});")
    return {r[1] for r in cur.fetchall()}


def _vencimiento(periodo: str, cuit: str) -> Optional[str]:
    """Estimado: día 18 del mes siguiente (terminación 0-1) y los cuatro días hábiles
    siguientes para 2-3, 4-5, 6-7 y 8-9. No contempla feriados."""
    digitos = re.sub(r"\D", "", cuit or "")
    if not digitos:
        return None
    grupo = int(digitos[-1]) // 2
    sig = _periodo_mas(periodo, 1)
    d = date(int(sig[:4]), int(sig[5:7]), 18)
    while d.weekday() >= 5:
        d += timedelta(days=1)
    for _ in range(grupo):
        d += timedelta(days=1)
        while d.weekday() >= 5:
            d += timedelta(days=1)
    return d.isoformat()


def _empresa(cur, empresa_id: int) -> dict:
    razon, cuit = "", ""
    try:
        cur.execute("SELECT razon_social, cuit FROM empresas WHERE id=?;", (empresa_id,))
        r = cur.fetchone()
        if r:
            razon, cuit = r["razon_social"] or "", r["cuit"] or ""
    except Exception:
        pass
    if not cuit:
        try:
            cur.execute("SELECT * FROM configuracion_empresa WHERE id=1;")
            r = cur.fetchone()
            if r:
                keys = r.keys()
                cuit = (r["cuit"] if "cuit" in keys else "") or ""
                razon = razon or ((r["razon_social"] if "razon_social" in keys else "") or "")
        except Exception:
            pass
    return {"razon_social": razon, "cuit": re.sub(r"\D", "", cuit)}


# ---------------------------------------------------------------- fuentes

def _overrides(cur, empresa_id: int) -> dict:
    cur.execute(
        "SELECT id, origen, origen_id, periodo_destino, motivo FROM iva_exclusiones WHERE empresa_id=?;",
        (empresa_id,),
    )
    return {(r["origen"], int(r["origen_id"])): dict(r) for r in cur.fetchall()}


def _ids_movidos_a(ovr: dict, origen: str, periodo: str) -> list:
    return [oid for (o, oid), v in ovr.items() if o == origen and (v["periodo_destino"] or "") == periodo]


def _in_clause(ids: list) -> str:
    return ",".join(str(int(i)) for i in ids) if ids else "NULL"


def _items_cuentas_corrientes(cur, empresa_id: int, periodo: str, ovr: dict) -> list:
    cols = _cols(cur, "cuentas_corrientes")
    col_libro = "cc.libro_iva" if "libro_iva" in cols else "NULL"
    extras = ", ".join(f"COALESCE(cc.{c},0) AS {c}" if c in cols else f"0 AS {c}"
                       for c in ("percepcion_iva", "no_gravado", "exento"))
    tiene_asientos = _tabla_existe(cur, "asientos_contables") and "asiento_id" in cols
    join_as = "LEFT JOIN asientos_contables a ON a.id = cc.asiento_id" if tiene_asientos else ""
    col_origen = "a.origen_modulo" if tiene_asientos else "NULL"
    p1, p2, p3 = _patrones_fecha(periodo)
    movidos = _ids_movidos_a(ovr, "cc", periodo) + _ids_movidos_a(ovr, "cc_perc", periodo)
    cur.execute(
        f"""
        SELECT cc.id, cc.entidad_id, cc.tipo_comprobante, cc.numero_comprobante, cc.fecha,
               COALESCE(cc.neto,0) AS neto, COALESCE(cc.iva,0) AS iva, COALESCE(cc.total,0) AS total,
               COALESCE(cc.debe,0) AS debe, COALESCE(cc.haber,0) AS haber, {extras},
               cc.observaciones, {col_libro} AS libro_iva, {col_origen} AS origen_modulo,
               (SELECT e.razon_social FROM entidades e
                 WHERE REPLACE(e.cuit,'-','') = REPLACE(cc.entidad_id,'-','') LIMIT 1) AS razon
        FROM cuentas_corrientes cc
        {join_as}
        WHERE COALESCE(cc.empresa_id,1)=?
          AND (cc.fecha LIKE ? OR cc.fecha LIKE ? OR cc.fecha LIKE ? OR cc.id IN ({_in_clause(movidos)}));
        """,
        (empresa_id, p1, p2, p3),
    )
    filas = [dict(r) for r in cur.fetchall()]

    comprobantes = [f for f in filas if (f["iva"] or f["neto"] or f["libro_iva"]) and _es_comprobante_iva(f["tipo_comprobante"])]
    imput = {}
    if comprobantes and _tabla_existe(cur, "factura_imputaciones"):
        ids = _in_clause([f["id"] for f in comprobantes])
        cur.execute(
            f"""SELECT cc_id, COALESCE(alicuota_iva,0) AS alic, SUM(COALESCE(neto,0)) AS neto, SUM(COALESCE(iva,0)) AS iva
                FROM factura_imputaciones WHERE cc_id IN ({ids}) GROUP BY cc_id, COALESCE(alicuota_iva,0);"""
        )
        for r in cur.fetchall():
            imput.setdefault(int(r["cc_id"]), []).append({"alicuota": _f(r["alic"]), "neto": _f(r["neto"]), "iva": _f(r["iva"])})

    items = []
    for f in comprobantes:
        libro = f["libro_iva"] or (
            "V" if (_norm(f["observaciones"]).startswith("venta ") or _norm(f["origen_modulo"]) == "ventas") else "C"
        )
        signo = -1 if _es_nc(f["tipo_comprobante"]) else 1
        neto, iva = _f(f["neto"]), _f(f["iva"])
        det = imput.get(int(f["id"]))
        if det and abs(sum(d["neto"] for d in det) - neto) <= max(1.0, abs(neto) * 0.005):
            tot_iva_det = sum(d["iva"] for d in det) or 0
            alicuotas = [
                {"alicuota": d["alicuota"], "neto": _r2(d["neto"] * signo),
                 "iva": _r2((d["iva"] * iva / tot_iva_det if tot_iva_det else d["iva"]) * signo)}
                for d in det
            ]
        else:
            a = _alicuota_inferida(neto, iva)
            alicuotas = [{"alicuota": a, "neto": _r2(neto * signo), "iva": _r2(iva * signo)}]
        base = {
            "fecha": _iso(f["fecha"]), "fecha_original": f["fecha"],
            "periodo_natural": _periodo_de(f["fecha"]),
            "tipo": f["tipo_comprobante"] or "", "letra": _letra(f["tipo_comprobante"]),
            "comprobante": f["numero_comprobante"] or "",
            "cuit": re.sub(r"\D", "", f["entidad_id"] or ""), "razon": f["razon"] or "",
        }
        items.append({
            **base, "origen": "cc", "origen_id": int(f["id"]),
            "seccion": "debito" if libro == "V" else "credito",
            "libro": libro,
            "neto": _r2(neto * signo), "iva": _r2(iva * signo), "total": _r2(_f(f["total"]) * signo),
            "no_gravado": _r2(_f(f["no_gravado"]) * signo), "exento": _r2(_f(f["exento"]) * signo),
            "importe": _r2(iva * signo), "alicuotas": alicuotas,
            "fuente": "Factura de venta" if libro == "V" else "Factura de compra",
        })
        perc = _f(f["percepcion_iva"])
        if perc and libro == "C":
            items.append({
                **base, "origen": "cc_perc", "origen_id": int(f["id"]), "seccion": "percepcion", "libro": "",
                "neto": 0.0, "iva": 0.0, "total": _r2(perc * signo), "importe": _r2(perc * signo),
                "alicuotas": [], "fuente": "Percepción IVA en factura de compra",
            })

    # Retenciones / percepciones / reintegros de IVA registrados en la cuenta corriente.
    for f in filas:
        t = _norm(f["tipo_comprobante"])
        if "iva" not in t:
            continue
        if re.search(r"reint|resint|devol", t):
            seccion, fuente = "devolucion", "Reintegro / devolución ARCA"
        elif re.search(r"\bret", t):
            seccion, fuente = "retencion", "Retención IVA sufrida"
        elif "percep" in t or re.search(r"\bperc\b", t):
            seccion, fuente = "percepcion", "Percepción IVA sufrida"
        else:
            continue
        importe = _f(f["debe"]) or _f(f["haber"]) or _f(f["total"])
        items.append({
            "origen": "cc", "origen_id": int(f["id"]), "seccion": seccion, "libro": "",
            "fecha": _iso(f["fecha"]), "fecha_original": f["fecha"],
            "periodo_natural": _periodo_de(f["fecha"]),
            "tipo": f["tipo_comprobante"] or "", "letra": "", "comprobante": f["numero_comprobante"] or "",
            "cuit": re.sub(r"\D", "", f["entidad_id"] or ""), "razon": f["razon"] or "",
            "neto": 0.0, "iva": 0.0, "total": _r2(importe), "importe": _r2(importe),
            "alicuotas": [], "fuente": fuente,
        })
    return items


def _items_bancos(cur, empresa_id: int, periodo: str, ovr: dict) -> list:
    if not _tabla_existe(cur, "movimientos_cta_cte_bancos"):
        return []
    cols = _cols(cur, "movimientos_cta_cte_bancos")
    filtro_proy = "AND COALESCE(m.generado_por_credito,0)=0" if "generado_por_credito" in cols else ""
    p1, p2, p3 = _patrones_fecha(periodo)
    movidos = _ids_movidos_a(ovr, "banco", periodo)
    fecha_sql = "COALESCE(NULLIF(m.fecha_debito,''), m.fecha_cobro)"
    cur.execute(
        f"""
        SELECT m.id, {fecha_sql} AS fecha, m.proveedor, m.tipo_operacion,
               COALESCE(m.haber,0) AS haber, COALESCE(m.debe,0) AS debe,
               b.banco, b.nro_cta_cte
        FROM movimientos_cta_cte_bancos m
        LEFT JOIN ctas_ctes_bancarias b ON b.id = m.cuenta_id
        WHERE COALESCE(b.empresa_id,1)=? {filtro_proy}
          AND (LOWER(m.proveedor) LIKE '%iva%' OR LOWER(m.proveedor) LIKE '%percep%')
          AND ({fecha_sql} LIKE ? OR {fecha_sql} LIKE ? OR {fecha_sql} LIKE ? OR m.id IN ({_in_clause(movidos)}));
        """,
        (empresa_id, p1, p2, p3),
    )
    hoy = date.today().isoformat()
    items = []
    for r in cur.fetchall():
        prov = _norm(r["proveedor"])
        if not re.search(r"\biva\b", prov):
            continue
        importe = _f(r["haber"]) - _f(r["debe"])
        if importe <= 0 or _iso(r["fecha"]) > hoy:
            continue
        if "percep" in prov or re.search(r"\bperc\b", prov):
            seccion, fuente = "percepcion", "Percepción IVA bancaria (RG 2408)"
        elif re.match(r"^ret", prov):
            seccion, fuente = "retencion", "Retención IVA bancaria"
        else:
            seccion, fuente = "credito", "IVA gastos bancarios (extracto, comprobante tipo 39)"
        banco = f"{r['banco'] or ''} {r['nro_cta_cte'] or ''}".strip()
        items.append({
            "origen": "banco", "origen_id": int(r["id"]), "seccion": seccion, "libro": "C" if seccion == "credito" else "",
            "fecha": _iso(r["fecha"]), "fecha_original": r["fecha"], "periodo_natural": _periodo_de(r["fecha"]),
            "tipo": "Extracto bancario" if seccion == "credito" else (r["proveedor"] or ""),
            "letra": "", "comprobante": r["proveedor"] or "", "cuit": "", "razon": banco,
            "neto": _r2(importe / 0.21) if seccion == "credito" else 0.0,
            "iva": _r2(importe) if seccion == "credito" else 0.0, "total": _r2(importe),
            "importe": _r2(importe),
            "alicuotas": [{"alicuota": 0.21, "neto": _r2(importe / 0.21), "iva": _r2(importe)}] if seccion == "credito" else [],
            "fuente": fuente,
        })
    return items


def _items_liquidaciones(cur, empresa_id: int, periodo: str, ovr: dict) -> list:
    items = []
    p1, p2, p3 = _patrones_fecha(periodo)
    mes_iva = f"{periodo[5:7]}/{periodo[:4]}"

    if _tabla_existe(cur, "liquidaciones_granos"):
        cols = _cols(cur, "liquidaciones_granos")
        if {"mes_iva", "iva_operacion", "subtotal"} <= cols:
            movidos = _ids_movidos_a(ovr, "granos", periodo)
            cur.execute(
                f"""
                SELECT * FROM liquidaciones_granos
                WHERE COALESCE(empresa_id,1)=?
                  AND (mes_iva=? OR (COALESCE(mes_iva,'')='' AND (fecha LIKE ? OR fecha LIKE ? OR fecha LIKE ?))
                       OR id IN ({_in_clause(movidos)}));
                """,
                (empresa_id, mes_iva, p1, p2, p3),
            )
            for r in cur.fetchall():
                r = dict(r)
                fecha = r.get("fecha") or r.get("fecha_pago") or ""
                per = _periodo_de(r.get("mes_iva")) or _periodo_de(fecha)
                base = {
                    "origen": "granos", "origen_id": int(r["id"]), "fecha": _iso(fecha), "fecha_original": fecha,
                    "periodo_natural": per, "letra": "", "comprobante": r.get("nro_lpg") or "",
                    "cuit": re.sub(r"\D", "", r.get("entidad_cuit") or ""), "razon": r.get("entidad_nombre") or "",
                }
                subtotal, iva_op = _f(r.get("subtotal")), _f(r.get("iva_operacion"))
                if iva_op or subtotal:
                    items.append({**base, "seccion": "debito", "libro": "V", "tipo": f"LPG {r.get('grano') or ''}".strip(),
                                  "neto": _r2(subtotal), "iva": _r2(iva_op), "total": _r2(subtotal + iva_op), "importe": _r2(iva_op),
                                  "alicuotas": [{"alicuota": _alicuota_inferida(subtotal, iva_op) or 0.105, "neto": _r2(subtotal), "iva": _r2(iva_op)}],
                                  "fuente": "Liquidación de granos (venta)"})
                iva_serv = _f(r.get("iva_servicios"))
                if iva_serv:
                    neto_serv = _f(r.get("gtos_adm")) + _f(r.get("fletes"))
                    items.append({**base, "seccion": "credito", "libro": "C", "tipo": "LPG servicios",
                                  "neto": _r2(neto_serv), "iva": _r2(iva_serv), "total": _r2(neto_serv + iva_serv), "importe": _r2(iva_serv),
                                  "alicuotas": [{"alicuota": _f(r.get("pct_iva_servicios")) / 100 or 0.105, "neto": _r2(neto_serv), "iva": _r2(iva_serv)}],
                                  "fuente": "Liquidación de granos (gastos del comprador)"})
                ret = _f(r.get("ret_iva"))
                if ret:
                    items.append({**base, "seccion": "retencion", "libro": "", "tipo": "Retención IVA RG 4310",
                                  "neto": 0.0, "iva": 0.0, "total": _r2(ret), "importe": _r2(ret), "alicuotas": [],
                                  "fuente": "Liquidación de granos (retención)"})

    for tabla, origen, fuente in (("liquidaciones_hacienda", "hacienda", "Liquidación de hacienda"),
                                  ("liquidaciones_leche", "leche", "Liquidación de leche")):
        if not _tabla_existe(cur, tabla):
            continue
        movidos = _ids_movidos_a(ovr, origen, periodo)
        cur.execute(
            f"""SELECT * FROM {tabla} WHERE COALESCE(empresa_id,1)=?
                AND (fecha LIKE ? OR fecha LIKE ? OR fecha LIKE ? OR id IN ({_in_clause(movidos)}));""",
            (empresa_id, p1, p2, p3),
        )
        for r in cur.fetchall():
            r = dict(r)
            iva = _f(r.get("iva"))
            if not iva:
                continue
            neto = _f(r.get("bruto"))
            items.append({
                "origen": origen, "origen_id": int(r["id"]), "seccion": "debito", "libro": "V",
                "fecha": _iso(r.get("fecha")), "fecha_original": r.get("fecha"), "periodo_natural": _periodo_de(r.get("fecha")),
                "tipo": fuente, "letra": "", "comprobante": r.get("nro_liquidacion") or "",
                "cuit": re.sub(r"\D", "", r.get("entidad_cuit") or ""), "razon": r.get("entidad_nombre") or "",
                "neto": _r2(neto), "iva": _r2(iva), "total": _r2(neto + iva), "importe": _r2(iva),
                "alicuotas": [{"alicuota": _alicuota_inferida(neto, iva), "neto": _r2(neto), "iva": _r2(iva)}],
                "fuente": fuente + " (venta)",
            })
    return items


def _items_ajustes(cur, empresa_id: int, periodo: str) -> list:
    cur.execute("SELECT * FROM iva_ajustes WHERE empresa_id=? AND periodo=? ORDER BY id;", (empresa_id, periodo))
    items = []
    for r in cur.fetchall():
        tipo = r["tipo"]
        importe = _f(r["importe"])
        items.append({
            "origen": "ajuste", "origen_id": int(r["id"]),
            "seccion": tipo, "libro": "V" if tipo == "debito" else ("C" if tipo == "credito" else ""),
            "fecha": f"{periodo}-01", "fecha_original": "", "periodo_natural": periodo,
            "tipo": "Ajuste manual", "letra": "", "comprobante": r["concepto"] or "", "cuit": "", "razon": r["concepto"] or "",
            "neto": _r2(r["neto"]), "iva": _r2(importe) if tipo in ("debito", "credito") else 0.0,
            "total": _r2(importe), "importe": _r2(importe),
            "alicuotas": [{"alicuota": _f(r["alicuota"]), "neto": _r2(r["neto"]), "iva": _r2(importe)}] if tipo in ("debito", "credito") else [],
            "fuente": "Ajuste manual",
        })
    return items


# ---------------------------------------------------------------- conciliación ARCA

def _pv_num_de_texto(s) -> tuple:
    s = str(s or "")
    m = re.search(r"(\d+)\s*[-/ ]\s*(\d+)\s*$", s)
    if m:
        return int(m.group(1)), int(m.group(2))
    d = re.sub(r"\D", "", s)
    if not d:
        return None, None
    if len(d) > 8:
        return int(d[:-8]), int(d[-8:])
    return None, int(d)


def _mismo_numero(texto, pv: int, numero: int) -> bool:
    p, n = _pv_num_de_texto(texto)
    return n == numero and (p is None or not pv or p == pv)


def _lado_arca(a: dict) -> str:
    venta_liq = a["tipo_codigo"] in iva_arca.TIPOS_VENTA_LIQUIDADA
    if a["origen"] == "R":
        return "V" if venta_liq else "C"
    return "C" if venta_liq else "V"


def _iva_computable_arca(a: dict, lado: str) -> float:
    if lado == "C" and (a.get("letra") or "") not in ("A", "M", ""):
        return 0.0
    return _f(a["iva"])


def _item_arca(a: dict, lado: str) -> dict:
    signo = -1 if a["es_nc"] else 1
    iva = _iva_computable_arca(a, lado)
    alic = [{"alicuota": x.get("alicuota"), "neto": _r2(_f(x["neto"]) * signo), "iva": _r2(_f(x["iva"]) * signo)}
            for x in (a.get("alicuotas") or [])] if iva else []
    return {
        "origen": "arca", "origen_id": int(a["id"]), "seccion": "debito" if lado == "V" else "credito", "libro": lado,
        "fecha": a["fecha"], "fecha_original": a["fecha"], "periodo_natural": a["periodo"],
        "tipo": a["tipo"] or "", "letra": a.get("letra") or "",
        "comprobante": f"{int(a['punto_venta'] or 0):05d}-{int(a['numero'] or 0):08d}",
        "cuit": a["cuit"] or "", "razon": a["denominacion"] or "",
        "neto": _r2(_f(a["neto_gravado"]) * signo), "iva": _r2(iva * signo), "total": _r2(_f(a["total"]) * signo),
        "no_gravado": _r2(_f(a["no_gravado"]) * signo), "exento": _r2(_f(a["exento"]) * signo),
        "importe": _r2(iva * signo), "alicuotas": alic,
        "fuente": "ARCA Mis Comprobantes (no cargado en el programa)",
    }


def _buscar_en_otros_periodos(cur, empresa_id: int, pendientes: list) -> dict:
    """Busca en toda la cuenta corriente los comprobantes de ARCA que no aparecen en el período."""
    cuits = sorted({a["cuit"] for a, _ in pendientes if a["cuit"]})
    if not cuits:
        return {}
    marcas = ",".join("?" for _ in cuits)
    cur.execute(
        f"""SELECT id, fecha, tipo_comprobante, numero_comprobante, COALESCE(iva,0) AS iva,
                   REPLACE(REPLACE(entidad_id,'-',''),' ','') AS cuit
            FROM cuentas_corrientes WHERE COALESCE(empresa_id,1)=? AND REPLACE(REPLACE(entidad_id,'-',''),' ','') IN ({marcas});""",
        (empresa_id, *cuits),
    )
    filas = [dict(r) for r in cur.fetchall() if _es_comprobante_iva(r["tipo_comprobante"])]
    out = {}
    for a, _ in pendientes:
        for f in filas:
            if f["cuit"] == a["cuit"] and _mismo_numero(f["numero_comprobante"], a["punto_venta"], a["numero"]):
                out[a["id"]] = f
                break
    return out


def _conciliar_arca(cur, empresa_id: int, periodo: str, items: list, ovr: dict) -> tuple:
    """Devuelve (conciliacion, items_arca) para el período."""
    comps = iva_arca.comprobantes_periodo(cur, empresa_id, periodo, _ids_movidos_a(ovr, "arca", periodo))
    if not comps:
        return None, []
    lados_importados = {("C" if a["origen"] == "R" else "V") for a in comps if a["periodo"] == periodo}
    prog = [it for it in items if it["origen"] in ("cc", "granos", "hacienda", "leche")
            and it["seccion"] in ("debito", "credito")]
    usados, filas, pendientes = set(), [], []

    for a in comps:
        lado = _lado_arca(a)
        signo = -1 if a["es_nc"] else 1
        cand = [it for it in prog if it["libro"] == lado and id(it) not in usados and it["cuit"] == a["cuit"]]
        match = next((it for it in cand if _mismo_numero(it["comprobante"], a["punto_venta"], a["numero"])), None)
        por_importe = False
        if match is None and a["total"]:
            match = next((it for it in cand if abs(abs(it["total"]) - abs(a["total"])) <= 1), None)
            por_importe = match is not None
        if match is None:
            pendientes.append((a, lado))
            continue
        usados.add(id(match))
        iva_arca_c = _r2(_iva_computable_arca(a, lado) * signo)
        dif = _r2(match["iva"] - iva_arca_c)
        alic_arca = [x for x in (a.get("alicuotas") or []) if x.get("alicuota") is not None]
        completado = False
        if alic_arca and abs(dif) <= 1 and (any(x["alicuota"] is None for x in match["alicuotas"]) or len(alic_arca) > len(match["alicuotas"])):
            match["alicuotas"] = [{"alicuota": x["alicuota"], "neto": _r2(_f(x["neto"]) * signo), "iva": _r2(_f(x["iva"]) * signo)}
                                  for x in alic_arca]
            completado = True
        estado = "ok" if abs(dif) <= 1 else "diferencia"
        notas = []
        if por_importe:
            notas.append("Vinculado por CUIT e importe (el número no coincide).")
        if completado:
            notas.append("Alícuotas completadas con el detalle de ARCA.")
        if match["estado"] in ("excluido", "movido_a"):
            notas.append("En el programa está excluido o movido a otro período.")
        if lado == "C" and a.get("letra") in ("B", "C") and a["iva"]:
            notas.append(f"Comprobante {a['letra']}: el IVA no es computable como crédito fiscal.")
        filas.append({"estado": estado, "lado": lado, "arca": _resumen_arca(a, iva_arca_c),
                      "programa": _resumen_item(match), "dif_iva": dif, "nota": " ".join(notas)})

    otros = _buscar_en_otros_periodos(cur, empresa_id, pendientes)
    items_arca = []
    for a, lado in pendientes:
        signo = -1 if a["es_nc"] else 1
        iva_c = _r2(_iva_computable_arca(a, lado) * signo)
        f = otros.get(a["id"])
        if f and lado == "C":
            filas.append({"estado": "otro_periodo", "lado": lado, "arca": _resumen_arca(a, iva_c),
                          "programa": {"origen": "cc", "origen_id": int(f["id"]), "fecha": _iso(f["fecha"]),
                                       "periodo": _periodo_de(f["fecha"]), "tipo": f["tipo_comprobante"],
                                       "comprobante": f["numero_comprobante"], "iva": _r2(_f(f["iva"]) * signo)},
                          "dif_iva": 0.0, "nota": f"Cargado en la cuenta corriente con fecha {_iso(f['fecha'])}."})
            continue
        it = _item_arca(a, lado)
        items_arca.append(it)
        filas.append({"estado": "solo_arca", "lado": lado, "arca": _resumen_arca(a, iva_c), "programa": None,
                      "dif_iva": iva_c, "nota": "No está cargado en el programa.", "item": {"origen": "arca", "origen_id": int(a["id"])}})

    for it in prog:
        if id(it) in usados or it["libro"] not in lados_importados or it["estado"] not in ("computa", "movido_desde"):
            continue
        filas.append({"estado": "solo_programa", "lado": it["libro"], "arca": None, "programa": _resumen_item(it),
                      "dif_iva": _r2(it["iva"]), "nota": "No figura en Mis Comprobantes de ARCA."})

    orden = {"diferencia": 0, "solo_arca": 1, "otro_periodo": 2, "solo_programa": 3, "ok": 4}
    filas.sort(key=lambda x: (orden[x["estado"]], x["lado"], (x["arca"] or x["programa"] or {}).get("fecha") or ""))
    resumen = {k: sum(1 for x in filas if x["estado"] == k) for k in orden}
    iva_arca_tot = {l: _r2(sum(x["arca"]["iva_computable"] for x in filas if x["arca"] and x["lado"] == l)) for l in ("C", "V")}
    return {"importado": sorted(lados_importados), "resumen": resumen, "iva_arca": iva_arca_tot, "filas": filas}, items_arca


def _resumen_arca(a: dict, iva_computable: float) -> dict:
    signo = -1 if a["es_nc"] else 1
    return {"id": int(a["id"]), "origen": a["origen"], "fecha": a["fecha"], "tipo": a["tipo"], "letra": a.get("letra") or "",
            "comprobante": f"{int(a['punto_venta'] or 0):05d}-{int(a['numero'] or 0):08d}", "cuit": a["cuit"],
            "denominacion": a["denominacion"], "neto": _r2(_f(a["neto_gravado"]) * signo), "iva": _r2(_f(a["iva"]) * signo),
            "iva_computable": iva_computable, "total": _r2(_f(a["total"]) * signo), "alicuotas": a.get("alicuotas") or []}


def _resumen_item(it: dict) -> dict:
    return {"origen": it["origen"], "origen_id": it["origen_id"], "fecha": it["fecha"], "periodo": it["periodo_natural"],
            "tipo": it["tipo"], "comprobante": it["comprobante"], "razon": it["razon"], "neto": it["neto"],
            "iva": it["iva"], "total": it["total"], "estado": it["estado"]}


# ---------------------------------------------------------------- cálculo

def _cabecera(cur, empresa_id: int, periodo: str) -> Optional[dict]:
    cur.execute("SELECT * FROM iva_posiciones WHERE empresa_id=? AND periodo=?;", (empresa_id, periodo))
    r = cur.fetchone()
    return dict(r) if r else None


def _saldos_cierre_anterior(cur, empresa_id: int, periodo: str, memo: dict, profundidad: int) -> Optional[dict]:
    prev = _periodo_mas(periodo, -1)
    cab = _cabecera(cur, empresa_id, prev)
    if not cab or profundidad > 36:
        return None
    if cab.get("estado") == "presentada" and cab.get("snapshot_json"):
        try:
            snap = json.loads(cab["snapshot_json"])["resultado"]
            return {"st_a_favor": snap["st_a_favor"], "sld_a_favor": snap["sld_a_favor"], "periodo": prev, "fuente": "presentada"}
        except Exception:
            pass
    res = calcular_posicion(cur, empresa_id, prev, memo=memo, profundidad=profundidad + 1, con_items=False)["resultado"]
    return {"st_a_favor": res["st_a_favor"], "sld_a_favor": res["sld_a_favor"], "periodo": prev, "fuente": "borrador"}


def _aplicar_override(it: dict, ovr: dict, periodo: str) -> bool:
    """Fija el estado del ítem en el período; False si no corresponde mostrarlo."""
    o = ovr.get((it["origen"], it["origen_id"]))
    it["override_id"] = o["id"] if o else None
    it["motivo"] = (o or {}).get("motivo") or ""
    destino = (o["periodo_destino"] or "") if o else None
    if o is None:
        if it["periodo_natural"] != periodo:
            return False
        # Los reintegros cobrados y los comprobantes que sólo están en ARCA se muestran
        # pero no computan hasta confirmarlos.
        it["estado"] = "informativo" if it["seccion"] == "devolucion" or it["origen"] == "arca" else "computa"
    elif destino == "":
        if it["periodo_natural"] != periodo:
            return False
        it["estado"] = "excluido"
    elif destino == periodo:
        it["estado"] = "movido_desde"
    else:
        if it["periodo_natural"] != periodo:
            return False
        it["estado"] = "movido_a"
        it["periodo_destino"] = destino
    return True


def calcular_posicion(cur, empresa_id: int, periodo: str, memo: Optional[dict] = None,
                      profundidad: int = 0, con_items: bool = True) -> dict:
    memo = {} if memo is None else memo
    clave = (empresa_id, periodo, con_items)
    if clave in memo:
        return memo[clave]
    init_iva_schema(cur)
    ovr = _overrides(cur, empresa_id)

    crudos = (
        _items_cuentas_corrientes(cur, empresa_id, periodo, ovr)
        + _items_bancos(cur, empresa_id, periodo, ovr)
        + _items_liquidaciones(cur, empresa_id, periodo, ovr)
    )
    items = [it for it in crudos if _aplicar_override(it, ovr, periodo)]
    conciliacion, items_arca = _conciliar_arca(cur, empresa_id, periodo, items, ovr)
    items += [it for it in items_arca if _aplicar_override(it, ovr, periodo)]
    items += [dict(it, estado="computa", override_id=None, motivo="") for it in _items_ajustes(cur, empresa_id, periodo)]
    items.sort(key=lambda x: (x["seccion"], x["fecha"], x["comprobante"]))

    computan = [it for it in items if it["estado"] in ("computa", "movido_desde")]
    tot = {s: _r2(sum(it["importe"] for it in computan if it["seccion"] == s)) for s in TIPOS_AJUSTE}

    por_fuente = {}
    for it in computan:
        k = (it["seccion"], it["fuente"])
        por_fuente[k] = por_fuente.get(k, 0) + it["importe"]
    fuentes = [{"seccion": s, "fuente": f, "importe": _r2(v)} for (s, f), v in sorted(por_fuente.items())]

    alic = {}
    for it in computan:
        if it["seccion"] not in ("debito", "credito"):
            continue
        for a in it["alicuotas"] or []:
            k = (it["seccion"], a["alicuota"])
            acc = alic.setdefault(k, {"seccion": it["seccion"], "alicuota": a["alicuota"], "neto": 0.0, "iva": 0.0})
            acc["neto"] += a["neto"]
            acc["iva"] += a["iva"]
    por_alicuota = [dict(v, neto=_r2(v["neto"]), iva=_r2(v["iva"])) for v in sorted(alic.values(), key=lambda x: (x["seccion"], -(x["alicuota"] or -1)))]

    cab = _cabecera(cur, empresa_id, periodo) or {}
    anterior = None
    if cab.get("st_anterior") is None or cab.get("sld_anterior") is None:
        anterior = _saldos_cierre_anterior(cur, empresa_id, periodo, memo, profundidad)
    st_ant = _f(cab["st_anterior"]) if cab.get("st_anterior") is not None else _f((anterior or {}).get("st_a_favor"))
    sld_ant = _f(cab["sld_anterior"]) if cab.get("sld_anterior") is not None else _f((anterior or {}).get("sld_a_favor"))
    pagos_cuenta = _r2(_f(cab.get("pagos_cuenta")) + tot["pago_cuenta"])
    devoluciones = _r2(_f(cab.get("devoluciones")) + tot["devolucion"])

    df, cf = tot["debito"], tot["credito"]
    saldo_tecnico = _r2(df - cf - st_ant)
    impuesto_determinado = max(saldo_tecnico, 0.0)
    st_a_favor = max(-saldo_tecnico, 0.0)
    ingresos_directos = _r2(tot["retencion"] + tot["percepcion"] + pagos_cuenta)
    sld_ant_neto = _r2(sld_ant - devoluciones)
    saldo = _r2(impuesto_determinado - ingresos_directos - sld_ant_neto)
    resultado = {
        "debito_fiscal": df, "credito_fiscal": cf, "st_anterior": _r2(st_ant),
        "saldo_tecnico": saldo_tecnico, "impuesto_determinado": _r2(impuesto_determinado), "st_a_favor": _r2(st_a_favor),
        "retenciones": tot["retencion"], "percepciones": tot["percepcion"], "pagos_cuenta": pagos_cuenta,
        "ingresos_directos": ingresos_directos, "sld_anterior": _r2(sld_ant), "devoluciones": devoluciones,
        "sld_anterior_neto": sld_ant_neto, "saldo": saldo,
        "a_pagar": _r2(max(saldo, 0.0)), "sld_a_favor": _r2(max(-saldo, 0.0)),
    }

    emp = _empresa(cur, empresa_id)
    out = {
        "periodo": periodo,
        "empresa": emp,
        "vencimiento": _vencimiento(periodo, emp["cuit"]),
        "cabecera": {
            "estado": cab.get("estado") or "sin_guardar",
            "st_anterior": cab.get("st_anterior"), "sld_anterior": cab.get("sld_anterior"),
            "pagos_cuenta": _f(cab.get("pagos_cuenta")), "devoluciones": _f(cab.get("devoluciones")),
            "fecha_presentacion": cab.get("fecha_presentacion") or "", "nro_transaccion": cab.get("nro_transaccion") or "",
            "importe_pagado": _f(cab.get("importe_pagado")), "fecha_pago": cab.get("fecha_pago") or "",
            "observaciones": cab.get("observaciones") or "",
        },
        "anterior": anterior,
        "resultado": resultado,
        "totales": tot,
        "por_fuente": fuentes,
        "por_alicuota": por_alicuota,
    }
    if cab.get("estado") == "presentada" and cab.get("snapshot_json"):
        try:
            out["presentado"] = json.loads(cab["snapshot_json"])["resultado"]
        except Exception:
            out["presentado"] = None
    if con_items:
        if conciliacion:
            por_id = {it["origen_id"]: it for it in items if it["origen"] == "arca"}
            for fila in conciliacion["filas"]:
                it = por_id.get((fila.get("item") or {}).get("origen_id"))
                if it:
                    fila["item"].update(estado=it["estado"], override_id=it["override_id"])
        out["conciliacion"] = conciliacion
        out["items"] = items
        out["alertas"] = _alertas(items, out)
    memo[clave] = out
    return out


def _alertas(items: list, pos: dict) -> list:
    al = []
    comp = [it for it in items if it["estado"] in ("computa", "movido_desde")]
    vistos = {}
    for it in comp:
        if it["origen"] != "cc" or it["seccion"] not in ("debito", "credito"):
            continue
        clave = (it["seccion"], it["cuit"], re.sub(r"\D", "", it["comprobante"]).lstrip("0"), it["tipo"].lower())
        if clave in vistos:
            al.append({"nivel": "error", "texto": f"Posible comprobante duplicado: {it['tipo']} {it['comprobante']} de {it['razon'] or it['cuit']}.",
                       "origen": "cc", "origen_id": it["origen_id"]})
        vistos[clave] = it
        if it["seccion"] == "credito" and it["letra"] == "A" and not it["iva"] and it["neto"]:
            al.append({"nivel": "aviso", "texto": f"Factura A sin IVA: {it['comprobante']} de {it['razon'] or it['cuit']}.",
                       "origen": "cc", "origen_id": it["origen_id"]})
        if it["seccion"] == "credito" and it["letra"] in ("B", "C") and it["iva"]:
            al.append({"nivel": "error", "texto": f"{it['tipo']} con IVA discriminado ({it['comprobante']}): una factura {it['letra']} no da crédito fiscal.",
                       "origen": "cc", "origen_id": it["origen_id"]})
        if not it["cuit"]:
            al.append({"nivel": "aviso", "texto": f"Comprobante sin CUIT: {it['tipo']} {it['comprobante']}.",
                       "origen": "cc", "origen_id": it["origen_id"]})
        if any(a["alicuota"] is None for a in it["alicuotas"] or []):
            al.append({"nivel": "info", "texto": f"No se pudo identificar la alícuota de {it['tipo']} {it['comprobante']} (IVA/neto no coincide con 21%, 10,5% ni 27%). Puede incluir percepciones o varias alícuotas.",
                       "origen": "cc", "origen_id": it["origen_id"]})
    if any(it["origen"] == "granos" and it["seccion"] == "retencion" for it in comp) and \
       any(it["origen"] == "cc" and it["seccion"] == "retencion" for it in comp):
        al.append({"nivel": "aviso", "texto": "Hay retenciones de IVA de granos en las liquidaciones y también en la cuenta corriente: revisá que no estén duplicadas."})
    conc = pos.get("conciliacion")
    if conc:
        r = conc["resumen"]
        pend = [it for it in items if it["origen"] == "arca" and it["estado"] == "informativo"]
        if pend:
            al.append({"nivel": "error", "texto": f"ARCA informa {len(pend)} comprobante(s) que no están cargados en el programa "
                                                  f"(IVA $ {_fmt(sum(it['importe'] for it in pend))}). Revisá la solapa Conciliación ARCA."})
        if r["diferencia"]:
            al.append({"nivel": "aviso", "texto": f"{r['diferencia']} comprobante(s) con IVA distinto al informado por ARCA."})
        if r["solo_programa"]:
            al.append({"nivel": "aviso", "texto": f"{r['solo_programa']} comprobante(s) cargados que no figuran en Mis Comprobantes de ARCA."})
    pres = pos.get("presentado")
    if pres and abs(_f(pres.get("saldo")) - _f(pos["resultado"]["saldo"])) > 1:
        al.append({"nivel": "error", "texto": "Los datos cambiaron después de presentar: el saldo recalculado difiere del presentado. Si corresponde, hacé una rectificativa."})
    return al


# ---------------------------------------------------------------- excel

def excel_posicion(pos: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Determinación"
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="EDE9FE")
    fmt = '#,##0.00;[Red]-#,##0.00'
    r = pos["resultado"]
    emp = pos["empresa"]
    filas = [
        (f"Posición de IVA {pos['periodo'][5:7]}/{pos['periodo'][:4]}", None),
        (f"{emp['razon_social']}  CUIT {emp['cuit']}", None),
        (f"Vencimiento estimado: {pos.get('vencimiento') or '-'}", None),
        ("", None),
        ("Débito fiscal", r["debito_fiscal"]),
        ("Crédito fiscal", r["credito_fiscal"]),
        ("Saldo técnico a favor del período anterior", r["st_anterior"]),
        ("Saldo técnico del período", r["saldo_tecnico"]),
        ("Impuesto determinado", r["impuesto_determinado"]),
        ("Saldo técnico a favor (pasa al período siguiente)", r["st_a_favor"]),
        ("", None),
        ("Retenciones sufridas", r["retenciones"]),
        ("Percepciones sufridas", r["percepciones"]),
        ("Pagos a cuenta", r["pagos_cuenta"]),
        ("Saldo de libre disponibilidad del período anterior", r["sld_anterior"]),
        ("Devoluciones / reintegros / transferencias", r["devoluciones"]),
        ("", None),
        ("SALDO A PAGAR", r["a_pagar"]),
        ("SALDO DE LIBRE DISPONIBILIDAD A FAVOR", r["sld_a_favor"]),
    ]
    for i, (txt, val) in enumerate(filas, start=1):
        ws.cell(i, 1, txt)
        if val is not None:
            c = ws.cell(i, 2, val)
            c.number_format = fmt
        if i <= 2 or txt.isupper() or txt in ("Impuesto determinado",):
            ws.cell(i, 1).font = bold
            ws.cell(i, 2).font = bold
    fila = len(filas) + 2
    ws.cell(fila, 1, "Detalle por alícuota").font = bold
    fila += 1
    for j, h in enumerate(("Libro", "Alícuota", "Neto gravado", "IVA"), start=1):
        c = ws.cell(fila, j, h)
        c.font = bold
        c.fill = head_fill
    for a in pos["por_alicuota"]:
        fila += 1
        ws.cell(fila, 1, "Ventas" if a["seccion"] == "debito" else "Compras")
        ws.cell(fila, 2, f"{a['alicuota'] * 100:g}%" if a["alicuota"] is not None else "s/identificar")
        ws.cell(fila, 3, a["neto"]).number_format = fmt
        ws.cell(fila, 4, a["iva"]).number_format = fmt
    ws.column_dimensions["A"].width = 52
    for col in "BCD":
        ws.column_dimensions[col].width = 18

    hojas = (
        ("Libro IVA Compras", lambda it: it["seccion"] == "credito"),
        ("Libro IVA Ventas", lambda it: it["seccion"] == "debito"),
        ("Retenciones y percepciones", lambda it: it["seccion"] in ("retencion", "percepcion", "devolucion")),
    )
    cab = ("Fecha", "Tipo", "Comprobante", "CUIT", "Razón social", "Neto gravado", "Alícuota", "IVA / Importe", "Total", "Fuente", "Estado")
    estados = {"computa": "Computa", "excluido": "Excluido", "movido_a": "Movido a otro período",
               "movido_desde": "Computado en este período", "informativo": "Informativo (no computa)"}
    for titulo, filtro in hojas:
        wsx = wb.create_sheet(titulo)
        for j, h in enumerate(cab, start=1):
            c = wsx.cell(1, j, h)
            c.font = bold
            c.fill = head_fill
            c.alignment = Alignment(horizontal="center")
        i = 1
        for it in (x for x in pos.get("items", []) if filtro(x)):
            i += 1
            alics = ", ".join(f"{a['alicuota'] * 100:g}%" for a in it["alicuotas"] if a["alicuota"] is not None)
            vals = (it["fecha"], it["tipo"], it["comprobante"], it["cuit"], it["razon"], it["neto"], alics,
                    it["importe"], it["total"], it["fuente"], estados.get(it["estado"], it["estado"]))
            for j, v in enumerate(vals, start=1):
                c = wsx.cell(i, j, v)
                if j in (6, 8, 9):
                    c.number_format = fmt
        for j, w in enumerate((11, 22, 18, 14, 34, 16, 10, 16, 16, 40, 22), start=1):
            wsx.column_dimensions[get_column_letter(j)].width = w
        wsx.freeze_panes = "A2"

    conc = pos.get("conciliacion")
    if conc:
        wsc = wb.create_sheet("Conciliación ARCA")
        cab_c = ("Resultado", "Libro", "Fecha", "Tipo", "Comprobante", "CUIT", "Razón social",
                 "IVA ARCA (computable)", "IVA programa", "Diferencia", "Observaciones")
        for j, h in enumerate(cab_c, start=1):
            c = wsc.cell(1, j, h)
            c.font = bold
            c.fill = head_fill
        nombres = {"ok": "Coincide", "diferencia": "Diferencia de IVA", "solo_arca": "Sólo en ARCA",
                   "otro_periodo": "Cargado en otro período", "solo_programa": "Sólo en el programa"}
        for i, f in enumerate(conc["filas"], start=2):
            a, p = f["arca"] or {}, f["programa"] or {}
            vals = (nombres.get(f["estado"], f["estado"]), "Compras" if f["lado"] == "C" else "Ventas",
                    a.get("fecha") or p.get("fecha"), a.get("tipo") or p.get("tipo"), a.get("comprobante") or p.get("comprobante"),
                    a.get("cuit") or "", a.get("denominacion") or p.get("razon") or "",
                    a.get("iva_computable") if a else None, p.get("iva") if p else None, f["dif_iva"], f["nota"])
            for j, v in enumerate(vals, start=1):
                c = wsc.cell(i, j, v)
                if j in (8, 9, 10):
                    c.number_format = fmt
        for j, w in enumerate((24, 10, 11, 26, 18, 14, 34, 18, 16, 16, 50), start=1):
            wsc.column_dimensions[get_column_letter(j)].width = w
        wsc.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- API

class CabeceraModel(BaseModel):
    st_anterior: Optional[float] = None
    sld_anterior: Optional[float] = None
    pagos_cuenta: float = 0
    devoluciones: float = 0
    observaciones: Optional[str] = ""


class PresentarModel(BaseModel):
    fecha_presentacion: str
    nro_transaccion: Optional[str] = ""
    importe_pagado: float = 0
    fecha_pago: Optional[str] = ""


class AjusteModel(BaseModel):
    periodo: str
    tipo: str
    concepto: str
    alicuota: float = 0
    neto: float = 0
    importe: float


class ArcaImportModel(BaseModel):
    nombre: str
    contenido_b64: str
    origen: Optional[str] = None


class ExclusionModel(BaseModel):
    origen: str
    origen_id: int
    periodo_destino: Optional[str] = ""
    motivo: Optional[str] = ""


def _ahora() -> str:
    return datetime.now().isoformat(timespec="seconds")


def _exigir_abierto(cur, empresa_id: int, periodo: str) -> None:
    cab = _cabecera(cur, empresa_id, periodo)
    if cab and cab.get("estado") == "presentada":
        raise HTTPException(400, f"El período {periodo[5:7]}/{periodo[:4]} está presentado. Reabrilo para modificarlo.")


def _upsert_cabecera(cur, empresa_id: int, periodo: str, campos: dict) -> None:
    cur.execute("INSERT OR IGNORE INTO iva_posiciones (empresa_id, periodo, actualizado_en) VALUES (?,?,?);",
                (empresa_id, periodo, _ahora()))
    sets = ", ".join(f"{k}=?" for k in campos) + ", actualizado_en=?"
    cur.execute(f"UPDATE iva_posiciones SET {sets} WHERE empresa_id=? AND periodo=?;",
                (*campos.values(), _ahora(), empresa_id, periodo))


def register_iva_routes(app, get_db, get_empresa_activa_id):
    @app.on_event("startup")
    def _init_iva():
        conn = get_db()
        try:
            init_iva_schema(conn.cursor())
            conn.commit()
        finally:
            conn.close()

    @app.get("/api/iva/posicion")
    def api_posicion(periodo: str = Query(...)):
        periodo = _validar_periodo(periodo)
        conn = get_db()
        try:
            return calcular_posicion(conn.cursor(), get_empresa_activa_id(), periodo)
        finally:
            conn.close()

    @app.put("/api/iva/posicion/{periodo}")
    def api_guardar_cabecera(periodo: str, data: CabeceraModel):
        periodo = _validar_periodo(periodo)
        emp = get_empresa_activa_id()
        conn = get_db()
        try:
            cur = conn.cursor()
            init_iva_schema(cur)
            _exigir_abierto(cur, emp, periodo)
            _upsert_cabecera(cur, emp, periodo, {
                "st_anterior": data.st_anterior, "sld_anterior": data.sld_anterior,
                "pagos_cuenta": data.pagos_cuenta or 0, "devoluciones": data.devoluciones or 0,
                "observaciones": (data.observaciones or "").strip(),
            })
            cab = _cabecera(cur, emp, periodo)
            if cab.get("estado") in (None, "", "sin_guardar"):
                cur.execute("UPDATE iva_posiciones SET estado='borrador' WHERE id=?;", (cab["id"],))
            conn.commit()
            return calcular_posicion(cur, emp, periodo)
        finally:
            conn.close()

    @app.post("/api/iva/posicion/{periodo}/presentar")
    def api_presentar(periodo: str, data: PresentarModel):
        periodo = _validar_periodo(periodo)
        emp = get_empresa_activa_id()
        conn = get_db()
        try:
            cur = conn.cursor()
            init_iva_schema(cur)
            _exigir_abierto(cur, emp, periodo)
            pos = calcular_posicion(cur, emp, periodo, con_items=False)
            snapshot = {"resultado": pos["resultado"], "totales": pos["totales"], "por_alicuota": pos["por_alicuota"],
                        "por_fuente": pos["por_fuente"], "fecha": _ahora()}
            _upsert_cabecera(cur, emp, periodo, {
                "estado": "presentada", "fecha_presentacion": data.fecha_presentacion[:10],
                "nro_transaccion": (data.nro_transaccion or "").strip(), "importe_pagado": data.importe_pagado or 0,
                "fecha_pago": (data.fecha_pago or "")[:10],
                "st_anterior": pos["resultado"]["st_anterior"], "sld_anterior": pos["resultado"]["sld_anterior"],
                "snapshot_json": json.dumps(snapshot, ensure_ascii=False),
            })
            conn.commit()
            return calcular_posicion(cur, emp, periodo)
        finally:
            conn.close()

    @app.post("/api/iva/posicion/{periodo}/reabrir")
    def api_reabrir(periodo: str):
        periodo = _validar_periodo(periodo)
        emp = get_empresa_activa_id()
        conn = get_db()
        try:
            cur = conn.cursor()
            init_iva_schema(cur)
            cur.execute("UPDATE iva_posiciones SET estado='borrador', snapshot_json=NULL, actualizado_en=? WHERE empresa_id=? AND periodo=?;",
                        (_ahora(), emp, periodo))
            conn.commit()
            return calcular_posicion(cur, emp, periodo)
        finally:
            conn.close()

    @app.post("/api/iva/ajustes")
    def api_ajuste(data: AjusteModel):
        periodo = _validar_periodo(data.periodo)
        if data.tipo not in TIPOS_AJUSTE:
            raise HTTPException(400, "Tipo de ajuste inválido.")
        if not (data.concepto or "").strip():
            raise HTTPException(400, "Indicá el concepto.")
        emp = get_empresa_activa_id()
        conn = get_db()
        try:
            cur = conn.cursor()
            init_iva_schema(cur)
            _exigir_abierto(cur, emp, periodo)
            cur.execute(
                "INSERT INTO iva_ajustes (empresa_id, periodo, tipo, concepto, alicuota, neto, importe, creado_en) VALUES (?,?,?,?,?,?,?,?);",
                (emp, periodo, data.tipo, data.concepto.strip(), data.alicuota or 0, data.neto or 0, data.importe, _ahora()),
            )
            conn.commit()
            return {"id": cur.lastrowid}
        finally:
            conn.close()

    @app.delete("/api/iva/ajustes/{ajuste_id}")
    def api_borrar_ajuste(ajuste_id: int):
        emp = get_empresa_activa_id()
        conn = get_db()
        try:
            cur = conn.cursor()
            init_iva_schema(cur)
            cur.execute("SELECT periodo FROM iva_ajustes WHERE id=? AND empresa_id=?;", (ajuste_id, emp))
            r = cur.fetchone()
            if not r:
                raise HTTPException(404, "Ajuste no encontrado.")
            _exigir_abierto(cur, emp, r["periodo"])
            cur.execute("DELETE FROM iva_ajustes WHERE id=?;", (ajuste_id,))
            conn.commit()
            return {"status": "ok"}
        finally:
            conn.close()

    @app.post("/api/iva/exclusiones")
    def api_exclusion(data: ExclusionModel):
        if data.origen not in ORIGENES_ITEMS:
            raise HTTPException(400, "Origen inválido.")
        destino = (data.periodo_destino or "").strip()
        if destino:
            destino = _validar_periodo(destino)
        emp = get_empresa_activa_id()
        conn = get_db()
        try:
            cur = conn.cursor()
            init_iva_schema(cur)
            ovr = _overrides(cur, emp)
            previo = ovr.get((data.origen, data.origen_id))
            if previo and previo["periodo_destino"]:
                _exigir_abierto(cur, emp, previo["periodo_destino"])
            if destino:
                _exigir_abierto(cur, emp, destino)
            cur.execute(
                """INSERT INTO iva_exclusiones (empresa_id, origen, origen_id, periodo_destino, motivo, creado_en)
                   VALUES (?,?,?,?,?,?)
                   ON CONFLICT(empresa_id, origen, origen_id) DO UPDATE SET
                     periodo_destino=excluded.periodo_destino, motivo=excluded.motivo, creado_en=excluded.creado_en;""",
                (emp, data.origen, data.origen_id, destino, (data.motivo or "").strip(), _ahora()),
            )
            conn.commit()
            return {"status": "ok"}
        finally:
            conn.close()

    @app.delete("/api/iva/exclusiones/{exclusion_id}")
    def api_borrar_exclusion(exclusion_id: int):
        emp = get_empresa_activa_id()
        conn = get_db()
        try:
            cur = conn.cursor()
            init_iva_schema(cur)
            cur.execute("SELECT periodo_destino FROM iva_exclusiones WHERE id=? AND empresa_id=?;", (exclusion_id, emp))
            r = cur.fetchone()
            if not r:
                raise HTTPException(404, "No encontrado.")
            if r["periodo_destino"]:
                _exigir_abierto(cur, emp, r["periodo_destino"])
            cur.execute("DELETE FROM iva_exclusiones WHERE id=?;", (exclusion_id,))
            conn.commit()
            return {"status": "ok"}
        finally:
            conn.close()

    @app.get("/api/iva/resumen")
    def api_resumen(anio: int = Query(...)):
        if anio < 2000 or anio > 2100:
            raise HTTPException(400, "Año inválido.")
        emp = get_empresa_activa_id()
        conn = get_db()
        try:
            cur = conn.cursor()
            memo = {}
            meses = []
            for m in range(1, 13):
                p = f"{anio:04d}-{m:02d}"
                pos = calcular_posicion(cur, emp, p, memo=memo, con_items=False)
                meses.append({"periodo": p, "estado": pos["cabecera"]["estado"], "vencimiento": pos["vencimiento"],
                              **pos["resultado"], "importe_pagado": pos["cabecera"]["importe_pagado"]})
            return {"anio": anio, "meses": meses}
        finally:
            conn.close()

    @app.post("/api/iva/arca/importar")
    def api_arca_importar(data: ArcaImportModel):
        try:
            contenido = base64.b64decode(data.contenido_b64.split(",")[-1])
        except Exception:
            raise HTTPException(400, "Archivo inválido.")
        if not contenido:
            raise HTTPException(400, "El archivo está vacío.")
        try:
            res = iva_arca.parsear_mis_comprobantes(data.nombre or "archivo.csv", contenido,
                                                    data.origen if data.origen in ("R", "E") else None)
        except Exception as e:
            raise HTTPException(400, f"No se pudo leer el archivo: {e}")
        if not res["registros"]:
            raise HTTPException(400, "No se encontraron comprobantes. " + " ".join(res["avisos"]))
        emp = get_empresa_activa_id()
        conn = get_db()
        try:
            cur = conn.cursor()
            init_iva_schema(cur)
            periodos = sorted({r["periodo"] for r in res["registros"]})
            bloqueados = [p for p in periodos if (_cabecera(cur, emp, p) or {}).get("estado") == "presentada"]
            guardado = iva_arca.guardar_comprobantes(cur, emp, res["registros"])
            conn.commit()
            return {**guardado, "origen": res["origen"], "periodos": periodos, "avisos": res["avisos"],
                    "presentados": bloqueados, "total": len(res["registros"])}
        finally:
            conn.close()

    @app.get("/api/iva/arca/importaciones")
    def api_arca_importaciones():
        conn = get_db()
        try:
            cur = conn.cursor()
            out = iva_arca.resumen_importaciones(cur, get_empresa_activa_id())
            conn.commit()
            return out
        finally:
            conn.close()

    @app.delete("/api/iva/arca")
    def api_arca_borrar(periodo: str = Query(...), origen: str = Query(...)):
        periodo = _validar_periodo(periodo)
        if origen not in ("R", "E"):
            raise HTTPException(400, "Origen inválido (R = recibidos, E = emitidos).")
        emp = get_empresa_activa_id()
        conn = get_db()
        try:
            cur = conn.cursor()
            init_iva_schema(cur)
            _exigir_abierto(cur, emp, periodo)
            n = iva_arca.borrar_periodo(cur, emp, periodo, origen)
            conn.commit()
            return {"borrados": n}
        finally:
            conn.close()

    @app.get("/api/iva/posicion/{periodo}/excel")
    def api_excel(periodo: str):
        periodo = _validar_periodo(periodo)
        conn = get_db()
        try:
            pos = calcular_posicion(conn.cursor(), get_empresa_activa_id(), periodo)
        finally:
            conn.close()
        return Response(
            excel_posicion(pos),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="posicion_iva_{periodo}.xlsx"'},
        )
