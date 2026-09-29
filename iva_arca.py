# -*- coding: utf-8 -*-
"""Importación de "Mis Comprobantes" de ARCA (Recibidos y Emitidos).

Acepta el ZIP que descarga ARCA, el CSV o el Excel. Reconoce las columnas por nombre,
tanto en el formato anterior (un solo neto e IVA) como en el de septiembre de 2025
(punto de venta y número juntos, neto e IVA por alícuota).
"""
from __future__ import annotations

import csv
import io
import json
import re
import unicodedata
import zipfile
from datetime import date, datetime
from typing import Optional

# Código ARCA → (nombre, letra). Las notas de crédito restan.
TIPOS = {
    1: ("Factura", "A"), 2: ("Nota de Débito", "A"), 3: ("Nota de Crédito", "A"), 4: ("Recibo", "A"),
    5: ("Nota de Venta al contado", "A"), 6: ("Factura", "B"), 7: ("Nota de Débito", "B"), 8: ("Nota de Crédito", "B"),
    9: ("Recibo", "B"), 11: ("Factura", "C"), 12: ("Nota de Débito", "C"), 13: ("Nota de Crédito", "C"), 15: ("Recibo", "C"),
    19: ("Factura", "E"), 20: ("Nota de Débito", "E"), 21: ("Nota de Crédito", "E"),
    33: ("Liquidación Primaria de Granos", ""), 331: ("Liquidación Secundaria de Granos", ""),
    51: ("Factura", "M"), 52: ("Nota de Débito", "M"), 53: ("Nota de Crédito", "M"), 54: ("Recibo", "M"),
    60: ("Cuenta de Venta y Líquido Producto", "A"), 61: ("Cuenta de Venta y Líquido Producto", "B"),
    63: ("Liquidación", "A"), 64: ("Liquidación", "B"),
    81: ("Tique Factura", "A"), 82: ("Tique Factura", "B"), 83: ("Tique", ""), 111: ("Tique Factura", "C"),
    112: ("Tique Nota de Crédito", "A"), 113: ("Tique Nota de Crédito", "B"), 114: ("Tique Nota de Crédito", "C"),
    118: ("Tique Factura", "M"), 119: ("Tique Nota de Crédito", "M"),
    201: ("Factura de Crédito Electrónica MiPyME", "A"), 202: ("Nota de Débito Electrónica MiPyME", "A"),
    203: ("Nota de Crédito Electrónica MiPyME", "A"), 206: ("Factura de Crédito Electrónica MiPyME", "B"),
    207: ("Nota de Débito Electrónica MiPyME", "B"), 208: ("Nota de Crédito Electrónica MiPyME", "B"),
    211: ("Factura de Crédito Electrónica MiPyME", "C"), 212: ("Nota de Débito Electrónica MiPyME", "C"),
    213: ("Nota de Crédito Electrónica MiPyME", "C"),
}
# Comprobantes que emite el comprador por una venta nuestra (granos, consignación, liquidaciones).
TIPOS_VENTA_LIQUIDADA = {33, 331, 60, 61, 63, 64}
ALICUOTAS_ARCA = (0.0, 0.025, 0.05, 0.105, 0.21, 0.27)


def init_arca_schema(cur) -> None:
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS arca_comprobantes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            empresa_id INTEGER NOT NULL,
            origen TEXT NOT NULL,
            fecha TEXT,
            periodo TEXT,
            tipo_codigo INTEGER,
            tipo TEXT,
            letra TEXT,
            es_nc INTEGER DEFAULT 0,
            punto_venta INTEGER,
            numero INTEGER,
            cae TEXT,
            cuit TEXT,
            denominacion TEXT,
            moneda TEXT,
            tipo_cambio REAL DEFAULT 1,
            neto_gravado REAL DEFAULT 0,
            no_gravado REAL DEFAULT 0,
            exento REAL DEFAULT 0,
            otros_tributos REAL DEFAULT 0,
            iva REAL DEFAULT 0,
            total REAL DEFAULT 0,
            alicuotas_json TEXT,
            archivo TEXT,
            importado_en TEXT,
            UNIQUE(empresa_id, origen, cuit, tipo_codigo, punto_venta, numero)
        );
        """
    )
    cur.execute("CREATE INDEX IF NOT EXISTS idx_arca_comp_periodo ON arca_comprobantes(empresa_id, periodo);")


def _norm(s) -> str:
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode().lower()
    s = re.sub(r"[.:()\[\]$]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def _num(v) -> float:
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("$", "").replace(" ", "")
    if not s:
        return 0.0
    neg = s.startswith("-") or (s.startswith("(") and s.endswith(")"))
    s = s.strip("-()")
    coma, punto = s.rfind(","), s.rfind(".")
    if coma >= 0 and punto >= 0:
        dec = max(coma, punto)
        s = re.sub(r"[.,]", "", s[:dec]) + "." + re.sub(r"[.,]", "", s[dec + 1:])
    elif coma >= 0:
        s = s.replace(".", "").replace(",", ".")
    elif s.count(".") > 1:
        s = s.replace(".", "")
    try:
        n = float(s)
    except ValueError:
        return 0.0
    return -n if neg else n


def _fecha(v) -> str:
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v or "").strip()
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return m.group(0)
    m = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{4})", s)
    if m:
        return f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"
    return ""


def _tasa(h: str) -> Optional[float]:
    m = re.search(r"(\d+(?:[,.]\d+)?)\s*%", h)
    if not m:
        return None
    t = float(m.group(1).replace(",", ".")) / 100
    return min(ALICUOTAS_ARCA, key=lambda a: abs(a - t))


def _mapear_columnas(headers: list) -> dict:
    col = {"neto_tasa": {}, "iva_tasa": {}}
    for i, raw in enumerate(headers):
        h = _norm(raw)
        if not h:
            continue
        tasa = _tasa(h)
        if tasa is not None and ("iva" in h or "neto" in h):
            (col["neto_tasa"] if "neto" in h else col["iva_tasa"])[tasa] = i
        elif h.startswith("fecha"):
            col.setdefault("fecha", i)
        elif h in ("tipo", "tipo comprobante", "tipo de comprobante") or h.startswith("tipo de comprobante"):
            col.setdefault("tipo", i)
        elif "punto de venta" in h or h in ("pto vta", "punto venta"):
            col.setdefault("pv", i)
        elif h.startswith("numero desde") or h.startswith("nro desde") or h in ("numero", "nro", "numero de comprobante"):
            col.setdefault("numero", i)
        elif "autorizacion" in h or h == "cae":
            col.setdefault("cae", i)
        elif ("doc" in h or "documento" in h) and ("emisor" in h or "receptor" in h) and not h.startswith("tipo"):
            col.setdefault("doc", i)
            col["lado"] = "emisor" if "emisor" in h else "receptor"
        elif h.startswith("denominacion"):
            col.setdefault("denominacion", i)
            col.setdefault("lado", "emisor" if "emisor" in h else ("receptor" if "receptor" in h else None))
        elif "tipo cambio" in h or "tipo de cambio" in h:
            col.setdefault("tc", i)
        elif h == "moneda":
            col.setdefault("moneda", i)
        elif "no gravado" in h:
            col.setdefault("no_gravado", i)
        elif "exent" in h:
            col.setdefault("exento", i)
        elif "neto gravado" in h:
            col.setdefault("neto_gravado", i)
        elif "otros tributos" in h:
            col.setdefault("otros_tributos", i)
        elif h in ("iva", "total iva", "imp iva", "importe iva"):
            col.setdefault("iva", i)
        elif h in ("imp total", "importe total", "total") or h.startswith("imp total"):
            col.setdefault("total", i)
    return col


def _filas_de_archivo(nombre: str, contenido: bytes) -> list:
    """Devuelve una lista de (nombre, filas) con filas como listas de celdas."""
    ext = nombre.lower().rsplit(".", 1)[-1] if "." in nombre else ""
    if ext == "zip" or contenido[:2] == b"PK" and ext not in ("xlsx", "xlsm"):
        salida = []
        with zipfile.ZipFile(io.BytesIO(contenido)) as z:
            for info in z.infolist():
                n = info.filename
                if n.lower().endswith((".csv", ".xlsx", ".txt")) and not info.is_dir():
                    salida += _filas_de_archivo(n, z.read(info))
        return salida
    if ext in ("xlsx", "xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        return [(nombre, [list(r) for r in ws.iter_rows(values_only=True)])]
    for enc in ("utf-8-sig", "latin-1"):
        try:
            texto = contenido.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    muestra = texto[:4000]
    delim = max((";", ",", "\t"), key=lambda d: muestra.count(d))
    return [(nombre, list(csv.reader(io.StringIO(texto), delimiter=delim)))]


def _parse_tipo(v) -> tuple:
    s = str(v or "").strip()
    m = re.match(r"^\s*(\d+)", s)
    codigo = int(m.group(1)) if m else None
    if codigo is None:
        n = _norm(s)
        for c, (nom, letra) in TIPOS.items():
            if n == _norm(f"{nom} {letra}".strip()):
                codigo = c
                break
    nom, letra = TIPOS.get(codigo, (re.sub(r"^\s*\d+\s*-\s*", "", s) or f"Tipo {codigo}", ""))
    if not letra:
        m2 = re.search(r"\b([ABCEM])\s*$", s)
        letra = m2.group(1) if m2 else ""
    nombre = f"{nom} {letra}".strip()
    es_nc = "credito" in _norm(nom)
    return codigo, nombre, letra, es_nc


def _pv_numero(pv_raw, num_raw) -> tuple:
    pv_s = str(pv_raw if pv_raw is not None else "").strip()
    m = re.match(r"^\s*(\d+)\s*-\s*(\d+)\s*$", pv_s)
    if m:
        return int(m.group(1)), int(m.group(2))
    pv = int(_num(pv_s)) if pv_s else 0
    num_s = str(num_raw if num_raw is not None else "").strip()
    m = re.match(r"^\s*(\d+)\s*-\s*(\d+)\s*$", num_s)
    if m:
        return int(m.group(1)), int(m.group(2))
    return pv, int(_num(num_s)) if num_s else 0


def parsear_mis_comprobantes(nombre: str, contenido: bytes, origen_forzado: Optional[str] = None) -> dict:
    registros, avisos = [], []
    origen_detectado = None
    for arch, filas in _filas_de_archivo(nombre, contenido):
        idx_head = None
        for i, fila in enumerate(filas[:10]):
            hs = [_norm(c) for c in fila]
            if any(h.startswith("fecha") for h in hs) and any(h.startswith("tipo") for h in hs):
                idx_head = i
                break
        if idx_head is None:
            avisos.append(f"{arch}: no se encontró la fila de encabezados (Fecha, Tipo...).")
            continue
        col = _mapear_columnas(filas[idx_head])
        if "fecha" not in col or "tipo" not in col or ("total" not in col and not col["iva_tasa"]):
            avisos.append(f"{arch}: faltan columnas obligatorias (Fecha, Tipo, Importe total).")
            continue
        titulo = " ".join(_norm(c) for fila in filas[:idx_head] for c in fila if c) + " " + _norm(arch)
        if origen_forzado in ("R", "E"):
            origen = origen_forzado
        elif "recibid" in titulo or col.get("lado") == "emisor":
            origen = "R"
        elif "emitid" in titulo or col.get("lado") == "receptor":
            origen = "E"
        else:
            avisos.append(f"{arch}: no se pudo saber si son Recibidos o Emitidos; se tomó Recibidos.")
            origen = "R"
        origen_detectado = origen_detectado or origen

        def celda(fila, clave):
            i = col.get(clave)
            return fila[i] if i is not None and i < len(fila) else None

        for fila in filas[idx_head + 1:]:
            if not fila or not any(c not in (None, "") for c in fila):
                continue
            fecha = _fecha(celda(fila, "fecha"))
            if not fecha:
                continue
            codigo, tipo, letra, es_nc = _parse_tipo(celda(fila, "tipo"))
            pv, numero = _pv_numero(celda(fila, "pv"), celda(fila, "numero"))
            tc = _num(celda(fila, "tc")) or 1.0
            moneda = str(celda(fila, "moneda") or "").strip()
            factor = tc if moneda and _norm(moneda) not in ("$", "pes", "ars", "pesos", "") else 1.0
            alic = []
            for t in sorted(set(col["neto_tasa"]) | set(col["iva_tasa"])):
                neto_t = _num(fila[col["neto_tasa"][t]]) if t in col["neto_tasa"] and col["neto_tasa"][t] < len(fila) else 0.0
                iva_t = _num(fila[col["iva_tasa"][t]]) if t in col["iva_tasa"] and col["iva_tasa"][t] < len(fila) else 0.0
                if neto_t or iva_t:
                    alic.append({"alicuota": t, "neto": round(neto_t * factor, 2), "iva": round(iva_t * factor, 2)})
            neto_g = _num(celda(fila, "neto_gravado")) * factor or sum(a["neto"] for a in alic)
            iva = _num(celda(fila, "iva")) * factor or sum(a["iva"] for a in alic)
            if not alic and (neto_g or iva):
                ratio = iva / neto_g if neto_g else 0
                t = min(ALICUOTAS_ARCA, key=lambda a: abs(a - ratio)) if neto_g else None
                alic.append({"alicuota": t if t is not None and abs(t - ratio) <= 0.0015 else None,
                             "neto": round(neto_g, 2), "iva": round(iva, 2)})
            doc = re.sub(r"\D", "", str(celda(fila, "doc") or ""))
            registros.append({
                "origen": origen, "fecha": fecha, "periodo": fecha[:7], "tipo_codigo": codigo, "tipo": tipo,
                "letra": letra, "es_nc": 1 if es_nc else 0, "punto_venta": pv, "numero": numero,
                "cae": str(celda(fila, "cae") or "").strip(), "cuit": doc,
                "denominacion": str(celda(fila, "denominacion") or "").strip(),
                "moneda": moneda or "$", "tipo_cambio": tc,
                "neto_gravado": round(neto_g, 2), "no_gravado": round(_num(celda(fila, "no_gravado")) * factor, 2),
                "exento": round(_num(celda(fila, "exento")) * factor, 2),
                "otros_tributos": round(_num(celda(fila, "otros_tributos")) * factor, 2),
                "iva": round(iva, 2), "total": round(_num(celda(fila, "total")) * factor, 2),
                "alicuotas": alic, "archivo": arch,
            })
    return {"origen": origen_detectado, "registros": registros, "avisos": avisos}


def guardar_comprobantes(cur, empresa_id: int, registros: list) -> dict:
    init_arca_schema(cur)
    nuevos = actualizados = 0
    ahora = datetime.now().isoformat(timespec="seconds")
    for r in registros:
        cur.execute(
            """SELECT id FROM arca_comprobantes WHERE empresa_id=? AND origen=? AND cuit=?
               AND tipo_codigo IS ? AND punto_venta=? AND numero=?;""",
            (empresa_id, r["origen"], r["cuit"], r["tipo_codigo"], r["punto_venta"], r["numero"]),
        )
        existe = cur.fetchone()
        valores = (
            r["fecha"], r["periodo"], r["tipo"], r["letra"], r["es_nc"], r["cae"], r["denominacion"], r["moneda"],
            r["tipo_cambio"], r["neto_gravado"], r["no_gravado"], r["exento"], r["otros_tributos"], r["iva"], r["total"],
            json.dumps(r["alicuotas"]), r["archivo"], ahora,
        )
        if existe:
            cur.execute(
                """UPDATE arca_comprobantes SET fecha=?, periodo=?, tipo=?, letra=?, es_nc=?, cae=?, denominacion=?, moneda=?,
                   tipo_cambio=?, neto_gravado=?, no_gravado=?, exento=?, otros_tributos=?, iva=?, total=?, alicuotas_json=?,
                   archivo=?, importado_en=? WHERE id=?;""",
                valores + (existe[0],),
            )
            actualizados += 1
        else:
            cur.execute(
                """INSERT INTO arca_comprobantes (fecha, periodo, tipo, letra, es_nc, cae, denominacion, moneda, tipo_cambio,
                   neto_gravado, no_gravado, exento, otros_tributos, iva, total, alicuotas_json, archivo, importado_en,
                   empresa_id, origen, cuit, tipo_codigo, punto_venta, numero)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?);""",
                valores + (empresa_id, r["origen"], r["cuit"], r["tipo_codigo"], r["punto_venta"], r["numero"]),
            )
            nuevos += 1
    return {"nuevos": nuevos, "actualizados": actualizados}


def comprobantes_periodo(cur, empresa_id: int, periodo: str, ids_extra: Optional[list] = None) -> list:
    init_arca_schema(cur)
    extra = ",".join(str(int(i)) for i in (ids_extra or [])) or "NULL"
    cur.execute(
        f"SELECT * FROM arca_comprobantes WHERE empresa_id=? AND (periodo=? OR id IN ({extra})) ORDER BY fecha, id;",
        (empresa_id, periodo),
    )
    out = []
    for r in cur.fetchall():
        d = dict(r)
        try:
            d["alicuotas"] = json.loads(d.pop("alicuotas_json") or "[]")
        except ValueError:
            d["alicuotas"] = []
        out.append(d)
    return out


def resumen_importaciones(cur, empresa_id: int) -> list:
    init_arca_schema(cur)
    cur.execute(
        """SELECT periodo, origen, COUNT(*) AS n, ROUND(SUM(CASE WHEN es_nc=1 THEN -iva ELSE iva END),2) AS iva,
                  MAX(importado_en) AS importado_en
           FROM arca_comprobantes WHERE empresa_id=? GROUP BY periodo, origen ORDER BY periodo DESC, origen;""",
        (empresa_id,),
    )
    return [dict(r) for r in cur.fetchall()]


def borrar_periodo(cur, empresa_id: int, periodo: str, origen: str) -> int:
    init_arca_schema(cur)
    cur.execute("DELETE FROM arca_comprobantes WHERE empresa_id=? AND periodo=? AND origen=?;", (empresa_id, periodo, origen))
    return cur.rowcount
